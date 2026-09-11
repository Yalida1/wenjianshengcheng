from __future__ import annotations

import io

from docx import Document
from fastapi.testclient import TestClient
from pytest import MonkeyPatch
from sqlalchemy import select

import backend.scripts.seed as seed_module
from backend.app.db import SessionLocal
from backend.app.models import TemplateVersion
from backend.app.services.storage import get_storage, sha256_bytes
from backend.tests.test_document_flow import (
    _confirm_required_fields,
    _create_project,
    _generate,
    _review_all_blocks,
)


def _docx_bytes() -> bytes:
    stream = io.BytesIO()
    document = Document()
    document.add_heading("来源材料", level=1)
    document.add_paragraph("项目名称：扩展接口测试项目")
    document.save(stream)
    return stream.getvalue()


def _feasibility_docx_bytes() -> bytes:
    stream = io.BytesIO()
    document = Document()
    document.add_paragraph("北京地铁智能制度审查项目")
    document.add_paragraph("可行性研究报告")
    document.add_heading("项目概况", level=1)
    document.add_paragraph("项目名称为“北京地铁智能制度审查项目”，拟建设制度智能审查平台。")
    document.add_paragraph("方案测算建设投资330.00万元，上述数字不是已批复招标预算。")
    document.add_paragraph("项目全部建设范围包括知识底座、规则审查和人工复核能力。")
    document.save(stream)
    return stream.getvalue()


def test_file_versions_parse_results_and_stage_source_options(
    authenticated_client: TestClient,
) -> None:
    project_id = _create_project(authenticated_client)
    uploaded = authenticated_client.post(
        f"/api/v1/files?project_id={project_id}&stage=requirement",
        files={
            "upload": (
                "source.docx",
                _docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert uploaded.status_code == 201, uploaded.text
    file_id = uploaded.json()["id"]
    versions = authenticated_client.get(f"/api/v1/files/{file_id}/versions")
    assert versions.status_code == 200
    assert versions.json()[0]["sha256"]
    jobs = authenticated_client.get(f"/api/v1/files/{file_id}/parse-jobs").json()
    result = authenticated_client.get(f"/api/v1/parse-jobs/{jobs[0]['id']}/result")
    assert result.status_code == 200, result.text
    assert result.json()["blocks"][0]["text"] == "来源材料"
    sources = authenticated_client.get(f"/api/v1/projects/{project_id}/stages/requirement/sources")
    assert any(item["id"] == versions.json()[0]["id"] for item in sources.json())


def test_parsed_feasibility_file_creates_only_traceable_tender_candidates(
    authenticated_client: TestClient,
) -> None:
    project_id = _create_project(authenticated_client)
    uploaded = authenticated_client.post(
        f"/api/v1/files?project_id={project_id}&stage=tender",
        files={
            "upload": (
                "可行性研究报告.docx",
                _feasibility_docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert uploaded.status_code == 201, uploaded.text
    assert uploaded.json()["status"] == "parsed"

    fields = authenticated_client.get(f"/api/v1/field-values?project_id={project_id}&stage=tender")
    assert fields.status_code == 200, fields.text
    assert [field["field_key"] for field in fields.json()] == ["project_name"]
    project_name = fields.json()[0]
    assert project_name["value"] == "北京地铁智能制度审查项目"
    assert project_name["status"] == "extracted"
    assert project_name["source_type"] == "extracted"
    assert project_name["evidence"][0]["source_file_id"] == uploaded.json()["id"]
    assert "项目名称" in project_name["evidence"][0]["excerpt"]

    not_acknowledged = authenticated_client.post(
        f"/api/v1/field-values/{project_name['id']}/confirm",
        json={"revision": project_name["revision"], "evidence_acknowledged": False},
    )
    assert not_acknowledged.status_code == 422
    assert not_acknowledged.json()["error"]["code"] == "evidence_not_acknowledged"

    refreshed = authenticated_client.post(f"/api/v1/files/{uploaded.json()['id']}/field-candidates")
    assert refreshed.status_code == 200, refreshed.text
    assert refreshed.json()["created"] == 0
    assert refreshed.json()["existing_keys"] == ["project_name"]


def test_no_bid_bond_confirmation_requires_and_retains_basis(
    authenticated_client: TestClient,
) -> None:
    project_id = _create_project(authenticated_client)
    created = authenticated_client.post(
        f"/api/v1/field-values?project_id={project_id}&stage=tender",
        json={
            "field_key": "doc::group-1::bid_bond",
            "field_label": "投标保证金",
            "data_type": "text",
            "value": "本项目不要求投标保证金",
            "normalized_value": "本项目不要求投标保证金",
            "criticality": "P0",
            "status": "missing",
            "source_type": "user_input",
        },
    )
    assert created.status_code == 201, created.text

    blocked = authenticated_client.post(
        f"/api/v1/field-values/{created.json()['id']}/confirm",
        json={"revision": created.json()["revision"], "evidence_acknowledged": False},
    )
    assert blocked.status_code == 422, blocked.text
    assert blocked.json()["error"]["code"] == "bid_bond_basis_required"

    with_basis = authenticated_client.patch(
        f"/api/v1/field-values/{created.json()['id']}",
        json={
            "value": "本项目不要求投标保证金",
            "normalized_value": "本项目不要求投标保证金",
            "unit": None,
            "status": "missing",
            "source_type": "user_input",
            "revision": created.json()["revision"],
            "evidence": {
                "excerpt": "采购审批意见第 6 条明确本项目不收取投标保证金",
                "extraction_method": "manual",
                "section_path": "投标保证金确认依据",
            },
        },
    )
    assert with_basis.status_code == 200, with_basis.text
    assert with_basis.json()["evidence"][0]["excerpt"].startswith("采购审批意见")

    confirmed = authenticated_client.post(
        f"/api/v1/field-values/{with_basis.json()['id']}/confirm",
        json={"revision": with_basis.json()["revision"], "evidence_acknowledged": False},
    )
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["status"] == "user_confirmed"


def test_single_feasibility_upload_automatically_generates_controlled_tender_draft(
    authenticated_client: TestClient,
) -> None:
    project_id = _create_project(authenticated_client)
    uploaded = authenticated_client.post(
        f"/api/v1/files?project_id={project_id}&stage=tender&auto_generate_draft=true",
        files={
            "upload": (
                "可行性研究报告.docx",
                _feasibility_docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert uploaded.status_code == 201, uploaded.text
    file_record = uploaded.json()
    assert file_record["status"] == "parsed"
    assert file_record["auto_generate_draft"] is True
    assert file_record["auto_generation_error"] is None
    assert file_record["auto_generation_job_id"]

    generation = authenticated_client.get(f"/api/v1/generation-jobs/{file_record['auto_generation_job_id']}")
    assert generation.status_code == 200, generation.text
    assert generation.json()["status"] == "succeeded"
    repeated = authenticated_client.post(f"/api/v1/files/{file_record['id']}/auto-draft")
    assert repeated.status_code == 202, repeated.text
    assert repeated.json()["id"] == generation.json()["id"]

    stages = authenticated_client.get(f"/api/v1/projects/{project_id}/stages").json()
    tender_stage = next(item for item in stages if item["stage"] == "tender")
    versions = authenticated_client.get(f"/api/v1/files/{file_record['id']}/versions").json()
    assert tender_stage["source_type"] == "uploaded_file"
    assert tender_stage["source_file_version_id"] == versions[0]["id"]

    documents = authenticated_client.get(f"/api/v1/documents?project_id={project_id}").json()
    tender_document = next(item for item in documents if item["stage"] == "tender")
    detail = authenticated_client.get(f"/api/v1/documents/{tender_document['id']}").json()
    version_id = detail["versions"][0]["id"]
    version = authenticated_client.get(f"/api/v1/documents/versions/{version_id}").json()
    templates = authenticated_client.get("/api/v1/templates?stage=tender").json()
    selected_template = next(item for item in templates if item["id"] == version["provenance"]["template_id"])
    assert "设备采购" in selected_template["name"]
    assert selected_template["source_kind"] == "platform_reference_template"
    generated_text = "".join(
        str(block["content"].get("text", ""))
        for section in version["sections"]
        for block in section["blocks"]
        if isinstance(block["content"], dict)
    )
    assert "北京地铁智能制度审查项目" in generated_text
    assert "【待确认】" in generated_text
    assert "建设投资330.00万元" not in generated_text

    validation = authenticated_client.post(f"/api/v1/documents/versions/{version_id}/validate")
    assert validation.status_code == 200, validation.text
    assert validation.json()["issue_counts"]["P0"] > 0
    finalize = authenticated_client.post(
        f"/api/v1/documents/versions/{version_id}/finalize",
        json={
            "revision": authenticated_client.get(f"/api/v1/documents/versions/{version_id}").json()[
                "revision"
            ],
            "declaration": "关键字段尚未确认时申请定稿",
        },
    )
    assert finalize.status_code == 422
    assert finalize.json()["error"]["code"] == "finalization_blocked"


def test_demo_template_has_docx_source_sections_variables_and_profile(
    authenticated_client: TestClient,
) -> None:
    templates = authenticated_client.get("/api/v1/templates?stage=tender&generation_only=true").json()
    template = next(item for item in templates if item["source_kind"] == "platform_reference_template")
    detail = authenticated_client.get(f"/api/v1/templates/{template['id']}")
    assert detail.status_code == 200
    version = detail.json()["versions"][0]
    assert version["sha256"] and version["storage_key"]
    preflight = authenticated_client.get(f"/api/v1/templates/{template['id']}/versions/1/preflight")
    assert preflight.json()["valid"] is True
    sections = authenticated_client.get(f"/api/v1/templates/{template['id']}/versions/1/sections")
    variables = authenticated_client.get(f"/api/v1/templates/{template['id']}/versions/1/variables")
    section_keys = {item["key"] for item in sections.json()}
    assert len(sections.json()) >= 5
    assert {"announcement", "instructions", "evaluation", "contract_terms", "requirements"}.issubset(
        section_keys
    )
    assert {item["variable_key"] for item in variables.json()} >= {
        "procurement_budget",
        "maximum_price",
    }
    profiles = authenticated_client.get("/api/v1/format-profiles")
    assert profiles.status_code == 200
    assert profiles.json()[0]["strict_compliance"] is False


def test_builtin_template_categories_and_generation_boundary(
    authenticated_client: TestClient,
) -> None:
    templates = authenticated_client.get("/api/v1/templates?current_only=true")
    assert templates.status_code == 200
    items = templates.json()
    assert {item["source_kind"] for item in items} >= {
        "national_official_text",
        "adapted_from_official_outline",
        "platform_reference_template",
    }
    official = next(item for item in items if item["source_kind"] == "national_official_text")
    assert official["generation_enabled"] is False
    assert official["issuing_authority"]
    assert official["source_url"].startswith("https://")
    adapted_feasibility = next(
        item for item in items if item["name"] == "政府投资项目可研报告适配模板（2023年大纲）"
    )
    adapted_sections = authenticated_client.get(
        f"/api/v1/templates/{adapted_feasibility['id']}/versions/1/sections"
    )
    assert adapted_sections.status_code == 200
    adapted_payload = adapted_sections.json()
    assert len(adapted_payload) == 38
    assert any(item["level"] == 2 for item in adapted_payload)
    assert any(item["parent_id"] for item in adapted_payload)
    assert adapted_payload[0]["level"] == 1
    assert adapted_payload[0]["parent_id"] is None
    assert adapted_payload[1]["level"] == 2
    assert adapted_payload[1]["parent_id"] == adapted_payload[0]["id"]
    generation_templates = authenticated_client.get(
        "/api/v1/templates?current_only=true&generation_only=true"
    ).json()
    assert all(item["generation_enabled"] for item in generation_templates)
    assert all(item["source_kind"] != "national_official_text" for item in generation_templates)

    project_id = _create_project(authenticated_client)
    official_feasibility = next(
        item
        for item in items
        if item["source_kind"] == "national_official_text" and item["stage"] == "feasibility"
    )
    blocked = authenticated_client.post(
        f"/api/v1/generation-jobs?project_id={project_id}&stage=feasibility",
        json={
            "template_id": official_feasibility["id"],
            "template_version": official_feasibility["current_version"],
            "idempotency_key": "official-text-must-not-generate",
        },
    )
    assert blocked.status_code == 422
    assert blocked.json()["error"]["code"] == "template_not_applicable"

    missing_source = authenticated_client.post(
        "/api/v1/templates",
        json={
            "name": "待补来源的正式模板",
            "stage": "contract",
            "source_kind": "other_official_template",
            "format_profile": {},
        },
    )
    assert missing_source.status_code == 422
    assert missing_source.json()["error"]["code"] == "official_source_metadata_missing"


def test_repeated_seed_keeps_hash_bound_to_stored_template(monkeypatch: MonkeyPatch) -> None:
    with SessionLocal() as db:
        version = db.scalar(select(TemplateVersion).where(TemplateVersion.storage_key.is_not(None)))
        assert version is not None
        assert version.storage_key is not None
        storage_key = version.storage_key
        stored_sha256 = sha256_bytes(get_storage().get(storage_key))

    monkeypatch.setattr(seed_module, "build_demo_template", lambda _stage, _name: b"new bytes")
    seed_module.seed()

    with SessionLocal() as db:
        version = db.scalar(select(TemplateVersion).where(TemplateVersion.storage_key == storage_key))
        assert version is not None
        assert version.sha256 == stored_sha256
        assert version.sha256 == sha256_bytes(get_storage().get(storage_key))


def test_comments_comparison_events_and_finalized_upstream_source(
    authenticated_client: TestClient,
) -> None:
    project_id = _create_project(authenticated_client)
    _confirm_required_fields(authenticated_client, project_id)
    document_id, version_id = _generate(authenticated_client, project_id)
    detail = authenticated_client.get(f"/api/v1/documents/versions/{version_id}").json()
    block_id = detail["sections"][0]["blocks"][0]["id"]
    comment = authenticated_client.post(
        f"/api/v1/documents/versions/{version_id}/comments",
        json={"content_block_id": block_id, "body": "请复核本节字段来源。"},
    )
    assert comment.status_code == 201, comment.text
    resolved = authenticated_client.post(
        f"/api/v1/documents/comments/{comment.json()['id']}/resolve",
        json={"revision": comment.json()["revision"]},
    )
    assert resolved.json()["status"] == "resolved"

    revision = authenticated_client.post(f"/api/v1/documents/versions/{version_id}/revisions")
    assert revision.status_code == 201, revision.text
    comparison = authenticated_client.post(
        "/api/v1/comparisons",
        json={"left_version_id": version_id, "right_version_id": revision.json()["id"]},
    )
    assert comparison.status_code == 201, comparison.text
    compared = authenticated_client.get(f"/api/v1/comparisons/{comparison.json()['id']}")
    assert compared.json()["summary"]["unchanged"] > 0

    _review_all_blocks(authenticated_client, version_id)
    validation = authenticated_client.post(f"/api/v1/documents/versions/{version_id}/validate").json()
    assert validation["status"] == "passed", [item["rule_key"] for item in validation["issues"]]
    current = authenticated_client.get(f"/api/v1/documents/{document_id}").json()
    source_version = next(item for item in current["versions"] if item["id"] == version_id)
    finalized = authenticated_client.post(
        f"/api/v1/documents/versions/{version_id}/finalize",
        json={"revision": source_version["revision"], "declaration": "字段和内容已复核"},
    )
    assert finalized.status_code == 200, finalized.text
    sources = authenticated_client.get(f"/api/v1/projects/{project_id}/stages/feasibility/sources")
    assert any(item["id"] == version_id for item in sources.json())


def test_contract_payment_check_rejects_inconsistent_totals(
    authenticated_client: TestClient,
) -> None:
    response = authenticated_client.post(
        "/api/v1/contract-payment-check",
        json={
            "final_contract_amount": 1000,
            "items": [
                {"label": "首款", "ratio": 40, "amount": 400, "trigger": "合同生效"},
                {"label": "尾款", "ratio": 50, "amount": 500, "trigger": "验收完成"},
            ],
        },
    )
    assert response.status_code == 200
    assert response.json()["valid"] is False
    assert len(response.json()["errors"]) == 2


def test_project_object_permission_denies_non_member(
    authenticated_client: TestClient, client: TestClient
) -> None:
    project_id = _create_project(authenticated_client)
    created_user = authenticated_client.post(
        "/api/v1/users",
        json={
            "email": "viewer@example.com",
            "display_name": "只读用户",
            "password": "ViewerPass123!",
            "role_keys": ["viewer"],
        },
    )
    assert created_user.status_code == 201, created_user.text
    login = client.post(
        "/api/v1/auth/login",
        json={"email": "viewer@example.com", "password": "ViewerPass123!"},
    )
    assert login.status_code == 200
    client.headers["X-CSRF-Token"] = login.json()["csrf_token"]
    denied = client.get(f"/api/v1/projects/{project_id}")
    assert denied.status_code == 403
    assert denied.json()["error"]["code"] == "project_forbidden"
