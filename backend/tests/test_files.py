from __future__ import annotations

import io

from docx import Document
from fastapi.testclient import TestClient


def _docx_bytes() -> bytes:
    stream = io.BytesIO()
    document = Document()
    document.add_heading("项目概况", level=1)
    document.add_paragraph("项目名称：测试项目")
    document.save(stream)
    return stream.getvalue()


def _project(client: TestClient) -> str:
    return client.post("/api/v1/projects", json={"code": "FILE-001", "name": "文件测试项目"}).json()["id"]


def test_upload_validates_signature_and_parses_docx(authenticated_client: TestClient) -> None:
    project_id = _project(authenticated_client)
    response = authenticated_client.post(
        f"/api/v1/files?project_id={project_id}&stage=requirement",
        files={
            "upload": (
                "source.docx",
                _docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 201, response.text
    jobs = authenticated_client.get(f"/api/v1/files/{response.json()['id']}/parse-jobs")
    assert jobs.status_code == 200
    assert jobs.json()[0]["status"] == "succeeded"


def test_upload_rejects_extension_signature_mismatch(authenticated_client: TestClient) -> None:
    project_id = _project(authenticated_client)
    response = authenticated_client.post(
        f"/api/v1/files?project_id={project_id}&stage=requirement",
        files={"upload": ("fake.pdf", b"not a pdf", "application/pdf")},
    )
    assert response.status_code == 415
    assert response.json()["error"]["code"] == "file_signature_mismatch"
