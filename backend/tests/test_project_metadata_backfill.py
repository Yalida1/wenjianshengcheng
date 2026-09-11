from __future__ import annotations

from types import SimpleNamespace

from fastapi.testclient import TestClient

from backend.app.db import SessionLocal
from backend.app.models import Project
from backend.app.services.field_extraction import (
    backfill_project_metadata_from_blocks,
    extract_project_code_from_blocks,
    extract_project_description_from_blocks,
    is_temporary_project_code,
)


def _block(sequence: int, text: str) -> SimpleNamespace:
    return SimpleNamespace(sequence=sequence, text=text)


def test_is_temporary_project_code() -> None:
    assert is_temporary_project_code("TMP-1710000000000")
    assert is_temporary_project_code("TMP-1710000000000-2")
    assert not is_temporary_project_code("XM-2026-0048")
    assert not is_temporary_project_code(None)


def test_extract_project_code_and_description_from_blocks() -> None:
    blocks = [
        _block(1, "项目名称：区域绿色数据中心节能改造"),
        _block(2, "项目编号：XM-2026-0048"),
        _block(3, "项目概况：对机房暖通与供配电系统实施节能改造，提升整体能效。"),
    ]
    assert extract_project_code_from_blocks(blocks) == "XM-2026-0048"
    assert "机房暖通" in (extract_project_description_from_blocks(blocks) or "")


def test_backfill_updates_temporary_code_and_empty_description(
    authenticated_client: TestClient,
) -> None:
    created = authenticated_client.post(
        "/api/v1/projects",
        json={"name": "待补齐项目"},
    )
    assert created.status_code == 201, created.text
    project_id = created.json()["id"]
    assert created.json()["code"].startswith("TMP-")

    with SessionLocal() as db:
        project = db.get(Project, project_id)
        assert project is not None
        result = backfill_project_metadata_from_blocks(
            db,
            project_id=project.id,
            blocks=[
                _block(1, "项目编号：BTA-UI-1788969351118"),
                _block(2, "项目说明：内置模板专项验收固定数据 UI 主链路"),
            ],
            actor_id=project.updated_by or project.created_by,
        )
        db.commit()
        db.refresh(project)
        assert result == {"updated": True, "code": True, "description": True}
        assert project.code == "BTA-UI-1788969351118"
        assert project.description == "内置模板专项验收固定数据 UI 主链路"

    refreshed = authenticated_client.get(f"/api/v1/projects/{project_id}")
    assert refreshed.status_code == 200
    assert refreshed.json()["code"] == "BTA-UI-1788969351118"
    assert refreshed.json()["description"] == "内置模板专项验收固定数据 UI 主链路"


def test_backfill_does_not_overwrite_formal_code(authenticated_client: TestClient) -> None:
    created = authenticated_client.post(
        "/api/v1/projects",
        json={"code": "FORMAL-001", "name": "已有正式编号", "description": "已有说明"},
    )
    assert created.status_code == 201, created.text
    project_id = created.json()["id"]

    with SessionLocal() as db:
        project = db.get(Project, project_id)
        assert project is not None
        revision = project.revision
        result = backfill_project_metadata_from_blocks(
            db,
            project_id=project.id,
            blocks=[
                _block(1, "项目编号：SHOULD-NOT-APPLY"),
                _block(2, "项目说明：也不应覆盖"),
            ],
            actor_id=project.updated_by or project.created_by,
        )
        db.commit()
        db.refresh(project)
        assert result == {"updated": False, "code": False, "description": False}
        assert project.code == "FORMAL-001"
        assert project.description == "已有说明"
        assert project.revision == revision


def test_parse_source_material_backfills_project_metadata(
    authenticated_client: TestClient,
) -> None:
    import io

    from docx import Document

    created = authenticated_client.post(
        "/api/v1/projects",
        json={"name": "解析回填项目"},
    )
    assert created.status_code == 201, created.text
    project = created.json()
    assert project["code"].startswith("TMP-")
    assert project["description"] is None

    stream = io.BytesIO()
    document = Document()
    document.add_paragraph("项目编号：KY-2026-7788")
    document.add_paragraph("项目概况：区域绿色数据中心节能改造可研编制依据材料。")
    document.add_paragraph("项目名称：解析回填项目")
    document.save(stream)

    uploaded = authenticated_client.post(
        f"/api/v1/files?project_id={project['id']}&stage=feasibility",
        files={
            "upload": (
                "feasibility.docx",
                stream.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert uploaded.status_code == 201, uploaded.text
    assert uploaded.json()["status"] in {"parsed", "needs_ocr"}

    refreshed = authenticated_client.get(f"/api/v1/projects/{project['id']}")
    assert refreshed.status_code == 200
    body = refreshed.json()
    assert body["code"] == "KY-2026-7788"
    assert "节能改造" in (body["description"] or "")
