from __future__ import annotations

import io
import uuid
import zipfile

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import select

from backend.app import api as api_module
from backend.app.db import SessionLocal
from backend.app.models import TemplateSection
from backend.app.tasks import export_document_task

REQUIREMENT_FIELDS = {
    "project_name": ("项目名称", "宁夏政务服务平台建设项目", None),
    "project_owner": ("项目单位", "宁夏示例单位", None),
    "construction_scope": ("建设范围", "建设统一申报和事项管理能力", None),
    "project_period": ("项目总建设周期", 18, "月"),
}


def _create_project(client: TestClient) -> str:
    response = client.post(
        "/api/v1/projects",
        json={"code": f"FLOW-{uuid.uuid4().hex[:8]}", "name": "宁夏政务服务平台建设项目"},
    )
    assert response.status_code == 201, response.text
    return response.json()["id"]


def _confirm_required_fields(client: TestClient, project_id: str) -> None:
    for key, (label, value, unit) in REQUIREMENT_FIELDS.items():
        created = client.post(
            f"/api/v1/field-values?project_id={project_id}&stage=requirement",
            json={
                "field_key": key,
                "field_label": label,
                "value": value,
                "normalized_value": value,
                "unit": unit,
                "criticality": "P0",
                "status": "missing",
                "source_type": "user_input",
            },
        )
        assert created.status_code == 201, created.text
        field = created.json()
        confirmed = client.post(
            f"/api/v1/field-values/{field['id']}/confirm",
            json={"revision": field["revision"], "evidence_acknowledged": True},
        )
        assert confirmed.status_code == 200, confirmed.text


def _generate(client: TestClient, project_id: str) -> tuple[str, str]:
    templates = client.get("/api/v1/templates?stage=requirement&generation_only=true")
    assert templates.status_code == 200
    template = templates.json()[0]
    generated = client.post(
        f"/api/v1/generation-jobs?project_id={project_id}&stage=requirement",
        json={
            "template_id": template["id"],
            "template_version": template["current_version"],
            "idempotency_key": f"generation-{uuid.uuid4()}",
        },
    )
    assert generated.status_code == 202, generated.text
    assert generated.json()["status"] == "succeeded"
    documents = client.get(f"/api/v1/documents?project_id={project_id}").json()
    document_id = documents[0]["id"]
    detail = client.get(f"/api/v1/documents/{document_id}").json()
    return document_id, detail["versions"][0]["id"]


def _review_all_blocks(client: TestClient, version_id: str) -> None:
    detail = client.get(f"/api/v1/documents/versions/{version_id}").json()
    for section in detail["sections"]:
        for block in section["blocks"]:
            response = client.patch(
                f"/api/v1/documents/blocks/{block['id']}",
                json={"content": block["content"], "reviewed": True, "revision": block["revision"]},
            )
            assert response.status_code == 200, response.text


def test_export_rejects_a_document_revision_changed_after_queueing(
    authenticated_client: TestClient, monkeypatch
) -> None:
    project_id = _create_project(authenticated_client)
    _document_id, version_id = _generate(authenticated_client, project_id)
    monkeypatch.setattr(api_module.export_document_task, "delay", lambda _job_id: None)

    queued = authenticated_client.post(
        f"/api/v1/exports?version_id={version_id}",
        json={
            "output_format": "docx",
            "idempotency_key": f"stale-export-{uuid.uuid4()}",
        },
    )
    assert queued.status_code == 202, queued.text
    job = queued.json()
    assert job["version_revision"] == 1

    version = authenticated_client.get(f"/api/v1/documents/versions/{version_id}").json()
    block = next(
        block
        for section in version["sections"]
        for block in section["blocks"]
        if block["source_kind"] != "fixed_template"
    )
    changed = authenticated_client.patch(
        f"/api/v1/documents/blocks/{block['id']}",
        json={
            "content": block["content"],
            "reviewed": True,
            "revision": block["revision"],
        },
    )
    assert changed.status_code == 200, changed.text

    assert export_document_task.run(job["id"]) == "stale"
    export_status = authenticated_client.get(f"/api/v1/exports/{job['id']}")
    assert export_status.status_code == 200
    assert export_status.json()["status"] == "failed"
    assert "已更新" in export_status.json()["error"]


def test_p0_gate_blocks_missing_fields(authenticated_client: TestClient) -> None:
    project_id = _create_project(authenticated_client)
    _document_id, version_id = _generate(authenticated_client, project_id)
    validation = authenticated_client.post(f"/api/v1/documents/versions/{version_id}/validate")
    assert validation.status_code == 200
    assert validation.json()["issue_counts"]["P0"] >= len(REQUIREMENT_FIELDS)
    finalize = authenticated_client.post(
        f"/api/v1/documents/versions/{version_id}/finalize",
        json={
            "revision": authenticated_client.get(f"/api/v1/documents/versions/{version_id}").json()[
                "revision"
            ],
            "declaration": "已完成审核并申请定稿",
        },
    )
    assert finalize.status_code == 422
    assert finalize.json()["error"]["code"] == "finalization_blocked"


def test_full_flow_finalizes_and_exports_real_files(authenticated_client: TestClient) -> None:
    project_id = _create_project(authenticated_client)
    _confirm_required_fields(authenticated_client, project_id)
    document_id, version_id = _generate(authenticated_client, project_id)
    _review_all_blocks(authenticated_client, version_id)
    validation = authenticated_client.post(f"/api/v1/documents/versions/{version_id}/validate")
    assert validation.status_code == 200
    assert validation.json()["status"] == "passed", [
        (item["rule_key"], item["location"]) for item in validation.json()["issues"]
    ]
    finalized = authenticated_client.post(
        f"/api/v1/documents/versions/{version_id}/finalize",
        json={
            "revision": authenticated_client.get(f"/api/v1/documents/versions/{version_id}").json()[
                "revision"
            ],
            "declaration": "关键字段和生成内容已完成人工复核",
        },
    )
    assert finalized.status_code == 200, finalized.text

    for output_format in ("docx", "pdf", "xlsx"):
        response = authenticated_client.post(
            f"/api/v1/exports?version_id={version_id}",
            json={
                "output_format": output_format,
                "idempotency_key": f"export-{output_format}-{uuid.uuid4()}",
            },
        )
        assert response.status_code == 202, response.text
        job = response.json()
        assert job["status"] == "succeeded", job
        download = authenticated_client.get(f"/api/v1/exports/{job['id']}/download")
        assert download.status_code == 200, download.text
        content = download.content
        if output_format == "docx":
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                assert "word/document.xml" in archive.namelist()
        elif output_format == "pdf":
            assert len(PdfReader(io.BytesIO(content)).pages) >= 2
        else:
            workbook = load_workbook(io.BytesIO(content), read_only=True)
            assert workbook.sheetnames == ["字段来源"]

    detail = authenticated_client.get(f"/api/v1/documents/{document_id}").json()
    assert detail["versions"][0]["immutable"] is True


def test_outline_selected_generation_accumulates_sections_and_supports_ai_optimization(
    authenticated_client: TestClient,
) -> None:
    project_id = _create_project(authenticated_client)
    templates = authenticated_client.get("/api/v1/templates?stage=requirement&generation_only=true").json()
    template = templates[0]
    outline = authenticated_client.get(
        f"/api/v1/templates/{template['id']}/versions/{template['current_version']}/sections"
    )
    assert outline.status_code == 200, outline.text
    section_keys = [item["key"] for item in outline.json()]
    assert len(section_keys) >= 2

    first_job = authenticated_client.post(
        f"/api/v1/generation-jobs?project_id={project_id}&stage=requirement",
        json={
            "template_id": template["id"],
            "template_version": template["current_version"],
            "idempotency_key": f"selected-generation-{uuid.uuid4()}",
            "selected_section_keys": [section_keys[0]],
        },
    )
    assert first_job.status_code == 202, first_job.text
    assert first_job.json()["status"] == "succeeded"
    steps = authenticated_client.get(f"/api/v1/generation-jobs/{first_job.json()['id']}/steps").json()
    assert next(item for item in steps if item["step_key"] == section_keys[0])["status"] == "succeeded"
    assert next(item for item in steps if item["step_key"] == section_keys[1])["status"] == "skipped"

    document_id = authenticated_client.get(f"/api/v1/documents?project_id={project_id}").json()[0]["id"]
    document = authenticated_client.get(f"/api/v1/documents/{document_id}").json()
    first_version = authenticated_client.get(
        f"/api/v1/documents/versions/{document['versions'][0]['id']}"
    ).json()
    first_by_key = {section["key"]: section for section in first_version["sections"]}
    assert set(first_by_key) == set(section_keys)
    assert first_by_key[section_keys[0]]["blocks"]
    assert first_by_key[section_keys[1]]["blocks"] == []
    retained_text = first_by_key[section_keys[0]]["blocks"][0]["content"]["text"]

    second_job = authenticated_client.post(
        f"/api/v1/generation-jobs?project_id={project_id}&stage=requirement",
        json={
            "template_id": template["id"],
            "template_version": template["current_version"],
            "idempotency_key": f"selected-generation-{uuid.uuid4()}",
            "selected_section_keys": [section_keys[1]],
        },
    )
    assert second_job.status_code == 202, second_job.text
    document = authenticated_client.get(f"/api/v1/documents/{document_id}").json()
    second_version = authenticated_client.get(
        f"/api/v1/documents/versions/{document['versions'][0]['id']}"
    ).json()
    second_by_key = {section["key"]: section for section in second_version["sections"]}
    assert second_by_key[section_keys[0]]["blocks"][0]["content"]["text"] == retained_text
    generated_block = second_by_key[section_keys[1]]["blocks"][0]

    selected_text = generated_block["content"]["text"][:30]
    optimized = authenticated_client.post(
        f"/api/v1/documents/blocks/{generated_block['id']}/ai-optimize",
        json={"selected_text": selected_text, "prompt": "语言更正式、简洁", "action": "polish"},
    )
    assert optimized.status_code == 200, optimized.text
    assert optimized.json()["optimized_prompt"]
    assert optimized.json()["suggestion"]
    assert optimized.json()["provider"] == "demo"


def test_leaf_generation_and_explicit_selection_preserve_unselected_body(
    authenticated_client: TestClient,
) -> None:
    client = authenticated_client
    project_id = _create_project(client)
    template = client.get("/api/v1/templates?stage=requirement&generation_only=true").json()[0]
    outline_url = f"/api/v1/templates/{template['id']}/versions/{template['current_version']}/sections"
    outline = client.get(outline_url).json()
    root, first_child, second_child = outline[:3]
    with SessionLocal() as db:
        for child in (first_child, second_child):
            section = db.scalar(select(TemplateSection).where(TemplateSection.id == child["id"]))
            assert section is not None
            section.parent_id = root["id"]
            section.level = 2
        db.commit()

    def generate(keys: list[str]) -> dict:
        response = client.post(
            f"/api/v1/generation-jobs?project_id={project_id}&stage=requirement",
            json={
                "template_id": template["id"],
                "template_version": template["current_version"],
                "idempotency_key": f"leaf-generation-{uuid.uuid4()}",
                "selected_section_keys": keys,
                "include_descendants": False,
            },
        )
        assert response.status_code == 202, response.text
        assert response.json()["status"] == "succeeded"
        steps = client.get(f"/api/v1/generation-jobs/{response.json()['id']}/steps").json()
        assert {step["step_key"] for step in steps if step["status"] == "succeeded"} == set(keys)
        document_id = client.get(f"/api/v1/documents?project_id={project_id}").json()[0]["id"]
        document = client.get(f"/api/v1/documents/{document_id}").json()
        return client.get(f"/api/v1/documents/versions/{document['versions'][0]['id']}").json()

    first = generate([first_child["key"]])
    sections = {section["key"]: section for section in first["sections"]}
    assert sections[root["key"]]["blocks"] == []
    assert sections[second_child["key"]]["blocks"] == []
    assert sections[first_child["key"]]["blocks"]
    assert sections[first_child["key"]]["parent_id"] == sections[root["key"]]["id"]
    retained_body = [block["content"] for block in sections[first_child["key"]]["blocks"]]

    second = generate([root["key"], second_child["key"]])
    sections = {section["key"]: section for section in second["sections"]}
    assert sections[root["key"]]["blocks"]
    assert sections[second_child["key"]]["blocks"]
    assert [block["content"] for block in sections[first_child["key"]]["blocks"]] == retained_body
    assert second["provenance"]["selected_section_keys"] == [root["key"], second_child["key"]]


def test_forbidden_cross_stage_mapping_is_rejected(authenticated_client: TestClient) -> None:
    project_id = _create_project(authenticated_client)
    response = authenticated_client.post(
        f"/api/v1/field-values?project_id={project_id}&stage=tender",
        json={
            "field_key": "procurement_budget",
            "field_label": "招标预算",
            "value": 1000000,
            "criticality": "P0",
            "status": "extracted",
            "source_type": "extracted",
            "evidence": {
                "excerpt": "可研总投资 100 万元",
                "source_stage": "feasibility",
                "source_field_key": "total_investment",
            },
        },
    )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "forbidden_field_mapping"
