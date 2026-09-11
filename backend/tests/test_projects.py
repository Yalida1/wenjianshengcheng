from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient


def _docx_bytes() -> bytes:
    stream = io.BytesIO()
    document = Document()
    document.add_heading("项目概况", level=1)
    document.add_paragraph("项目名称：待删除项目")
    document.save(stream)
    return stream.getvalue()


def test_project_creation_creates_four_stages(authenticated_client: TestClient) -> None:
    created = authenticated_client.post(
        "/api/v1/projects",
        json={"code": "NX-001", "name": "宁夏政务服务平台建设项目"},
    )
    assert created.status_code == 201, created.text
    project = created.json()
    stages = authenticated_client.get(f"/api/v1/projects/{project['id']}/stages")
    assert stages.status_code == 200
    assert {item["stage"] for item in stages.json()} == {
        "requirement",
        "feasibility",
        "tender",
        "contract",
    }


def test_project_creation_without_code_assigns_temporary_code(
    authenticated_client: TestClient,
) -> None:
    created = authenticated_client.post(
        "/api/v1/projects",
        json={"name": "仅名称新建项目", "project_type": "enterprise_investment"},
    )
    assert created.status_code == 201, created.text
    project = created.json()
    assert project["name"] == "仅名称新建项目"
    assert project["description"] is None
    assert project["code"].startswith("TMP-")
    assert project["project_type"] == "enterprise_investment"


def test_project_patch_can_update_code(authenticated_client: TestClient) -> None:
    project = authenticated_client.post(
        "/api/v1/projects",
        json={"name": "待回填编号项目"},
    ).json()
    updated = authenticated_client.patch(
        f"/api/v1/projects/{project['id']}",
        json={"code": "XM-2026-0099", "revision": project["revision"]},
    )
    assert updated.status_code == 200, updated.text
    assert updated.json()["code"] == "XM-2026-0099"
    assert updated.json()["revision"] == project["revision"] + 1


def test_revision_conflict_returns_409(authenticated_client: TestClient) -> None:
    project = authenticated_client.post(
        "/api/v1/projects",
        json={"code": "NX-002", "name": "并发测试项目"},
    ).json()
    updated = authenticated_client.patch(
        f"/api/v1/projects/{project['id']}",
        json={"name": "第一次修改", "revision": project["revision"]},
    )
    assert updated.status_code == 200
    conflict = authenticated_client.patch(
        f"/api/v1/projects/{project['id']}",
        json={"name": "覆盖修改", "revision": project["revision"]},
    )
    assert conflict.status_code == 409
    assert conflict.json()["error"]["code"] == "revision_conflict"


def test_admin_can_delete_complete_project_graph(authenticated_client: TestClient) -> None:
    project = authenticated_client.post(
        "/api/v1/projects",
        json={"code": "DEL-001", "name": "待删除项目"},
    ).json()
    uploaded = authenticated_client.post(
        f"/api/v1/files?project_id={project['id']}&stage=requirement",
        files={
            "upload": (
                "source.docx",
                _docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert uploaded.status_code == 201, uploaded.text
    storage_dir = (
        Path("data/test-objects").resolve()
        / project["organization_id"]
        / uploaded.json()["id"]
    )
    assert storage_dir.is_dir()

    deleted = authenticated_client.request(
        "DELETE",
        f"/api/v1/projects/{project['id']}",
        json={"revision": project["revision"], "confirmation_code": project["code"]},
    )
    assert deleted.status_code == 200, deleted.text
    assert deleted.json() == {
        "project_id": project["id"],
        "deleted": True,
        "storage_objects_deleted": 1,
        "storage_cleanup_failed": 0,
    }
    assert authenticated_client.get(f"/api/v1/projects/{project['id']}").status_code == 404
    assert not storage_dir.exists()

    logs = authenticated_client.get("/api/v1/audit-logs?page=1&page_size=50")
    assert logs.status_code == 200
    assert any(
        item["action"] == "project.delete" and item["object_id"] == project["id"]
        for item in logs.json()["items"]
    )


def test_project_delete_checks_revision_and_confirmation_code(
    authenticated_client: TestClient,
) -> None:
    project = authenticated_client.post(
        "/api/v1/projects",
        json={"code": "DEL-002", "name": "删除校验项目"},
    ).json()
    revision_conflict = authenticated_client.request(
        "DELETE",
        f"/api/v1/projects/{project['id']}",
        json={"revision": project["revision"] + 1, "confirmation_code": project["code"]},
    )
    assert revision_conflict.status_code == 409
    assert revision_conflict.json()["error"]["code"] == "revision_conflict"

    wrong_code = authenticated_client.request(
        "DELETE",
        f"/api/v1/projects/{project['id']}",
        json={"revision": project["revision"], "confirmation_code": "WRONG-CODE"},
    )
    assert wrong_code.status_code == 422
    assert wrong_code.json()["error"]["code"] == "project_confirmation_mismatch"
    assert authenticated_client.get(f"/api/v1/projects/{project['id']}").status_code == 200


def test_project_delete_requires_system_admin(authenticated_client: TestClient) -> None:
    project = authenticated_client.post(
        "/api/v1/projects",
        json={"code": "DEL-003", "name": "权限校验项目"},
    ).json()
    created_user = authenticated_client.post(
        "/api/v1/users",
        json={
            "email": "editor@example.com",
            "display_name": "项目编制人员",
            "password": "editor-pass-123",
            "role_keys": ["project_editor"],
        },
    )
    assert created_user.status_code == 201, created_user.text
    login = authenticated_client.post(
        "/api/v1/auth/login",
        json={"email": "editor@example.com", "password": "editor-pass-123"},
    )
    assert login.status_code == 200, login.text
    authenticated_client.headers["X-CSRF-Token"] = login.json()["csrf_token"]

    forbidden = authenticated_client.request(
        "DELETE",
        f"/api/v1/projects/{project['id']}",
        json={"revision": project["revision"], "confirmation_code": project["code"]},
    )
    assert forbidden.status_code == 403
    assert forbidden.json()["error"]["code"] == "forbidden"
