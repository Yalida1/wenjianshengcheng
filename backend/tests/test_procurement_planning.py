from __future__ import annotations

import io
import shutil
import uuid
import zipfile
from datetime import timedelta
from typing import Any, cast

import pytest
from docx import Document as WordDocument
from fastapi.testclient import TestClient
from pypdf import PdfReader
from sqlalchemy import select

from backend.app import api as api_module
from backend.app import tasks as tasks_module
from backend.app.db import SessionLocal
from backend.app.main import app
from backend.app.models import (
    Document,
    DocumentContentBlock,
    DocumentSection,
    DocumentVersion,
    ParsedDocument,
    ProcurementAnalysisCheckpoint,
    Project,
    ProjectStage,
    utc_now,
)
from backend.app.services import procurement_planning as procurement_planning_service


def _docx(lines: list[str]) -> bytes:
    document = WordDocument()
    document.add_heading("可行性研究报告", level=1)
    for line in lines:
        document.add_paragraph(line)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _create_project(client: TestClient) -> str:
    response = client.post(
        "/api/v1/projects",
        json={"code": f"PROC-{uuid.uuid4().hex[:8]}", "name": "采购方案测试项目"},
    )
    assert response.status_code == 201, response.text
    return cast(str, response.json()["id"])


def _upload_and_bind(client: TestClient, project_id: str, lines: list[str]) -> tuple[str, str]:
    uploaded = client.post(
        f"/api/v1/files?project_id={project_id}&stage=tender",
        files={
            "upload": (
                "feasibility.docx",
                _docx(lines),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert uploaded.status_code == 201, uploaded.text
    file_id = uploaded.json()["id"]
    versions = client.get(f"/api/v1/files/{file_id}/versions")
    assert versions.status_code == 200, versions.text
    version_id = versions.json()[0]["id"]
    stages = client.get(f"/api/v1/projects/{project_id}/stages")
    tender_stage = next(item for item in stages.json() if item["stage"] == "tender")
    bound = client.put(
        f"/api/v1/projects/{project_id}/stages/tender/source",
        json={
            "source_type": "uploaded_file",
            "source_file_version_id": version_id,
            "revision": tender_stage["revision"],
        },
    )
    assert bound.status_code == 200, bound.text
    return file_id, version_id


def _analyze(client: TestClient, project_id: str) -> dict[str, Any]:
    response = client.post(
        f"/api/v1/projects/{project_id}/procurement-analyses",
        json={"include_alternative": True},
    )
    assert response.status_code == 202, response.text
    assert response.json()["status"] == "succeeded_demo"
    plans = client.get(f"/api/v1/projects/{project_id}/procurement-plans")
    assert plans.status_code == 200, plans.text
    return cast(dict[str, Any], plans.json()[0])


def _plan(client: TestClient, lines: list[str]) -> tuple[str, dict[str, Any]]:
    project_id = _create_project(client)
    _upload_and_bind(client, project_id, lines)
    return project_id, _analyze(client, project_id)


def _structure_payload(plan: dict[str, Any], *, groups: list[dict[str, Any]] | None = None) -> dict[str, Any]:
    packages = [
        {
            "id": item["id"],
            "code": item["code"],
            "name": item["name"],
            "procurement_category": item["procurement_category"],
            "business_subcategory": item["business_subcategory"],
            "procurement_method": item["procurement_method"],
            "scope": item["scope"],
            "exclusions": item["exclusions"],
            "deliverables": item["deliverables"],
            "implementation_period": item["implementation_period"],
            "estimated_amount": item["estimated_amount"],
            "confirmed_budget": item["confirmed_budget"],
            "maximum_price": item["maximum_price"],
            "currency": item["currency"],
            "original_unit": item["original_unit"],
            "tax_included": item["tax_included"],
            "budget_period": item["budget_period"],
            "budget_status": item["budget_status"],
            "budget_basis": item["budget_basis"],
            "evidence_status": item["evidence_status"],
            "content_item_ids": item["content_item_ids"],
        }
        for item in plan["packages"]
    ]
    document_groups = groups or [
        {
            "id": item["id"],
            "code": item["code"],
            "name": item["name"],
            "procurement_category": item["procurement_category"],
            "business_subcategory": item["business_subcategory"],
            "procurement_method": item["procurement_method"],
            "organization_method": item["organization_method"],
            "scope": item["scope"],
            "exclusions": item["exclusions"],
            "deliverables": item["deliverables"],
            "implementation_period": item["implementation_period"],
            "rationale": item["rationale"],
            "template_id": item["template_id"],
            "template_version": item["template_version"],
            "template_match_basis": item["template_match_basis"],
            "status": item["status"],
            "package_ids": item["package_ids"],
        }
        for item in plan["document_groups"]
    ]
    return {
        "revision": plan["revision"],
        "name": plan["name"],
        "packages": packages,
        "document_groups": document_groups,
    }


def _confirm(client: TestClient, plan: dict[str, Any]) -> dict[str, Any]:
    response = client.post(
        f"/api/v1/procurement-plans/{plan['id']}/confirm",
        json={"revision": plan["revision"], "decision_note": "已核对采购范围和文件划分"},
    )
    assert response.status_code == 200, response.text
    return cast(dict[str, Any], response.json())


def test_live_analysis_checkpoints_are_reused_after_retry(
    authenticated_client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    project_id = _create_project(authenticated_client)
    _file_id, version_id = _upload_and_bind(
        authenticated_client,
        project_id,
        [
            "主招标文件：断点续跑测试招标文件|平台建设|整体交付",
            "采购包：P1|平台建设包|软件、硬件及集成服务",
        ],
    )
    calls: list[str] = []

    def fake_request_model(
        _provider: object,
        model_type: type[Any],
        *,
        schema_name: str,
        system: str,
        user_payload: dict[str, Any],
    ) -> Any:
        del system, user_payload
        calls.append(schema_name)
        if model_type is procurement_planning_service.ProcurementFacts:
            return procurement_planning_service.ProcurementFacts()
        return procurement_planning_service._demo_analysis(blocks, "采购方案测试项目")

    monkeypatch.setattr(
        procurement_planning_service.OpenAICompatibleProvider,
        "_request_model",
        fake_request_model,
    )
    monkeypatch.setattr(
        procurement_planning_service.OpenAICompatibleProvider,
        "__init__",
        lambda self: None,
    )
    with SessionLocal() as db:
        project = db.get(Project, project_id)
        assert project is not None
        run = procurement_planning_service.build_analysis_run(
            db,
            project=project,
            source_kind="uploaded_file",
            source_version_id=version_id,
            user_id=project.created_by or "test-user",
        )
        db.commit()
        blocks, _sha256, _coverage = procurement_planning_service._source_for_run(db, run)
        monkeypatch.setattr(
            procurement_planning_service,
            "_chunks",
            lambda source_blocks: [[source_blocks[0]], source_blocks[1:]],
        )

        first = procurement_planning_service._openai_analysis(db, run, blocks, project.name)
        second = procurement_planning_service._openai_analysis(db, run, blocks, project.name)

        assert first.model_dump(mode="json") == second.model_dump(mode="json")
        assert len(calls) == 3
        assert set(calls[:2]) == {"procurement_facts_1", "procurement_facts_2"}
        assert calls[-1] == "procurement_plan_candidates"
        checkpoints = list(
            db.scalars(
                select(ProcurementAnalysisCheckpoint).where(ProcurementAnalysisCheckpoint.run_id == run.id)
            )
        )
        assert len(checkpoints) == 3
        assert all(item.status == "succeeded" for item in checkpoints)
        db.refresh(run)
        assert run.coverage_json["completed_chunk_count"] == 2
        assert run.coverage_json["processed_block_count"] == len(blocks)


def test_failed_analysis_can_be_requeued_without_creating_a_new_run(
    authenticated_client: TestClient,
) -> None:
    project_id = _create_project(authenticated_client)
    _file_id, version_id = _upload_and_bind(
        authenticated_client,
        project_id,
        ["主招标文件：重试测试招标文件|测试范围|测试依据", "采购包：P1|测试包|测试范围"],
    )
    with SessionLocal() as db:
        project = db.get(Project, project_id)
        assert project is not None
        run = procurement_planning_service.build_analysis_run(
            db,
            project=project,
            source_kind="uploaded_file",
            source_version_id=version_id,
            user_id=project.created_by or "test-user",
        )
        run.status = "failed"
        run.task_id = "orphaned-task"
        run.error = "SoftTimeLimitExceeded()"
        db.commit()
        run_id = run.id

    retried = authenticated_client.post(f"/api/v1/procurement-analyses/{run_id}/retry")

    assert retried.status_code == 202, retried.text
    assert retried.json()["id"] == run_id
    assert retried.json()["task_id"] != "orphaned-task"
    assert retried.json()["status"] == "succeeded_demo"
    runs = authenticated_client.get(f"/api/v1/projects/{project_id}/procurement-analyses")
    assert runs.status_code == 200
    assert [item["id"] for item in runs.json()] == [run_id]


def test_watchdog_recovers_analysis_orphaned_after_worker_restart(
    authenticated_client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    project_id = _create_project(authenticated_client)
    _file_id, version_id = _upload_and_bind(
        authenticated_client,
        project_id,
        ["主招标文件：看门任务测试招标文件|测试范围|测试依据", "采购包：P1|测试包|测试范围"],
    )
    with SessionLocal() as db:
        project = db.get(Project, project_id)
        assert project is not None
        run = procurement_planning_service.build_analysis_run(
            db,
            project=project,
            source_kind="uploaded_file",
            source_version_id=version_id,
            user_id=project.created_by or "test-user",
        )
        run.status = "running"
        run.task_id = "dead-worker-task"
        run.updated_at = utc_now() - timedelta(
            seconds=tasks_module.settings.procurement_analysis_stale_after_seconds + 1
        )
        db.commit()
        run_id = run.id

    dispatched: list[tuple[list[str], str]] = []

    def fake_apply_async(*, args: list[str], task_id: str) -> None:
        dispatched.append((args, task_id))

    monkeypatch.setattr(tasks_module.analyze_procurement_task, "apply_async", fake_apply_async)

    assert tasks_module.recover_stale_procurement_analyses() == 1
    assert len(dispatched) == 1
    assert dispatched[0][0] == [run_id]
    with SessionLocal() as db:
        recovered = db.get(procurement_planning_service.ProcurementAnalysisRun, run_id)
        assert recovered is not None
        assert recovered.status == "queued"
        assert recovered.task_id == dispatched[0][1]
        assert recovered.task_id != "dead-worker-task"


def test_three_explicit_independent_tender_documents(authenticated_client: TestClient) -> None:
    _project_id, plan = _plan(
        authenticated_client,
        [
            "独立主招标文件：网络设备采购招标文件|核心交换与接入网络|网络可独立验收",
            "采购包：P1|网络设备包|交换机及网络配套",
            "独立主招标文件：应用软件建设招标文件|业务应用软件|软件成果独立验收",
            "采购包：P2|应用软件包|应用软件开发与集成",
            "独立主招标文件：第三方测评服务招标文件|第三方测评|独立性要求",
            "采购包：P3|测评服务包|等保测评与验收测评",
        ],
    )
    assert plan["recommended_document_count"] == 3
    assert plan["procurement_package_count"] == 3
    assert [item["name"] for item in plan["document_groups"]] == [
        "网络设备采购招标文件",
        "应用软件建设招标文件",
        "第三方测评服务招标文件",
    ]
    assert all(item["package_ids"] for item in plan["document_groups"])
    assert plan["analysis_coverage"]["processed_block_count"] == plan["analysis_coverage"]["block_count"]


def test_one_main_document_with_three_packages_counts_one(authenticated_client: TestClient) -> None:
    _project_id, plan = _plan(
        authenticated_client,
        [
            "主招标文件：数据中心设备采购招标文件|数据中心设备统一采购|统一接口和整体验收",
            "采购包：P1|服务器包|计算服务器",
            "采购包：P2|网络设备包|核心网络设备",
            "采购包：P3|安全设备包|安全防护设备",
        ],
    )
    assert plan["recommended_document_count"] == 1
    assert plan["procurement_package_count"] == 3
    assert len(plan["document_groups"][0]["package_ids"]) == 3


def test_integrated_software_hardware_is_not_mechanically_split(
    authenticated_client: TestClient,
) -> None:
    _project_id, plan = _plan(
        authenticated_client,
        [
            "主招标文件：一体化平台建设招标文件|软件、硬件与集成服务整体交付|统一接口、责任和整体验收",
            "采购包：P1|一体化建设包|应用软件、服务器和系统集成整体交付",
        ],
    )
    assert plan["recommended_document_count"] == 1
    assert "整体" in plan["document_groups"][0]["rationale"]


def test_mixed_multiple_documents_and_multi_package_mapping(
    authenticated_client: TestClient,
) -> None:
    _project_id, plan = _plan(
        authenticated_client,
        [
            "独立主招标文件：平台建设招标文件|平台建设|总集成责任统一",
            "采购包：P1|软件包|应用软件",
            "采购包：P2|硬件包|配套硬件",
            "独立主招标文件：监理服务招标文件|项目监理|独立监督要求",
            "采购包：P3|监理包|全过程监理",
        ],
    )
    assert plan["recommended_document_count"] == 2
    assert [len(item["package_ids"]) for item in plan["document_groups"]] == [2, 1]


def test_total_investment_is_not_copied_to_package_budget_or_ceiling(
    authenticated_client: TestClient,
) -> None:
    _project_id, plan = _plan(
        authenticated_client,
        [
            "总投资：330万元",
            "主招标文件：平台建设招标文件|平台建设|统一交付",
            "采购包：P1|平台建设包|软件与配套服务",
        ],
    )
    package = plan["packages"][0]
    assert package["estimated_amount"] is None
    assert package["confirmed_budget"] is None
    assert package["maximum_price"] is None
    assert plan["budget_reconciliation"]["items"][0]["amount"] == "3300000.0000"
    assert plan["budget_reconciliation"]["total_investment_copied_to_packages"] is False
    assert plan["recommended_document_count"] == 1


def test_existing_assets_reuse_and_future_scope_are_excluded(
    authenticated_client: TestClient,
) -> None:
    _project_id, plan = _plan(
        authenticated_client,
        [
            "复用内容：现有服务器",
            "已采购：机房网络设备",
            "远期规划：二期数据治理能力",
            "主招标文件：本期应用建设招标文件|本期应用建设|按本期范围验收",
            "采购包：P1|本期应用包|本期应用功能",
            "本期应用功能在其他章节重复描述，但不是新增采购包。",
        ],
    )
    assert plan["procurement_package_count"] == 1
    assert plan["coverage_check"]["excluded_count"] == 3
    assert plan["coverage_check"]["duplicate_content_ids"] == []


def test_conflict_and_incomplete_parse_keep_document_count_unknown(
    authenticated_client: TestClient,
) -> None:
    project_id = _create_project(authenticated_client)
    _file_id, version_id = _upload_and_bind(
        authenticated_client,
        project_id,
        [
            "来源冲突：可研正文要求统一采购，批准材料附件要求拆为两份",
            "主招标文件：平台建设招标文件|平台建设|原文候选",
            "采购包：P1|平台建设包|平台建设",
        ],
    )
    with SessionLocal() as db:
        parsed = db.query(ParsedDocument).filter(ParsedDocument.file_version_id == version_id).one()
        parsed.metadata_json = {**parsed.metadata_json, "unparsed_pages": [18], "needs_ocr": True}
        db.commit()
    plan = _analyze(authenticated_client, project_id)
    assert plan["recommended_document_count"] is None
    codes = {item["code"] for item in plan["unresolved_items"]}
    assert {"source.conflict", "source.incomplete_parse"}.issubset(codes)
    blocked = authenticated_client.post(
        f"/api/v1/procurement-plans/{plan['id']}/confirm",
        json={"revision": plan["revision"], "decision_note": "尝试确认"},
    )
    assert blocked.status_code == 422


def test_non_tender_method_is_separately_counted_and_can_result_in_zero(
    authenticated_client: TestClient,
) -> None:
    _project_id, plan = _plan(
        authenticated_client,
        [
            "主招标文件：运维服务采购文件|年度运维服务|非招标采购|competitive_negotiation",
            "采购包：P1|运维服务包|年度运维服务|competitive_negotiation",
        ],
    )
    assert plan["recommended_document_count"] == 0
    assert plan["other_procurement_document_count"] == 1
    assert plan["procurement_package_count"] == 1


def test_structure_adjustment_recalculates_count_and_persists(
    authenticated_client: TestClient,
) -> None:
    _project_id, plan = _plan(
        authenticated_client,
        [
            "主招标文件：平台建设招标文件|平台建设|统一采购候选",
            "采购包：P1|软件包|应用软件",
            "采购包：P2|硬件包|硬件设备",
        ],
    )
    base_group = plan["document_groups"][0]
    groups = []
    for index, package in enumerate(plan["packages"], 1):
        groups.append(
            {
                "code": f"DOC-{index:02d}",
                "name": f"{package['name']}招标文件",
                "procurement_category": package["procurement_category"],
                "business_subcategory": package["business_subcategory"],
                "procurement_method": "public_tender",
                "organization_method": None,
                "scope": package["scope"],
                "exclusions": package["exclusions"],
                "deliverables": package["deliverables"],
                "implementation_period": package["implementation_period"],
                "rationale": "用户确认按独立交付和验收边界拆分",
                "template_id": base_group["template_id"],
                "template_version": base_group["template_version"],
                "template_match_basis": base_group["template_match_basis"],
                "status": "active",
                "package_ids": [package["id"]],
            }
        )
    payload = _structure_payload(plan, groups=groups)
    updated = authenticated_client.put(f"/api/v1/procurement-plans/{plan['id']}/structure", json=payload)
    assert updated.status_code == 200, updated.text
    assert updated.json()["recommended_document_count"] == 2
    refreshed = authenticated_client.get(f"/api/v1/procurement-plans/{plan['id']}")
    assert refreshed.json()["recommended_document_count"] == 2
    assert len(refreshed.json()["document_groups"]) == 2


def test_confirmed_plan_batch_generation_is_idempotent_and_source_change_marks_stale(
    authenticated_client: TestClient,
) -> None:
    project_id, plan = _plan(
        authenticated_client,
        [
            "独立主招标文件：软件建设招标文件|软件建设|独立交付",
            "采购包：P1|软件包|应用软件",
            "独立主招标文件：硬件采购招标文件|硬件采购|独立验收",
            "采购包：P2|硬件包|硬件设备",
        ],
    )
    confirmed = _confirm(authenticated_client, plan)
    key = f"batch-{uuid.uuid4()}"
    payload = {"group_ids": [item["id"] for item in confirmed["document_groups"]], "idempotency_key": key}
    first = authenticated_client.post(f"/api/v1/procurement-plans/{plan['id']}/generate-batch", json=payload)
    assert first.status_code == 202, first.text
    assert first.json()["succeeded_count"] == 2
    second = authenticated_client.post(f"/api/v1/procurement-plans/{plan['id']}/generate-batch", json=payload)
    assert second.status_code == 202
    assert second.json()["id"] == first.json()["id"]
    documents = authenticated_client.get(f"/api/v1/documents?project_id={project_id}").json()
    assert len([item for item in documents if item["stage"] == "tender"]) == 2

    _upload_and_bind(
        authenticated_client,
        project_id,
        ["主招标文件：新方案招标文件|新范围|来源已变化", "采购包：P9|新采购包|新范围"],
    )
    stale = authenticated_client.get(f"/api/v1/procurement-plans/{plan['id']}").json()
    assert stale["status"] == "stale"
    assert stale["draft_generation_allowed"] is False


def test_procurement_generated_draft_exports_real_docx_and_pdf(
    authenticated_client: TestClient,
) -> None:
    project_id, plan = _plan(
        authenticated_client,
        [
            "主招标文件：采购导出验收招标文件|应用系统建设和配套实施|统一交付验收",
            "采购包：P1|应用系统建设包|应用系统建设和配套实施",
        ],
    )
    confirmed = _confirm(authenticated_client, plan)
    generated = authenticated_client.post(
        f"/api/v1/procurement-plans/{plan['id']}/generate-batch",
        json={
            "group_ids": [confirmed["document_groups"][0]["id"]],
            "idempotency_key": f"batch-export-{uuid.uuid4()}",
        },
    )
    assert generated.status_code == 202, generated.text
    assert generated.json()["succeeded_count"] == 1
    tender = next(
        item
        for item in authenticated_client.get(f"/api/v1/documents?project_id={project_id}&stage=tender").json()
        if item["procurement_plan_id"] == confirmed["id"]
    )
    detail = authenticated_client.get(f"/api/v1/documents/{tender['id']}").json()
    version_id = detail["versions"][0]["id"]

    def export(output_format: str) -> bytes:
        response = authenticated_client.post(
            f"/api/v1/exports?version_id={version_id}",
            json={
                "output_format": output_format,
                "idempotency_key": f"procurement-export-{output_format}-{uuid.uuid4()}",
            },
        )
        assert response.status_code == 202, response.text
        job = response.json()
        assert job["status"] == "succeeded", job
        download = authenticated_client.get(f"/api/v1/exports/{job['id']}/download")
        assert download.status_code == 200, download.text
        return bytes(download.content)

    docx_content = export("docx")
    with zipfile.ZipFile(io.BytesIO(docx_content)) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert "草稿" in document_xml
    assert "应用系统建设和配套实施" in document_xml

    if shutil.which("soffice") is None and shutil.which("libreoffice") is None:
        return
    pdf_content = export("pdf")
    reader = PdfReader(io.BytesIO(pdf_content))
    assert len(reader.pages) >= 2
    pdf_text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "草稿" in pdf_text
    assert "应用系统建设和配套实施" in pdf_text


def test_batch_generation_keeps_successes_and_retries_only_failed_items(
    authenticated_client: TestClient,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _project_id, plan = _plan(
        authenticated_client,
        [
            "独立主招标文件：软件建设招标文件|软件建设|独立交付",
            "采购包：P1|软件包|应用软件",
            "独立主招标文件：硬件采购招标文件|硬件采购|独立验收",
            "采购包：P2|硬件包|硬件设备",
        ],
    )
    confirmed = _confirm(authenticated_client, plan)
    generation_task = cast(Any, api_module.generate_document_task)  # type: ignore[attr-defined]
    original_delay = generation_task.delay
    dispatched = 0

    def flaky_delay(job_id: str) -> object:
        nonlocal dispatched
        dispatched += 1
        if dispatched == 2:
            raise RuntimeError("synthetic queue dispatch failure")
        return original_delay(job_id)

    monkeypatch.setattr(generation_task, "delay", flaky_delay)
    created = authenticated_client.post(
        f"/api/v1/procurement-plans/{plan['id']}/generate-batch",
        json={
            "group_ids": [item["id"] for item in confirmed["document_groups"]],
            "idempotency_key": f"batch-partial-{uuid.uuid4()}",
        },
    )
    assert created.status_code == 202, created.text
    assert created.json()["status"] == "partial_failed"
    assert created.json()["succeeded_count"] == 1
    assert created.json()["failed_count"] == 1

    monkeypatch.setattr(generation_task, "delay", original_delay)
    retried = authenticated_client.post(
        f"/api/v1/procurement-generation-batches/{created.json()['id']}/retry-failed"
    )
    assert retried.status_code == 202, retried.text
    assert retried.json()["status"] == "succeeded"
    assert retried.json()["succeeded_count"] == 2
    assert retried.json()["failed_count"] == 0


def test_missing_template_blocks_generation_but_not_scope_confirmation(
    authenticated_client: TestClient,
) -> None:
    _project_id, plan = _plan(
        authenticated_client,
        ["主招标文件：平台招标文件|平台范围|整体交付", "采购包：P1|平台包|平台范围"],
    )
    groups = _structure_payload(plan)["document_groups"]
    groups[0]["template_id"] = None
    groups[0]["template_version"] = None
    groups[0]["template_match_basis"] = "未匹配"
    updated = authenticated_client.put(
        f"/api/v1/procurement-plans/{plan['id']}/structure",
        json={**_structure_payload(plan), "document_groups": groups},
    )
    assert updated.status_code == 200, updated.text
    confirmed = _confirm(authenticated_client, updated.json())
    assert confirmed["draft_generation_allowed"] is False
    generated = authenticated_client.post(
        f"/api/v1/procurement-plans/{plan['id']}/generate-batch",
        json={"group_ids": [groups[0]["id"]], "idempotency_key": f"batch-{uuid.uuid4()}"},
    )
    assert generated.status_code == 422


def test_contract_generation_inherits_only_selected_package_scope(
    authenticated_client: TestClient,
) -> None:
    project_id, plan = _plan(
        authenticated_client,
        [
            "主招标文件：信息化建设招标文件|应用系统和配套设备|统一采购、分包验收",
            "采购包：P1|应用系统包|应用系统建设范围",
            "采购包：P2|配套设备包|配套设备采购范围",
        ],
    )
    payload = _structure_payload(plan)
    for index, package in enumerate(payload["packages"], start=1):
        package["confirmed_budget"] = str(index * 100000)
        package["maximum_price"] = str(index * 90000)
        package["budget_status"] = "user_confirmed"
        package["budget_basis"] = "测试用户确认的脱敏预算依据"
    updated = authenticated_client.put(
        f"/api/v1/procurement-plans/{plan['id']}/structure",
        json=payload,
    )
    assert updated.status_code == 200, updated.text
    confirmed = _confirm(authenticated_client, updated.json())
    group = confirmed["document_groups"][0]
    selected_package, omitted_package = confirmed["packages"]

    with SessionLocal() as db:
        project = db.get(Project, project_id)
        assert project is not None
        tender_document = Document(
            organization_id=project.organization_id,
            project_id=project.id,
            stage="tender",
            title=group["name"],
            status="finalized",
            current_version=1,
            procurement_plan_id=confirmed["id"],
            procurement_document_group_id=group["id"],
            procurement_package_ids=group["package_ids"],
        )
        db.add(tender_document)
        db.flush()
        tender_version = DocumentVersion(
            organization_id=project.organization_id,
            document_id=tender_document.id,
            version=1,
            status="finalized",
            immutable=True,
            provenance={"procurement_plan_id": confirmed["id"]},
        )
        db.add(tender_version)
        db.flush()
        section = DocumentSection(
            organization_id=project.organization_id,
            document_version_id=tender_version.id,
            sequence=1,
            key="procurement_scope",
            title="采购范围",
            level=1,
        )
        db.add(section)
        db.flush()
        db.add(
            DocumentContentBlock(
                organization_id=project.organization_id,
                document_section_id=section.id,
                sequence=1,
                block_type="paragraph",
                content={"text": f"{selected_package['scope']}；{omitted_package['scope']}"},
                source_kind="confirmed_procurement_plan",
            )
        )
        contract_stage = db.scalar(
            select(ProjectStage).where(
                ProjectStage.project_id == project_id,
                ProjectStage.stage == "contract",
            )
        )
        assert contract_stage is not None
        contract_stage_revision = contract_stage.revision
        tender_version_id = tender_version.id
        db.commit()

    source = authenticated_client.put(
        f"/api/v1/projects/{project_id}/stages/contract/source",
        json={
            "source_type": "upstream_final",
            "source_file_version_id": tender_version_id,
            "revision": contract_stage_revision,
        },
    )
    assert source.status_code == 200, source.text
    templates = authenticated_client.get("/api/v1/templates?stage=contract&generation_only=true").json()
    assert templates
    template = templates[0]
    generated = authenticated_client.post(
        f"/api/v1/generation-jobs?project_id={project_id}&stage=contract",
        json={
            "template_id": template["id"],
            "template_version": template["current_version"],
            "idempotency_key": f"contract-selected-package-{uuid.uuid4()}",
            "procurement_plan_id": confirmed["id"],
            "procurement_document_group_id": group["id"],
            "procurement_package_ids": [selected_package["id"]],
        },
    )
    assert generated.status_code == 202, generated.text
    assert generated.json()["status"] == "succeeded"
    contract_documents = [
        item
        for item in authenticated_client.get(
            f"/api/v1/documents?project_id={project_id}&stage=contract"
        ).json()
        if item["stage"] == "contract"
    ]
    assert len(contract_documents) == 1
    assert contract_documents[0]["procurement_package_ids"] == [selected_package["id"]]
    detail = authenticated_client.get(f"/api/v1/documents/{contract_documents[0]['id']}").json()
    snapshot = detail["versions"][0]["provenance"]["procurement_snapshot"]
    assert snapshot["package_ids"] == [selected_package["id"]]
    assert selected_package["scope"] in snapshot["scope"]
    assert omitted_package["scope"] not in snapshot["scope"]
    assert "final_contract_amount" not in snapshot
    assert "supplier" not in snapshot


def test_project_isolation_blocks_non_member_from_plan(authenticated_client: TestClient) -> None:
    _project_id, plan = _plan(
        authenticated_client,
        ["主招标文件：隔离测试招标文件|测试范围|测试依据", "采购包：P1|测试包|测试范围"],
    )
    created = authenticated_client.post(
        "/api/v1/users",
        json={
            "email": "isolated@example.com",
            "display_name": "隔离用户",
            "password": "safe-password-123",
            "role_keys": ["project_editor"],
        },
    )
    assert created.status_code == 201, created.text
    isolated = TestClient(app)
    login = isolated.post(
        "/api/v1/auth/login",
        json={"email": "isolated@example.com", "password": "safe-password-123"},
    )
    assert login.status_code == 200, login.text
    isolated.headers["X-CSRF-Token"] = login.json()["csrf_token"]
    forbidden = isolated.get(f"/api/v1/procurement-plans/{plan['id']}")
    assert forbidden.status_code == 403


def test_published_regulatory_rule_is_versioned_and_traced_without_assuming_applicability(
    authenticated_client: TestClient,
) -> None:
    created = authenticated_client.post(
        "/api/v1/procurement-rule-sets",
        json={
            "key": "regional-procurement-policy",
            "name": "地区采购管理规则",
            "region": "测试地区",
            "funding_nature": "财政资金",
            "source_name": "脱敏测试采购管理规则",
            "effective_date": "2026-01-01",
            "rules_json": {
                "constraints": [
                    {
                        "code": "manual-review",
                        "description": "适用性由有权限用户结合主体和资金性质确认",
                    }
                ]
            },
        },
    )
    assert created.status_code == 201, created.text
    published = authenticated_client.post(
        f"/api/v1/procurement-rule-sets/{created.json()['id']}/publish",
        json={"revision": created.json()["revision"]},
    )
    assert published.status_code == 200, published.text
    assert published.json()["version"] == 1
    assert published.json()["status"] == "published"

    _project_id, plan = _plan(
        authenticated_client,
        [
            "主招标文件：规则追溯测试招标文件|测试采购范围|来源明确",
            "采购包：P1|测试采购包|测试采购范围",
        ],
    )
    traces = plan["analysis_coverage"]["candidate_rule_sets"]
    regional_trace = next(item for item in traces if item["key"] == "regional-procurement-policy")
    assert regional_trace == {
        "id": published.json()["id"],
        "key": "regional-procurement-policy",
        "name": "地区采购管理规则",
        "version": 1,
        "source_name": "脱敏测试采购管理规则",
        "source_url": None,
        "effective_date": "2026-01-01",
        "applicability_status": "requires_business_confirmation",
    }
    applicability = next(
        item for item in plan["unresolved_items"] if item["category"] == "regulatory_applicability"
    )
    assert applicability["severity"] == "P1"
    assert plan["recommended_document_count"] == 1
