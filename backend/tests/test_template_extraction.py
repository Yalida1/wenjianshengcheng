from __future__ import annotations

import io
import json
from typing import Any

from docx import Document
from fastapi.testclient import TestClient

from backend.app.services.providers import (
    OpenAICompatibleProvider,
    TemplateExtractionContext,
)


def _finished_feasibility_docx() -> bytes:
    stream = io.BytesIO()
    document = Document()
    document.add_heading("测试项目可行性研究报告", level=0)
    document.add_paragraph("项目名称：宁夏测试项目")
    document.add_paragraph("建设单位：测试建设单位")
    document.add_heading("第一章 项目概况", level=1)
    document.add_paragraph("本项目用于验证模板反向提取流程。")
    document.add_paragraph("建设地点位于测试园区。")
    document.add_heading("第二章 建设方案", level=1)
    document.add_paragraph("建设内容包括系统改造与制度审查能力建设。")
    document.add_heading("第三章 投资估算", level=1)
    document.add_paragraph("总投资：1000万元")
    document.save(stream)
    return stream.getvalue()


def _thin_unqualified_docx() -> bytes:
    stream = io.BytesIO()
    document = Document()
    document.add_paragraph("这是一份缺少正式章节结构的草稿备忘。")
    document.add_paragraph("金额待定。")
    document.save(stream)
    return stream.getvalue()


def test_finished_docx_becomes_reviewable_template_candidate(
    authenticated_client: TestClient,
) -> None:
    created = authenticated_client.post(
        "/api/v1/template-extractions",
        params={"stage": "feasibility", "authorized_external_processing": "true"},
        files={
            "upload": (
                "已完成可研.docx",
                _finished_feasibility_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert created.status_code == 201, created.text
    job = created.json()
    assert job["status"] == "review_required"
    assert job["provider_name"] == "demo"
    assert job["model_name"] == "deterministic-v1"
    assert job["result_json"]["summary"]["layout_preserved"] is True
    assert len(job["result_json"]["sections"]) >= 2
    assert {item["variable_key"] for item in job["result_json"]["variables"]} >= {
        "project_name",
        "project_owner",
        "total_investment",
    }

    selected_sections = [item["id"] for item in job["result_json"]["sections"]]
    selected_variables = [item["id"] for item in job["result_json"]["variables"]]
    confirmed = authenticated_client.post(
        f"/api/v1/template-extractions/{job['id']}/confirm",
        json={
            "revision": job["revision"],
            "template_name": "测试可研提取模板",
            "source_kind": "platform_reference_template",
            "issuing_authority": None,
            "document_number": None,
            "publish_year": None,
            "source_url": None,
            "applicability": "人工确认后的测试模板",
            "selected_section_ids": selected_sections,
            "selected_variable_ids": selected_variables,
        },
    )
    assert confirmed.status_code == 200, confirmed.text
    template = confirmed.json()
    assert template["status"] == "published"
    assert template["source_kind"] == "platform_reference_template"
    assert template["has_docx_source"] is True
    assert job["result_json"]["quality"]["passed"] is True
    assert job["result_json"]["quality"]["grade"] == "publishable"

    sections = authenticated_client.get(
        f"/api/v1/templates/{template['id']}/versions/1/sections"
    )
    variables = authenticated_client.get(
        f"/api/v1/templates/{template['id']}/versions/1/variables"
    )
    source = authenticated_client.get(
        f"/api/v1/templates/{template['id']}/versions/1/source"
    )
    assert sections.status_code == 200
    assert variables.status_code == 200
    assert source.status_code == 200
    assert len(sections.json()) == len(selected_sections)
    assert all("level" in item for item in sections.json())
    assert all(item.get("level", 1) >= 1 for item in sections.json())
    assert len(variables.json()) == len(selected_variables)
    extracted_document = Document(io.BytesIO(source.content))
    extracted_text = "\n".join(paragraph.text for paragraph in extracted_document.paragraphs)
    assert "{{DOCUMENT_BODY}}" in extracted_text
    assert "宁夏测试项目" not in extracted_text
    assert "1000万元" not in extracted_text


def test_unqualified_extraction_is_rejected_and_not_confirmable(
    authenticated_client: TestClient,
) -> None:
    created = authenticated_client.post(
        "/api/v1/template-extractions",
        params={"stage": "feasibility", "authorized_external_processing": "true"},
        files={
            "upload": (
                "不合格草稿.docx",
                _thin_unqualified_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert created.status_code == 201, created.text
    job = created.json()
    assert job["status"] == "quality_rejected"
    assert job["result_json"]["quality"]["passed"] is False
    assert job["result_json"]["quality"]["grade"] == "rejected"
    assert job["error"]

    confirmed = authenticated_client.post(
        f"/api/v1/template-extractions/{job['id']}/confirm",
        json={
            "revision": job["revision"],
            "template_name": "不应发布的模板",
            "source_kind": "platform_reference_template",
            "issuing_authority": None,
            "document_number": None,
            "publish_year": None,
            "source_url": None,
            "applicability": None,
            "selected_section_ids": ["section-1"],
            "selected_variable_ids": [],
        },
    )
    assert confirmed.status_code == 409
    assert confirmed.json()["error"]["code"] == "template_extraction_quality_rejected"

    templates = authenticated_client.get(
        "/api/v1/templates",
        params={"current_only": "false", "stage": "feasibility"},
    )
    assert templates.status_code == 200
    assert all(item["name"] != "不应发布的模板" for item in templates.json())


def test_template_extraction_requires_external_processing_consent(
    authenticated_client: TestClient,
) -> None:
    response = authenticated_client.post(
        "/api/v1/template-extractions",
        params={"stage": "feasibility", "authorized_external_processing": "false"},
        files={
            "upload": (
                "已完成可研.docx",
                _finished_feasibility_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "external_processing_consent_required"


def test_deepseek_provider_uses_json_object_and_validates_response(monkeypatch: Any) -> None:
    captured: dict[str, Any] = {}

    class FakeResponse:
        def raise_for_status(self) -> None:
            return None

        def json(self) -> dict[str, Any]:
            return {
                "choices": [
                    {
                        "message": {
                            "content": json.dumps(
                                {
                                    "sections": [
                                        {
                                            "key": "overview",
                                            "title": "项目概况",
                                            "source_block_ids": ["block-1"],
                                            "confidence": 0.9,
                                        }
                                    ],
                                    "variables": [],
                                    "warnings": [],
                                },
                                ensure_ascii=False,
                            )
                        }
                    }
                ]
            }

    class FakeClient:
        def __init__(self, timeout: int) -> None:
            captured["timeout"] = timeout

        def __enter__(self) -> FakeClient:
            return self

        def __exit__(self, *_args: object) -> None:
            return None

        def post(self, url: str, *, headers: dict[str, str], json: dict[str, Any]) -> FakeResponse:
            captured["url"] = url
            captured["headers"] = headers
            captured["payload"] = json
            return FakeResponse()

    monkeypatch.setattr("backend.app.services.providers.httpx.Client", FakeClient)
    provider = object.__new__(OpenAICompatibleProvider)
    provider.base_url = "https://api.deepseek.com"
    provider.api_key = "test-only-key"
    provider.model = "deepseek-v4-flash"
    result = provider.extract_template(
        TemplateExtractionContext(
            stage="feasibility",
            filename="sample.docx",
            blocks=[
                {
                    "id": "block-1",
                    "sequence": 1,
                    "kind": "heading",
                    "text": "项目概况",
                }
            ],
        )
    )
    assert result.sections[0].title == "项目概况"
    assert captured["payload"]["response_format"] == {"type": "json_object"}
    assert captured["url"] == "https://api.deepseek.com/chat/completions"
