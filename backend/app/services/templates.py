from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass

from docx import Document as WordDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DEMO_TEMPLATE_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


@dataclass(frozen=True)
class TemplatePreflight:
    valid: bool
    placeholders: tuple[str, ...]
    warnings: tuple[str, ...]


def _add_page_number(paragraph: object) -> None:
    target = paragraph  # python-docx has no public field API.
    run = target.add_run("第 ")  # type: ignore[attr-defined]
    field_run = target.add_run()  # type: ignore[attr-defined]
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run._r.extend([begin, instruction, end])
    run = target.add_run(" 页")  # type: ignore[attr-defined]
    run.font.size = Pt(9)


def build_template_document(
    stage: str,
    name: str,
    *,
    source_label: str,
    disclaimer: str,
) -> bytes:
    doc = WordDocument()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.6)
    section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.6)

    normal = doc.styles["Normal"]
    normal.font.name = "SimSun"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.5
    for style_name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 14)):
        style = doc.styles[style_name]
        style.font.name = "SimHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.add_run(f"{name} · {source_label}")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("项目文件链式生成平台 · ")
    _add_page_number(footer)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("{{DOCUMENT_TITLE}}")
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("版本 {{DOCUMENT_VERSION}}")
    warning = doc.add_paragraph()
    warning.alignment = WD_ALIGN_PARAGRAPH.CENTER
    warning.add_run(disclaimer).italic = True
    doc.add_page_break()  # type: ignore[no-untyped-call]
    doc.add_paragraph("{{DOCUMENT_BODY}}")

    doc.core_properties.title = name
    doc.core_properties.subject = f"{stage} {source_label}"
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def build_demo_template(stage: str, name: str) -> bytes:
    """Backward-compatible platform reference template builder."""
    return build_template_document(
        stage,
        name,
        source_label="平台参考模板",
        disclaimer="平台参考模板，仅用于生成和编制初稿，不属于国家正式文本",
    )


def preflight_docx_template(content: bytes) -> TemplatePreflight:
    warnings: list[str] = []
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if archive.testzip() is not None or "word/document.xml" not in archive.namelist():
                return TemplatePreflight(False, (), ("DOCX OOXML 包不完整",))
        doc = WordDocument(io.BytesIO(content))
    except (OSError, ValueError, zipfile.BadZipFile):
        return TemplatePreflight(False, (), ("文件不是可读取的 DOCX 模板",))
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    placeholders = tuple(
        token
        for token in ("{{DOCUMENT_TITLE}}", "{{DOCUMENT_VERSION}}", "{{DOCUMENT_BODY}}")
        if token in text
    )
    if "{{DOCUMENT_BODY}}" not in placeholders:
        warnings.append("模板缺少 DOCUMENT_BODY 装配锚点")
    if not doc.sections:
        warnings.append("模板没有页面节定义")
    return TemplatePreflight(not warnings, placeholders, tuple(warnings))
