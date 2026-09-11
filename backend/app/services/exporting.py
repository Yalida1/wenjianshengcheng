from __future__ import annotations

import hashlib
import io
import json
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from docx import Document as WordDocument
from docx.document import Document as DocxDocument
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import (
    Document,
    DocumentContentBlock,
    DocumentSection,
    DocumentVersion,
    FieldValue,
    Project,
    Template,
    TemplateVersion,
)
from .providers import format_field_value
from .storage import get_storage

FIELD_STATUS_LABELS = {
    "missing": "待补充",
    "extracted": "已提取",
    "ai_suggested": "AI 建议待确认",
    "user_confirmed": "已人工确认",
    "system_authoritative": "系统权威值",
    "template_default": "模板默认值",
    "conflict": "存在冲突",
    "invalid": "校验无效",
}
FIELD_SOURCE_LABELS = {
    "missing": "暂无来源",
    "extracted": "文件提取",
    "ai_suggestion": "AI 建议",
    "user_input": "人工录入",
    "system": "系统数据",
    "template_default": "模板默认值",
}


def _content_text(block: DocumentContentBlock) -> str:
    if isinstance(block.content, dict):
        return str(block.content.get("text", ""))
    return str(block.content)


def _normalize_formal_text(text: str) -> str:
    text = re.sub(r"。\s*[，；]\s*", "；", text)
    text = re.sub(r"；\s*。", "。", text)
    text = re.sub(r"。\s*。+", "。", text)
    text = re.sub(r"；\s*；+", "；", text)
    return text.strip()


def _content_search_text(block: DocumentContentBlock) -> str:
    if isinstance(block.content, (dict, list)):
        return json.dumps(block.content, ensure_ascii=False)
    return str(block.content)


def load_export_data(
    db: Session, version: DocumentVersion
) -> tuple[Document, list[tuple[DocumentSection, list[DocumentContentBlock]]], list[FieldValue]]:
    document = db.get(Document, version.document_id)
    if document is None:
        raise ValueError("Document not found")
    sections = list(
        db.scalars(
            select(DocumentSection)
            .where(DocumentSection.document_version_id == version.id)
            .order_by(DocumentSection.sequence)
        )
    )
    grouped = []
    for section in sections:
        blocks = list(
            db.scalars(
                select(DocumentContentBlock)
                .where(DocumentContentBlock.document_section_id == section.id)
                .order_by(DocumentContentBlock.sequence)
            )
        )
        grouped.append((section, blocks))
    fields = list(
        db.scalars(
            select(FieldValue)
            .where(
                FieldValue.project_id == document.project_id,
                FieldValue.stage == document.stage,
                FieldValue.is_current.is_(True),
            )
            .order_by(FieldValue.criticality, FieldValue.field_key)
        )
    )
    return document, grouped, fields


def _set_cell_shading(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    tc_pr.append(shade)


def _set_cell_margins(cell: Any, value: int = 100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_repeat_table_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def _prevent_row_split(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def _field_display_value(field: FieldValue) -> str:
    value = field.normalized_value if field.normalized_value is not None else field.value
    formatted = format_field_value(value)
    if isinstance(value, list):
        return formatted.replace("；", "\n")
    if field.unit == "%" and not formatted.rstrip().endswith("%"):
        return f"{formatted}%"
    if field.unit == "元" and isinstance(value, (int, float)):
        return f"{value:,.0f} 元"
    return formatted


def _add_page_field(paragraph: Any) -> None:
    paragraph.add_run("第 ")
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])
    paragraph.add_run(" 页")


def _load_template_document(db: Session, version: DocumentVersion) -> DocxDocument:
    template_id = version.provenance.get("template_id")
    template_version_number = version.provenance.get("template_version")
    if not isinstance(template_id, str) or not isinstance(template_version_number, int):
        raise RuntimeError("文档版本未锁定模板版本，不能导出 DOCX")
    template_version = db.scalar(
        select(TemplateVersion).where(
            TemplateVersion.template_id == template_id,
            TemplateVersion.version == template_version_number,
        )
    )
    if template_version is None or not template_version.storage_key:
        raise RuntimeError("锁定的 DOCX 模板源不存在，不能导出")
    content = get_storage().get(template_version.storage_key)
    return WordDocument(io.BytesIO(content))


def _replace_template_tokens(doc: DocxDocument, *, title: str, version: int) -> None:
    replacements = {
        "{{DOCUMENT_TITLE}}": title,
        "{{DOCUMENT_VERSION}}": str(version),
        "{{DOCUMENT_BODY}}": "",
    }
    for paragraph in doc.paragraphs:
        for token, value in replacements.items():
            if token in paragraph.text:
                for run in paragraph.runs:
                    if token in run.text:
                        run.text = run.text.replace(token, value)
                if token in paragraph.text:
                    paragraph.text = paragraph.text.replace(token, value)


def _clear_document_body(doc: DocxDocument) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _set_east_asia_font(run: Any, name: str, size: float | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)


def _set_table_borders(table: Any, color: str = "BFBFBF") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), color)


def _add_complex_field(paragraph: Any, instruction_text: str, display_text: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = f" {instruction_text} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = display_text
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def _bookmark_name(section_id: str) -> str:
    return f"toc_{section_id.replace('-', '_')}"


def _add_bookmark(paragraph: Any, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _heading_texts(
    sections: list[tuple[DocumentSection, list[DocumentContentBlock]]],
    labels: dict[str, str],
) -> dict[str, str]:
    chinese_chapters = "一二三四五六七八九十"
    root_index = 0
    result: dict[str, str] = {}
    for section, _blocks in sections:
        level = max(1, min(3, int(section.level or 1)))
        if level == 1:
            root_index += 1
            prefix = (
                f"第{chinese_chapters[root_index - 1]}章"
                if root_index <= len(chinese_chapters)
                else f"第{root_index}章"
            )
            result[section.id] = f"{prefix} {section.title}"
        else:
            result[section.id] = f"{labels.get(section.id, str(section.sequence))} {section.title}"
    return result


def _add_toc_entries(
    doc: DocxDocument,
    sections: list[tuple[DocumentSection, list[DocumentContentBlock]]],
    heading_texts: dict[str, str],
    page_numbers: dict[str, int] | None,
) -> None:
    for index, (section, _blocks) in enumerate(sections):
        level = max(1, min(3, int(section.level or 1)))
        paragraph = doc.add_paragraph()
        if index == 0:
            run = paragraph.add_run()
            begin = OxmlElement("w:fldChar")
            begin.set(qn("w:fldCharType"), "begin")
            instruction = OxmlElement("w:instrText")
            instruction.set(qn("xml:space"), "preserve")
            instruction.text = ' TOC \\o "1-3" \\h \\z \\u '
            separate = OxmlElement("w:fldChar")
            separate.set(qn("w:fldCharType"), "separate")
            run._r.extend([begin, instruction, separate])
        paragraph.paragraph_format.left_indent = Cm((level - 1) * 0.7)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.tab_stops.add_tab_stop(
            Cm(15.2), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )
        title_run = paragraph.add_run(heading_texts[section.id])
        _set_east_asia_font(title_run, "SimSun", 11 if level == 1 else 10.5)
        if level == 1:
            title_run.bold = True
        paragraph.add_run("\t")
        display_page = str((page_numbers or {}).get(section.id, 1))
        _add_complex_field(
            paragraph,
            f"PAGEREF {_bookmark_name(section.id)} \\h",
            display_page,
        )
        if index == len(sections) - 1:
            run = paragraph.add_run()
            end = OxmlElement("w:fldChar")
            end.set(qn("w:fldCharType"), "end")
            run._r.append(end)


def _resolve_heading_pages(pdf_content: bytes, heading_texts: dict[str, str]) -> dict[str, int]:
    reader = PdfReader(io.BytesIO(pdf_content))
    normalized_pages = [re.sub(r"\s+", "", page.extract_text() or "") for page in reader.pages]
    result: dict[str, int] = {}
    for section_id, heading_text in heading_texts.items():
        needle = re.sub(r"\s+", "", heading_text)
        for page_number, text in enumerate(normalized_pages, 1):
            if needle in text:
                result[section_id] = page_number
    return result


def _enable_field_updates(doc: DocxDocument) -> None:
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def _field_map(fields: list[FieldValue]) -> dict[str, FieldValue]:
    return {field.field_key: field for field in fields}


def _field_or_pending(fields: dict[str, FieldValue], key: str) -> str:
    field = fields.get(key)
    return _field_display_value(field) if field is not None else "【待确认】"


def _add_cover_line(doc: DocxDocument, label: str, value: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(7)
    label_run = paragraph.add_run(f"{label}：")
    label_run.bold = True
    _set_east_asia_font(label_run, "SimHei", 12)
    value_run = paragraph.add_run(value)
    _set_east_asia_font(value_run, "SimSun", 12)


def _add_structured_table(doc: DocxDocument, content: dict[str, Any]) -> None:
    caption = str(content.get("caption", "")).strip()
    headers = [str(item) for item in content.get("headers", [])]
    raw_rows = content.get("rows", [])
    rows = [[str(value) for value in row] for row in raw_rows if isinstance(row, list)]
    if caption:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(5)
        run = paragraph.add_run(caption)
        run.bold = True
        _set_east_asia_font(run, "SimHei", 11)
    if not headers:
        return
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table)
    _set_repeat_table_header(table.rows[0])
    width_values = content.get("column_widths", [])
    numeric_widths = [float(value) for value in width_values] if width_values else []
    if numeric_widths and sum(numeric_widths) > 15.4:
        scale = 15.4 / sum(numeric_widths)
        numeric_widths = [value * scale for value in numeric_widths]
    widths = [Cm(value) for value in numeric_widths]
    for index, (cell, header) in enumerate(zip(table.rows[0].cells, headers, strict=True)):
        cell.text = header
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_shading(cell, "E7E6E6")
        _set_cell_margins(cell, 90)
        if index < len(widths):
            cell.width = widths[index]
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            _set_east_asia_font(run, "SimHei", 9)
    for values in rows:
        row = table.add_row()
        _prevent_row_split(row)
        for index, cell in enumerate(row.cells):
            cell.text = values[index] if index < len(values) else ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell, 90)
            if index < len(widths):
                cell.width = widths[index]
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.15
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER if index == 0 or len(headers) <= 3 else WD_ALIGN_PARAGRAPH.LEFT
                )
                for run in paragraph.runs:
                    _set_east_asia_font(run, "SimSun", 9)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def _outline_labels(sections: list[tuple[DocumentSection, list[DocumentContentBlock]]]) -> dict[str, str]:
    counters: list[int] = []
    labels: dict[str, str] = {}
    for section, _blocks in sections:
        level = max(1, min(3, int(section.level or 1)))
        while len(counters) < level:
            counters.append(0)
        counters = counters[:level]
        counters[level - 1] += 1
        labels[section.id] = ".".join(str(value) for value in counters)
    return labels


def _configure_document_layout(doc: DocxDocument, *, preserve_customer_layout: bool) -> None:
    if not preserve_customer_layout:
        for section in doc.sections:
            section.page_width = Cm(21)
            section.page_height = Cm(29.7)
            section.top_margin = Cm(2.6)
            section.bottom_margin = Cm(2.4)
            section.left_margin = Cm(2.8)
            section.right_margin = Cm(2.6)
            section.header_distance = Cm(1.4)
            section.footer_distance = Cm(1.4)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "SimSun"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    for style_name, size in (("Title", 22), ("Heading 1", 18), ("Heading 2", 15), ("Heading 3", 13)):
        style = styles[style_name]
        style.font.name = "SimHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if style_name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(8 if style_name == "Heading 1" else 5)
        style_properties = style._element.get_or_add_pPr()
        paragraph_border = style_properties.find(qn("w:pBdr"))
        if paragraph_border is not None:
            style_properties.remove(paragraph_border)


def _assert_final_publication_content(
    sections: list[tuple[DocumentSection, list[DocumentContentBlock]]],
) -> None:
    prohibited = ("【待确认", "ai_generated", "人工录入", "请回到官方原文核对")
    for section, blocks in sections:
        if section.key == "announcement" and "或" in section.title:
            raise RuntimeError("定稿招标文件必须在招标公告与投标邀请书中确定一种")
        for block in blocks:
            text = _content_search_text(block)
            hit = next((token for token in prohibited if token in text), None)
            if hit:
                raise RuntimeError(f"定稿正文仍包含内部审查标记：{hit}")


def export_docx(
    db: Session,
    version: DocumentVersion,
    *,
    _toc_page_numbers: dict[str, int] | None = None,
    _resolve_toc: bool = True,
) -> bytes:
    document, sections, fields = load_export_data(db, version)
    project = db.get(Project, document.project_id)
    if project is None:
        raise ValueError("Project not found")
    if version.status == "finalized" and version.immutable:
        _assert_final_publication_content(sections)
    output = io.BytesIO()
    doc = _load_template_document(db, version)
    _clear_document_body(doc)
    template_id = version.provenance.get("template_id")
    template = db.get(Template, template_id) if isinstance(template_id, str) else None
    format_profile = version.provenance.get("format_profile")
    preserve_customer_layout = bool(
        template is not None
        and not template.is_builtin
        and isinstance(format_profile, dict)
        and format_profile.get("standard") == "customer_template"
    )
    _configure_document_layout(doc, preserve_customer_layout=preserve_customer_layout)
    _enable_field_updates(doc)
    is_draft = version.status != "finalized" or not version.immutable

    field_by_key = _field_map(fields)
    stage_labels = {
        "requirement": "项目建议书",
        "feasibility": "可行性研究报告",
        "tender": "招标文件",
        "contract": "合同文件",
    }
    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Pt(54)
    cover.paragraph_format.space_after = Pt(32)
    cover_title = (
        _field_or_pending(field_by_key, "project_name") if document.stage == "tender" else document.title
    )
    title_run = cover.add_run(cover_title)
    title_run.bold = True
    _set_east_asia_font(title_run, "SimHei", 24)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(38)
    subtitle_run = subtitle.add_run(stage_labels.get(document.stage, document.title))
    subtitle_run.bold = True
    _set_east_asia_font(subtitle_run, "SimHei", 26)
    if is_draft:
        draft_notice = doc.add_paragraph()
        draft_notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
        draft_notice.paragraph_format.space_after = Pt(26)
        run = draft_notice.add_run("草稿／不可用于发布")
        run.bold = True
        run.font.color.rgb = RGBColor(192, 0, 0)
        _set_east_asia_font(run, "SimHei", 16)
    _add_cover_line(doc, "项目编号", project.code)
    if document.stage == "tender":
        _add_cover_line(doc, "招标编号", _field_or_pending(field_by_key, "tender_number"))
        _add_cover_line(doc, "标段或包号", _field_or_pending(field_by_key, "package_number"))
        _add_cover_line(doc, "招标人", _field_or_pending(field_by_key, "tenderer"))
        _add_cover_line(doc, "招标代理机构", _field_or_pending(field_by_key, "tender_agency"))
        _add_cover_line(doc, "编制或发布日期", _field_or_pending(field_by_key, "issue_date"))

    labels = _outline_labels(sections)
    heading_texts = _heading_texts(sections, labels)

    doc.add_page_break()  # type: ignore[no-untyped-call]
    toc_heading = doc.add_paragraph(style="Title")
    toc_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    toc_heading.add_run("目录")
    _add_toc_entries(doc, sections, heading_texts, _toc_page_numbers)

    for bookmark_id, (item, blocks) in enumerate(sections, 1):
        heading_level = max(1, min(3, int(getattr(item, "level", 1) or 1)))
        heading = doc.add_heading(heading_texts[item.id], level=heading_level)
        heading.paragraph_format.keep_with_next = True
        # A paragraph-level page break starts a chapter without producing an
        # empty page when the previous chapter happens to end exactly at a
        # natural page boundary. An explicit break paragraph can skip a page
        # in that situation after LibreOffice pagination.
        heading.paragraph_format.page_break_before = heading_level == 1
        _add_bookmark(heading, _bookmark_name(item.id), bookmark_id)
        for block in blocks:
            if block.block_type == "table" and isinstance(block.content, dict):
                _add_structured_table(doc, block.content)
                continue
            if block.block_type == "page_break":
                doc.add_page_break()  # type: ignore[no-untyped-call]
                continue
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.first_line_indent = Cm(0.85)
            paragraph.paragraph_format.widow_control = True
            text = _content_text(block)
            if document.stage == "tender":
                text = _normalize_formal_text(text)
            run = paragraph.add_run(text)
            _set_east_asia_font(run, "SimSun", 12)

    for output_section in doc.sections:
        output_section.different_first_page_header_footer = True
        if not preserve_customer_layout or not any(p.text.strip() for p in output_section.header.paragraphs):
            header = output_section.header.paragraphs[0]
            header.alignment = WD_ALIGN_PARAGRAPH.CENTER
            header.clear()
            header_run = header.add_run(document.title)
            _set_east_asia_font(header_run, "SimSun", 9)
        if not preserve_customer_layout or not any(p.text.strip() for p in output_section.footer.paragraphs):
            footer = output_section.footer.paragraphs[0]
            footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
            footer.clear()
            _add_page_field(footer)
    doc.core_properties.title = document.title
    doc.core_properties.subject = "招标文件草稿" if is_draft else "招标文件定稿"
    doc.core_properties.keywords = (
        f"document-version:{version.id};template:{version.provenance.get('template_name', '')};"
        f"template-version:{version.provenance.get('template_version', '')};"
        f"template-sha256:{version.provenance.get('template_sha256', '')}"
    )
    doc.save(output)
    docx_content = output.getvalue()
    if _resolve_toc and (shutil.which("soffice") or shutil.which("libreoffice")):
        provisional_pdf = convert_docx_to_pdf(docx_content)
        resolved_pages = _resolve_heading_pages(provisional_pdf, heading_texts)
        if len(resolved_pages) == len(heading_texts):
            return export_docx(
                db,
                version,
                _toc_page_numbers=resolved_pages,
                _resolve_toc=False,
            )
    return docx_content


def convert_docx_to_pdf(docx_content: bytes) -> bytes:
    office_executable = shutil.which("soffice") or shutil.which("libreoffice")
    if office_executable is None:
        raise RuntimeError("未安装 LibreOffice，无法将 DOCX 转换为 PDF")

    with tempfile.TemporaryDirectory(prefix="docchain-pdf-") as temporary_directory:
        output_directory = Path(temporary_directory)
        docx_path = output_directory / "document.docx"
        pdf_path = output_directory / "document.pdf"
        profile_path = output_directory / "libreoffice-profile"
        docx_path.write_bytes(docx_content)
        completed = subprocess.run(  # noqa: S603 - executable is resolved from the deployment PATH
            [
                office_executable,
                "--headless",
                f"-env:UserInstallation={profile_path.resolve().as_uri()}",
                "--convert-to",
                "pdf:writer_pdf_Export",
                "--outdir",
                str(output_directory),
                str(docx_path),
            ],
            check=False,
            capture_output=True,
            text=True,
            timeout=120,
        )
        if completed.returncode != 0 or not pdf_path.is_file() or pdf_path.stat().st_size == 0:
            detail = (completed.stderr or completed.stdout or "未生成 PDF 文件").strip()
            raise RuntimeError(f"LibreOffice PDF 转换失败：{detail[:500]}")
        return pdf_path.read_bytes()


def export_pdf(db: Session, version: DocumentVersion) -> bytes:
    return convert_docx_to_pdf(export_docx(db, version))


def export_xlsx(db: Session, version: DocumentVersion) -> bytes:
    _document, sections, fields = load_export_data(db, version)
    workbook = Workbook()
    fields_sheet = workbook.active
    fields_sheet.title = "字段来源"
    fields_sheet.sheet_view.showGridLines = False
    fields_sheet.append(["字段键", "字段名称", "数据类型", "值", "单位", "级别", "状态", "来源类型"])
    for field in fields:
        fields_sheet.append(
            [
                field.field_key,
                field.field_label,
                field.data_type,
                str(field.normalized_value if field.normalized_value is not None else field.value),
                field.unit or "",
                field.criticality,
                field.status,
                field.source_type,
            ]
        )
    exported_captions: set[str] = set()
    allowed_captions = {"采购清单", "技术规格和性能指标", "商务、技术和报价评审因素表"}
    for _section, blocks in sections:
        for block in blocks:
            if block.block_type != "table" or not isinstance(block.content, dict):
                continue
            caption = str(block.content.get("caption", "")).strip()
            if caption not in allowed_captions or caption in exported_captions:
                continue
            exported_captions.add(caption)
            sheet = workbook.create_sheet(caption[:31])
            sheet.sheet_view.showGridLines = False
            headers = [str(item) for item in block.content.get("headers", [])]
            rows = [item for item in block.content.get("rows", []) if isinstance(item, list)]
            sheet.append(headers)
            for row in rows:
                sheet.append([str(value) for value in row])
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        border = Border(bottom=Side(style="thin", color="D9D9D9"))
        for row in sheet.iter_rows():
            for cell in row:
                cell.font = Font(name="Microsoft YaHei", size=10, bold=cell.row == 1, color=cell.font.color)
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                    horizontal="center" if cell.column <= 3 else "left",
                )
                cell.border = border
        widths = [16, 26, 18, 56, 18, 14, 20, 20]
        for index, width in enumerate(widths[: sheet.max_column], 1):
            sheet.column_dimensions[chr(64 + index)].width = width
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_version(db: Session, version: DocumentVersion, output_format: str) -> bytes:
    if output_format == "docx":
        return export_docx(db, version)
    if output_format == "pdf":
        return export_pdf(db, version)
    if output_format == "xlsx":
        return export_xlsx(db, version)
    raise ValueError(f"Unsupported export format: {output_format}")


def export_sha256(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def export_filename(db: Session, version: DocumentVersion, output_format: str) -> str:
    document = db.get(Document, version.document_id)
    if document is None:
        raise ValueError("Document not found")
    stem = "字段来源清单" if output_format == "xlsx" else document.title
    stem = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", stem).strip(" .") or "document"
    status_suffix = "" if version.status == "finalized" and version.immutable else "_草稿_不可用于发布"
    return f"{stem}_V{version.version}.0{status_suffix}.{output_format}"
