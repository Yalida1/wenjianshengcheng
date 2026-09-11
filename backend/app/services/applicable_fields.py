"""Resolve applicable fields for a tender document group.

Applicability belongs to the document group, not the whole tender stage catalog.
Priority: bound template variables → field profile → core / material-collection set.

Semantics:
- applicable: shown for this document group
- source_stage / input_role: where the value is expected to come from
- required_for_phase: material_collection vs formal_tender
- blocking: gates finalize when required and P0/P1
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Any, Literal

from sqlalchemy import select
from sqlalchemy.orm import Session

from ..field_catalog import BASE_FIELD_DEFINITIONS
from ..models import (
    FieldDefinition,
    FieldValue,
    ProcurementPlan,
    Template,
    TemplateVariable,
    TemplateVersion,
    TenderDocumentGroup,
)
from ..tender_field_profiles import (
    TENDER_MATERIAL_COLLECTION_FIELDS,
    InputRole,
    ProfileFieldSpec,
    ProfileName,
    RequiredForPhase,
    SourceStage,
    classify_empty_reason,
    empty_reason_message,
    field_input_role,
    field_required_for_phase,
    field_source_stage,
    profile_fields,
    resolve_profile_name,
)

ResolutionSource = Literal["template", "profile", "core", "mixed"]

_TEMPLATE_VAR_RE = re.compile(r"\{\{\s*([a-zA-Z_][a-zA-Z0-9_]*)\s*\}\}")
_ASSEMBLY_TOKENS = frozenset({"DOCUMENT_TITLE", "DOCUMENT_VERSION", "DOCUMENT_BODY"})
_CATALOG_KEYS = frozenset(item[0] for item in BASE_FIELD_DEFINITIONS.get("tender", []))
_DOCX_VAR_CACHE: dict[tuple[str, int, str], tuple[str, ...]] = {}
_TENDER_METHODS = frozenset({"public_tender", "invited_tender", "tender"})


@dataclass(frozen=True)
class ApplicableField:
    field_key: str
    label: str
    data_type: str
    unit: str | None
    level: str
    required: bool
    source: Literal["template", "profile", "core", "global"]
    reason: str
    profile: str
    document_group_id: str
    blocking: bool
    definition_missing: bool = False
    source_template_id: str | None = None
    source_profile: str | None = None
    display_order: int = 100
    empty_reason_code: str | None = None
    empty_reason_message: str | None = None
    source_stage: SourceStage = "upstream_or_material"
    required_for_phase: RequiredForPhase = "both"
    input_role: InputRole = "material"


@dataclass
class TemplateConfigIssue:
    variable_key: str
    message: str


@dataclass
class ApplicableFieldsResult:
    document_group_id: str
    profile: ProfileName
    resolution_source: ResolutionSource
    fields: list[ApplicableField] = field(default_factory=list)
    template_config_errors: list[TemplateConfigIssue] = field(default_factory=list)
    source_template_id: str | None = None
    source_template_version: int | None = None
    setup_incomplete: bool = False
    setup_incomplete_reason: str | None = None

    def required_keys(self) -> set[str]:
        return {item.field_key for item in self.fields if item.required}

    def blocking_keys(self) -> set[str]:
        return {item.field_key for item in self.fields if item.blocking}

    def field_keys(self) -> set[str]:
        return {item.field_key for item in self.fields}

    def as_dict(self) -> dict[str, Any]:
        return {
            "document_group_id": self.document_group_id,
            "profile": self.profile,
            "resolution_source": self.resolution_source,
            "source_template_id": self.source_template_id,
            "source_template_version": self.source_template_version,
            "setup_incomplete": self.setup_incomplete,
            "setup_incomplete_reason": self.setup_incomplete_reason,
            "template_config_errors": [
                {"variable_key": item.variable_key, "message": item.message}
                for item in self.template_config_errors
            ],
            "fields": [
                {
                    "field_key": item.field_key,
                    "label": item.label,
                    "data_type": item.data_type,
                    "unit": item.unit,
                    "level": item.level,
                    "required": item.required,
                    "source": item.source,
                    "reason": item.reason,
                    "profile": item.profile,
                    "document_group_id": item.document_group_id,
                    "blocking": item.blocking,
                    "definition_missing": item.definition_missing,
                    "source_template_id": item.source_template_id,
                    "source_profile": item.source_profile,
                    "display_order": item.display_order,
                    "empty_reason_code": item.empty_reason_code,
                    "empty_reason_message": item.empty_reason_message,
                    "source_stage": item.source_stage,
                    "required_for_phase": item.required_for_phase,
                    "input_role": item.input_role,
                }
                for item in self.fields
            ],
        }


def _catalog_meta(field_key: str) -> tuple[str, str, str | None, str, bool]:
    for key, label, data_type, unit, criticality, required in BASE_FIELD_DEFINITIONS.get(
        "tender", []
    ):
        if key == field_key:
            return label, data_type, unit, criticality, required
    return field_key, "string", None, "P2", False


def _definition_map(db: Session, organization_id: str) -> dict[str, FieldDefinition]:
    rows = list(
        db.scalars(
            select(FieldDefinition).where(
                FieldDefinition.organization_id == organization_id,
                FieldDefinition.stage == "tender",
                FieldDefinition.is_active.is_(True),
            )
        )
    )
    return {row.field_key: row for row in rows}


def _scan_docx_variables(content: bytes) -> tuple[str, ...]:
    try:
        from docx import Document

        document = Document(io.BytesIO(content))
    except Exception:  # noqa: BLE001
        return ()
    texts: list[str] = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.append(cell.text)
    blob = "\n".join(texts)
    found = [key for key in _TEMPLATE_VAR_RE.findall(blob) if key not in _ASSEMBLY_TOKENS]
    seen: set[str] = set()
    ordered: list[str] = []
    for key in found:
        if key not in seen:
            seen.add(key)
            ordered.append(key)
    return tuple(ordered)


def _docx_variables_cached(template_id: str, version: int, sha256: str, content: bytes) -> tuple[str, ...]:
    cache_key = (template_id, version, sha256)
    cached = _DOCX_VAR_CACHE.get(cache_key)
    if cached is not None:
        return cached
    result = _scan_docx_variables(content)
    _DOCX_VAR_CACHE[cache_key] = result
    return result


@dataclass
class _TemplateLoadResult:
    pairs: list[tuple[str, bool]]
    errors: list[TemplateConfigIssue]
    version: TemplateVersion | None
    db_variable_count: int = 0
    docx_variable_keys: tuple[str, ...] = ()
    docx_scan_attempted: bool = False
    docx_scan_failed: bool = False
    is_builtin_template: bool = False


def _load_template_variables(
    db: Session,
    group: TenderDocumentGroup,
    definitions: dict[str, FieldDefinition],
) -> _TemplateLoadResult:
    """Return template variable pairs from DB registry and/or DOCX scan."""

    if not group.template_id:
        return _TemplateLoadResult([], [], None)
    template = db.get(Template, group.template_id)
    if template is None:
        return _TemplateLoadResult(
            [],
            [
                TemplateConfigIssue(
                    variable_key="__template__",
                    message="绑定的模板不存在，属于模板配置问题，不能当作材料缺失处理。",
                )
            ],
            None,
        )

    version: TemplateVersion | None = None
    if group.template_version is not None:
        version = db.scalar(
            select(TemplateVersion).where(
                TemplateVersion.template_id == template.id,
                TemplateVersion.version == group.template_version,
            )
        )
    if version is None:
        version = db.scalar(
            select(TemplateVersion)
            .where(TemplateVersion.template_id == template.id)
            .order_by(TemplateVersion.version.desc())
            .limit(1)
        )
    if version is None:
        return _TemplateLoadResult(
            [],
            [
                TemplateConfigIssue(
                    variable_key="__template_version__",
                    message="模板版本缺失，属于模板配置问题，不能当作材料缺失处理。",
                )
            ],
            None,
            is_builtin_template=bool(getattr(template, "is_builtin", False)),
        )

    variables = list(
        db.scalars(
            select(TemplateVariable)
            .where(TemplateVariable.template_version_id == version.id)
            .order_by(TemplateVariable.variable_key)
        )
    )
    pairs: list[tuple[str, bool]] = []
    errors: list[TemplateConfigIssue] = []
    seen: set[str] = set()
    docx_keys: tuple[str, ...] = ()
    docx_attempted = False
    docx_failed = False

    def _register(key: str, required: bool) -> None:
        if not key or key in _ASSEMBLY_TOKENS:
            return
        known = key in _CATALOG_KEYS or key in definitions
        if not known:
            errors.append(
                TemplateConfigIssue(
                    variable_key=key,
                    message=(
                        f"模板变量“{key}”不在招标字段字典中，"
                        "属于模板配置问题，不能当作材料缺失处理。"
                    ),
                )
            )
            return
        if key not in seen:
            seen.add(key)
            pairs.append((key, required))

    for variable in variables:
        key = (variable.field_key or variable.variable_key or "").strip()
        _register(key, bool(variable.required))

    if version.storage_key:
        docx_attempted = True
        try:
            from .storage import get_storage

            content = get_storage().get(version.storage_key)
            sha = version.sha256 or version.storage_key
            docx_keys = _docx_variables_cached(template.id, version.version, sha, content)
            for key in docx_keys:
                _register(key, True)
        except Exception:  # noqa: BLE001
            docx_failed = True
            if not variables:
                errors.append(
                    TemplateConfigIssue(
                        variable_key="__template_parse__",
                        message="模板正文解析失败且无已登记变量，属于模板配置问题，不能当作材料缺失处理。",
                    )
                )

    return _TemplateLoadResult(
        pairs=pairs,
        errors=errors,
        version=version,
        db_variable_count=len(variables),
        docx_variable_keys=docx_keys,
        docx_scan_attempted=docx_attempted,
        docx_scan_failed=docx_failed,
        is_builtin_template=bool(
            getattr(template, "is_builtin", False)
            or getattr(template, "source_kind", "")
            in {"platform_reference_template", "adapted_from_official_outline"}
        ),
    )


def _is_builtin_bulk_registry(load: _TemplateLoadResult) -> bool:
    """True when stage-wide seed registered every catalog key (not real DOCX needs)."""

    if not load.pairs or not _CATALOG_KEYS:
        return False
    known = {key for key, _ in load.pairs if key in _CATALOG_KEYS}
    if known < _CATALOG_KEYS:
        return False
    # Real customer/DOCX templates that truly reference the whole catalog keep template mode.
    if load.docx_variable_keys and set(load.docx_variable_keys) >= _CATALOG_KEYS:
        return False
    if load.docx_scan_attempted and not load.docx_scan_failed and load.docx_variable_keys:
        # DOCX has a real (smaller or equal non-bulk) variable set already merged into pairs;
        # if pairs still equal full catalog solely from DB seed extras, treat as bulk when
        # DOCX subset is strictly smaller.
        if set(load.docx_variable_keys) < _CATALOG_KEYS:
            return True
    # Builtin / platform seed with full DB registry and no contradictory DOCX subset.
    return load.is_builtin_template or load.db_variable_count >= len(_CATALOG_KEYS)


def _compliance_context(db: Session, project_id: str) -> dict[str, bool]:
    """Detect upstream facts that must not auto-map into tender decision fields."""

    keys = ("total_investment", "construction_scope", "project_period")
    rows = list(
        db.execute(
            select(FieldValue.field_key, FieldValue.value, FieldValue.normalized_value).where(
                FieldValue.project_id == project_id,
                FieldValue.stage.in_(("feasibility", "requirement")),
                FieldValue.is_current.is_(True),
                FieldValue.field_key.in_(keys),
            )
        )
    )
    present: set[str] = set()
    for field_key, value, normalized in rows:
        payload = normalized if normalized not in (None, "", [], {}) else value
        if payload not in (None, "", [], {}):
            present.add(str(field_key))
    return {
        "has_total_investment": "total_investment" in present,
        "has_construction_scope": "construction_scope" in present,
        "has_project_period": "project_period" in present,
    }


def _build_field(
    *,
    field_key: str,
    required: bool,
    level: str,
    source: Literal["template", "profile", "core", "global"],
    reason: str,
    profile: ProfileName,
    group: TenderDocumentGroup,
    definitions: dict[str, FieldDefinition],
    template_id: str | None,
    display_order: int,
    compliance: dict[str, bool],
) -> ApplicableField:
    definition = definitions.get(field_key)
    if definition is not None:
        label = definition.field_label
        data_type = definition.data_type
        unit = definition.unit
        catalog_level = definition.criticality
        definition_missing = False
    else:
        label, data_type, unit, catalog_level, _ = _catalog_meta(field_key)
        definition_missing = field_key not in _CATALOG_KEYS
    effective_level = level or catalog_level
    if source == "template":
        source_stage: SourceStage = "template"
        input_role: InputRole = field_input_role(field_key)
        required_for_phase: RequiredForPhase = "formal_tender" if required else "material_collection"
    else:
        source_stage = field_source_stage(field_key)
        input_role = field_input_role(field_key)
        required_for_phase = field_required_for_phase(field_key, required=required)
    # Decision / process empties are not material-extraction failures.
    blocking = bool(required and effective_level in {"P0", "P1"})
    empty_code = classify_empty_reason(field_key, **compliance)
    return ApplicableField(
        field_key=field_key,
        label=label,
        data_type=data_type,
        unit=unit,
        level=effective_level,
        required=required,
        source=source,
        reason=reason,
        profile=profile,
        document_group_id=group.id,
        blocking=blocking,
        definition_missing=definition_missing,
        source_template_id=template_id,
        source_profile=profile,
        display_order=display_order,
        empty_reason_code=empty_code,
        empty_reason_message=empty_reason_message(empty_code, field_key),
        source_stage=source_stage,
        required_for_phase=required_for_phase,
        input_role=input_role,
    )


def _setup_state(
    *,
    profile: ProfileName,
    group: TenderDocumentGroup,
    load: _TemplateLoadResult,
) -> tuple[bool, str | None]:
    reasons: list[str] = []
    if not group.template_id:
        reasons.append("未绑定招标模板")
    elif load.version is None:
        reasons.append("模板版本缺失或无法解析")
    if profile == "generic":
        reasons.append("采购类别未知，招标配置未完成")
    if any(err.variable_key.startswith("__") for err in load.errors):
        reasons.append("模板配置存在阻断错误")
    if not reasons:
        return False, None
    return True, "；".join(reasons)


def resolve_applicable_fields(
    db: Session,
    *,
    project_id: str,
    document_group: TenderDocumentGroup,
    organization_id: str | None = None,
) -> ApplicableFieldsResult:
    org_id = organization_id or document_group.organization_id
    definitions = _definition_map(db, org_id)
    profile = resolve_profile_name(
        procurement_category=document_group.procurement_category,
        business_subcategory=document_group.business_subcategory,
        group_name=document_group.name,
    )
    compliance = _compliance_context(db, project_id)
    load = _load_template_variables(db, document_group, definitions)
    template_pairs = load.pairs
    template_errors = list(load.errors)
    template_version = load.version
    setup_incomplete, setup_reason = _setup_state(profile=profile, group=document_group, load=load)

    fields: list[ApplicableField] = []
    # Builtin full-catalog DB seed is not a real document field set. Prefer DOCX subset when present.
    if _is_builtin_bulk_registry(load):
        if load.docx_variable_keys:
            template_pairs = [
                (key, True)
                for key in load.docx_variable_keys
                if key in _CATALOG_KEYS or key in definitions
            ]
        else:
            template_pairs = []

    use_template = bool(template_pairs) and (
        not _is_builtin_bulk_registry(load) or bool(load.docx_variable_keys)
    )
    if use_template and template_pairs:
        for index, (key, required) in enumerate(template_pairs):
            meta = definitions.get(key)
            level = meta.criticality if meta is not None else _catalog_meta(key)[3]
            fields.append(
                _build_field(
                    field_key=key,
                    required=required,
                    level=level,
                    source="template",
                    reason="绑定模板变量",
                    profile=profile,
                    group=document_group,
                    definitions=definitions,
                    template_id=document_group.template_id,
                    display_order=(index + 1) * 10,
                    compliance=compliance,
                )
            )
        # Unknown template vars / fatal parse issues keep setup incomplete for gates.
        if template_errors:
            setup_incomplete = True
            setup_reason = setup_reason or "模板变量配置不完整"
        return ApplicableFieldsResult(
            document_group_id=document_group.id,
            profile=profile,
            resolution_source="template",
            fields=fields,
            template_config_errors=template_errors,
            source_template_id=document_group.template_id,
            source_template_version=template_version.version if template_version else None,
            setup_incomplete=setup_incomplete,
            setup_incomplete_reason=setup_reason,
        )

    # No usable template variable set: profile or minimal material collection.
    # Unknown category → material collection only. Known category keeps process clauses
    # for formal tender even before template bind; setup_incomplete still blocks finalize.
    if profile == "generic":
        specs: tuple[ProfileFieldSpec, ...] = TENDER_MATERIAL_COLLECTION_FIELDS
        source: ResolutionSource = "core"
        field_source: Literal["profile", "core"] = "core"
    else:
        specs = profile_fields(profile)
        source = "profile"
        field_source = "profile"

    for index, spec in enumerate(specs):
        fields.append(
            _build_field(
                field_key=spec.field_key,
                required=spec.required,
                level=spec.level,
                source=field_source,
                reason=spec.reason or f"profile:{profile}",
                profile=profile,
                group=document_group,
                definitions=definitions,
                template_id=None,
                display_order=(index + 1) * 10,
                compliance=compliance,
            )
        )
    return ApplicableFieldsResult(
        document_group_id=document_group.id,
        profile=profile,
        resolution_source=source,
        fields=fields,
        template_config_errors=template_errors,
        source_template_id=document_group.template_id,
        source_template_version=template_version.version if template_version else None,
        setup_incomplete=setup_incomplete,
        setup_incomplete_reason=setup_reason,
    )


def select_effective_procurement_plan(
    db: Session, project_id: str, *, plan_id: str | None = None
) -> ProcurementPlan | None:
    """Mirror frontend pickPreferredProcurementPlan: prefer confirmed/recommended with groups."""

    if plan_id:
        plan = db.get(ProcurementPlan, plan_id)
        if plan is not None and plan.project_id == project_id:
            return plan
        return None

    plans = list(
        db.scalars(select(ProcurementPlan).where(ProcurementPlan.project_id == project_id))
    )
    if not plans:
        return None

    groups_by_plan: dict[str, list[TenderDocumentGroup]] = {}
    plan_ids = [item.id for item in plans]
    for group in db.scalars(
        select(TenderDocumentGroup).where(
            TenderDocumentGroup.plan_id.in_(plan_ids),
            TenderDocumentGroup.status == "active",
        )
    ):
        groups_by_plan.setdefault(group.plan_id, []).append(group)

    def tender_group_count(plan: ProcurementPlan) -> int:
        return sum(
            1
            for item in groups_by_plan.get(plan.id, [])
            if item.procurement_method in _TENDER_METHODS
        )

    def with_groups(items: list[ProcurementPlan]) -> list[ProcurementPlan]:
        return [item for item in items if tender_group_count(item) > 0]

    confirmed_recommended = with_groups(
        [item for item in plans if item.status == "confirmed" and item.is_recommended]
    )
    if confirmed_recommended:
        return confirmed_recommended[0]

    any_confirmed = with_groups([item for item in plans if item.status == "confirmed"])
    if any_confirmed:
        return any_confirmed[0]

    recommended_with_groups = with_groups([item for item in plans if item.is_recommended])
    if recommended_with_groups:
        return recommended_with_groups[0]

    any_with_groups = with_groups(plans)
    if any_with_groups:
        return any_with_groups[0]

    recommended = next((item for item in plans if item.is_recommended), None)
    return recommended or plans[0]


def list_project_tender_groups(
    db: Session,
    project_id: str,
    *,
    plan_id: str | None = None,
) -> list[TenderDocumentGroup]:
    """Active tender document groups for the selected effective plan only."""

    plan = select_effective_procurement_plan(db, project_id, plan_id=plan_id)
    if plan is None:
        return []
    return list(
        db.scalars(
            select(TenderDocumentGroup)
            .where(
                TenderDocumentGroup.plan_id == plan.id,
                TenderDocumentGroup.status == "active",
            )
            .order_by(TenderDocumentGroup.code)
        )
    )


def resolve_project_applicable_field_union(
    db: Session,
    *,
    project_id: str,
    organization_id: str,
    plan_id: str | None = None,
) -> set[str]:
    """Union of applicable field keys across active groups of the effective plan."""

    groups = list_project_tender_groups(db, project_id, plan_id=plan_id)
    if not groups:
        return {spec.field_key for spec in TENDER_MATERIAL_COLLECTION_FIELDS}
    keys: set[str] = set()
    for group in groups:
        result = resolve_applicable_fields(
            db,
            project_id=project_id,
            document_group=group,
            organization_id=organization_id,
        )
        keys.update(result.field_keys())
    return keys


def get_document_group_for_project(
    db: Session, *, project_id: str, group_id: str
) -> TenderDocumentGroup | None:
    group = db.get(TenderDocumentGroup, group_id)
    if group is None:
        return None
    plan = db.get(ProcurementPlan, group.plan_id)
    if plan is None or plan.project_id != project_id:
        return None
    return group
