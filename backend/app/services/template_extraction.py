from __future__ import annotations

import io
import re
from collections.abc import Iterable
from typing import Any

from docx import Document as WordDocument
from docx.document import Document as DocxDocument

from .parsing import ParsedBlock, parse_file
from .providers import (
    GenerationProvider,
    TemplateExtractionContext,
    TemplateExtractionResponse,
    get_provider,
)
from .section_tree import build_heading_tree_candidates
from .templates import preflight_docx_template

MAX_BLOCK_TEXT = 1_500
MAX_CHUNK_CHARACTERS = 24_000
MAX_LLM_CHUNKS = 12
MIN_BLOCKS_BY_STAGE = {
    "requirement": 5,
    "feasibility": 8,
    "tender": 8,
    "contract": 5,
}
MIN_SECTIONS_BY_STAGE = {
    "requirement": 2,
    "feasibility": 3,
    "tender": 3,
    "contract": 2,
}


def _candidate_key(value: str, prefix: str, sequence: int) -> str:
    ascii_key = re.sub(r"[^a-z0-9]+", "_", value.lower()).strip("_")
    return ascii_key[:100] or f"{prefix}_{sequence + 1}"


def _block_payload(block: ParsedBlock) -> dict[str, Any]:
    return {
        "id": f"block-{block.sequence}",
        "sequence": block.sequence,
        "kind": block.kind,
        "section_path": block.section_path,
        "text": block.text[:MAX_BLOCK_TEXT],
        "locator": block.locator or {},
    }


def _chunk_blocks(blocks: list[dict[str, Any]]) -> list[list[dict[str, Any]]]:
    chunks: list[list[dict[str, Any]]] = []
    current: list[dict[str, Any]] = []
    current_size = 0
    for block in blocks:
        size = len(str(block.get("text", ""))) + 240
        if current and current_size + size > MAX_CHUNK_CHARACTERS:
            chunks.append(current)
            current = []
            current_size = 0
        current.append(block)
        current_size += size
    if current:
        chunks.append(current)
    return chunks


def _document_structure(content: bytes) -> dict[str, Any]:
    document = WordDocument(io.BytesIO(content))
    styles = sorted({paragraph.style.name for paragraph in document.paragraphs if paragraph.style})
    header_paragraphs = sum(
        len([paragraph for paragraph in section.header.paragraphs if paragraph.text.strip()])
        for section in document.sections
    )
    footer_paragraphs = sum(
        len([paragraph for paragraph in section.footer.paragraphs if paragraph.text.strip()])
        for section in document.sections
    )
    return {
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "section_count": len(document.sections),
        "header_paragraph_count": header_paragraphs,
        "footer_paragraph_count": footer_paragraphs,
        "style_names": styles[:100],
        "layout_preserved": True,
    }


def _merge_responses(
    blocks: list[dict[str, Any]], responses: Iterable[TemplateExtractionResponse]
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[str]]:
    source_text = {str(block["id"]): str(block.get("text", "")) for block in blocks}
    headings: list[dict[str, Any]] = []
    for block in blocks:
        if block.get("kind") != "heading":
            continue
        title = str(block.get("text", "")).strip()
        if not title:
            continue
        locator = block.get("locator") or {}
        level = int(locator.get("heading_level") or 1)
        headings.append(
            {
                "key": _candidate_key(title, "section", len(headings)),
                "title": title,
                "level": level,
                "id": str(block["id"]),
                "source_block_ids": [str(block["id"])],
                "confidence": 1.0,
                "basis": "program_style",
            }
        )
    section_candidates = build_heading_tree_candidates(headings)
    seen_keys = {str(item["key"]) for item in section_candidates}
    seen_under_parent = {
        (item.get("parent_key"), str(item["title"]).casefold()) for item in section_candidates
    }

    variable_candidates: list[dict[str, Any]] = []
    seen_variables: set[str] = set()
    warnings: list[str] = []
    for response in responses:
        for section_suggestion in response.sections:
            title = section_suggestion.title.strip()
            valid_ids = [
                block_id for block_id in section_suggestion.source_block_ids if block_id in source_text
            ]
            if not title or not valid_ids:
                continue
            level = max(1, min(3, int(section_suggestion.level or 1)))
            parent_key = section_suggestion.parent_key
            dedupe = (parent_key, title.casefold())
            if dedupe in seen_under_parent:
                continue
            key = section_suggestion.key or _candidate_key(title, "section", len(section_candidates))
            base_key = key
            suffix = 2
            while key in seen_keys:
                key = f"{base_key}_{suffix}"
                suffix += 1
            seen_keys.add(key)
            seen_under_parent.add(dedupe)
            section_candidates.append(
                {
                    "key": key,
                    "title": title,
                    "level": level,
                    "parent_key": parent_key,
                    "source_block_ids": valid_ids,
                    "confidence": section_suggestion.confidence,
                    "basis": "llm_suggested",
                }
            )
        for variable_suggestion in response.variables:
            valid_ids = [
                block_id for block_id in variable_suggestion.source_block_ids if block_id in source_text
            ]
            if not valid_ids:
                continue
            exact_text = variable_suggestion.exact_text.strip()
            if not exact_text or not any(exact_text in source_text[block_id] for block_id in valid_ids):
                warnings.append(f"已丢弃无法定位原文的变量候选：{variable_suggestion.variable_key}")
                continue
            if variable_suggestion.variable_key in seen_variables:
                continue
            seen_variables.add(variable_suggestion.variable_key)
            variable_candidates.append({**variable_suggestion.model_dump(), "source_block_ids": valid_ids})
        warnings.extend(response.warnings)

    for index, candidate in enumerate(section_candidates):
        candidate["id"] = f"section-{index + 1}"
        candidate["selected"] = True
        candidate.setdefault("level", 1)
        candidate.setdefault("parent_key", None)
    for index, candidate in enumerate(variable_candidates):
        candidate["id"] = f"variable-{index + 1}"
        candidate["selected"] = float(str(candidate["confidence"])) >= 0.75
    return section_candidates, variable_candidates, list(dict.fromkeys(warnings))


def assess_extraction_quality(
    *,
    content: bytes,
    filename: str,
    stage: str,
    sections: list[dict[str, Any]],
    summary: dict[str, Any],
) -> dict[str, Any]:
    """Deterministic gate: unfinished/low-structure files must not become templates."""
    reasons: list[str] = []
    checks: list[dict[str, Any]] = []
    min_blocks = MIN_BLOCKS_BY_STAGE.get(stage, 5)
    min_sections = MIN_SECTIONS_BY_STAGE.get(stage, 2)
    block_count = int(summary.get("block_count") or 0)
    section_count = len(sections)
    program_sections = sum(1 for item in sections if item.get("basis") == "program_style")
    heading_levels = {max(1, min(3, int(item.get("level") or 1))) for item in sections if item.get("title")}

    def _check(key: str, passed: bool, detail: str) -> None:
        checks.append({"key": key, "passed": passed, "detail": detail})
        if not passed:
            reasons.append(detail)

    _check(
        "min_blocks",
        block_count >= min_blocks,
        f"结构化内容块不足（{block_count}/{min_blocks}），不像完整正式文稿",
    )
    _check(
        "min_sections",
        section_count >= min_sections,
        f"可靠章节不足（{section_count}/{min_sections}），无法作为正式模板结构",
    )
    _check(
        "program_headings",
        program_sections >= 1,
        "缺少样式识别的标题结构，版式不足以沉淀为正式模板",
    )
    _check(
        "heading_levels",
        len(heading_levels) >= 1 and max(heading_levels) >= 1,
        "未识别到可用的标题层级",
    )
    _check(
        "layout_preserved",
        bool(summary.get("layout_preserved")),
        "未能保留 DOCX 版式，不能作为正式模板源",
    )

    candidate = build_candidate_template_docx(
        content,
        template_name=f"{filename.rsplit('.', 1)[0]}模板",
        source_filename=filename,
    )
    preflight = preflight_docx_template(candidate)
    _check(
        "candidate_preflight",
        preflight.valid,
        "候选模板预检未通过：" + "；".join(preflight.warnings)
        if preflight.warnings
        else "候选模板预检未通过",
    )

    score = round(sum(1 for item in checks if item["passed"]) / max(len(checks), 1), 3)
    passed = not reasons
    return {
        "passed": passed,
        "score": score,
        "grade": "publishable" if passed else "rejected",
        "min_blocks": min_blocks,
        "min_sections": min_sections,
        "program_section_count": program_sections,
        "checks": checks,
        "reasons": reasons,
        "preflight_placeholders": list(preflight.placeholders),
    }


def analyze_finished_document(
    content: bytes,
    *,
    filename: str,
    stage: str,
    provider: GenerationProvider | None = None,
) -> dict[str, Any]:
    if not filename.lower().endswith(".docx"):
        raise ValueError("当前模板反向提取仅支持 DOCX，以保证版式和样式可保留")
    parsed = parse_file(filename, content)
    if parsed.needs_ocr or not parsed.blocks:
        raise ValueError("文件没有可供模板提取的结构化文本")
    block_payloads = [_block_payload(block) for block in parsed.blocks]
    chunks = _chunk_blocks(block_payloads)
    selected_chunks = chunks[:MAX_LLM_CHUNKS]
    active_provider = provider or get_provider()
    responses = [
        active_provider.extract_template(
            TemplateExtractionContext(stage=stage, filename=filename, blocks=chunk)
        )
        for chunk in selected_chunks
    ]
    sections, variables, warnings = _merge_responses(block_payloads, responses)
    if len(chunks) > MAX_LLM_CHUNKS:
        warnings.append(f"文档内容较长，本次模型分析覆盖前 {MAX_LLM_CHUNKS} 个分块；章节结构仍由程序完整提取")
    if not sections:
        warnings.append("未识别到可靠章节，质量门禁将拒绝该文件进入模板库")
    warnings.append("模型结果仅为候选；确认发布时不会把原文中的金额、主体、日期等实例值写入模板默认值")
    summary = {
        "filename": filename,
        "stage": stage,
        "provider": active_provider.name,
        "model": active_provider.model,
        "block_count": len(block_payloads),
        "llm_chunk_count": len(selected_chunks),
        "truncated_for_llm": len(chunks) > MAX_LLM_CHUNKS,
        **_document_structure(content),
    }
    quality = assess_extraction_quality(
        content=content,
        filename=filename,
        stage=stage,
        sections=sections,
        summary=summary,
    )
    if not quality["passed"]:
        warnings.append("质量门禁未通过：该文件不会进入下方模板列表，需更换合格正式文稿后重试")
    else:
        warnings.append("质量门禁已通过：人工确认后将直接发布为可用于生成的正式模板")
    return {
        "summary": summary,
        "sections": sections,
        "variables": variables,
        "warnings": list(dict.fromkeys(warnings)),
        "quality": quality,
    }


def build_candidate_template_docx(content: bytes, *, template_name: str, source_filename: str) -> bytes:
    document: DocxDocument = WordDocument(io.BytesIO(content))
    body = document._element.body
    for element in list(body):
        if not element.tag.endswith("}sectPr"):
            body.remove(element)

    title_style = "Title" if "Title" in document.styles else None
    title = document.add_paragraph(style=title_style)
    title.add_run("{{DOCUMENT_TITLE}}")
    document.add_paragraph("版本 {{DOCUMENT_VERSION}}")
    document.add_paragraph(
        f"由《{source_filename}》提取形成的候选模板；原文件实例数据已从正文移除，须经人工确认后发布。"
    )
    document.add_page_break()  # type: ignore[no-untyped-call]
    document.add_paragraph("{{DOCUMENT_BODY}}")

    for section in document.sections:
        for paragraph in section.header.paragraphs:
            paragraph.clear()
        section.header.paragraphs[0].add_run(f"{template_name} · 人工确认候选模板")
        for paragraph in section.footer.paragraphs:
            paragraph.clear()
        section.footer.paragraphs[0].add_run("项目文件链式生成平台")

    document.core_properties.title = template_name
    document.core_properties.subject = "从成品 DOCX 提取的候选模板"
    document.core_properties.comments = "原文内容不得作为模板正式值；所有候选变量均需人工确认。"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
