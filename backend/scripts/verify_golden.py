from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from typing import Any

from docx import Document as WordDocument
from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[2]
CASE_DIR = ROOT / "golden_cases" / "demo_001"
ARTIFACT_DIR = ROOT / "artifacts" / "golden_case"
PLACEHOLDER = re.compile(r"\{\{[^{}]+\}\}|\[\[[^\[\]]+\]\]|__+[A-Za-z0-9_]+__+")
FINAL_PROHIBITED = (
    "【待确认】",
    "ai_generated",
    "P0",
    "P1",
    "人工录入",
    "平台内部字段",
    "请回到官方原文核对",
)
TENDER_TABLES_AND_FORMS = (
    "投标人须知前附表",
    "评标办法前附表",
    "资格审查表",
    "符合性审查表",
    "商务、技术和报价评审因素表",
    "合同协议书",
    "采购清单",
    "技术规格和性能指标",
    "投标函",
    "法定代表人身份证明",
    "授权委托书",
    "开标一览表",
    "分项报价表",
    "商务偏离表",
    "技术偏离表",
    "资格证明文件目录",
    "业绩表",
    "项目团队表",
    "服务方案",
    "承诺函",
)
STAGES = {
    "项目建议书_V1.0": ("requirement", "项目建议书"),
    "可行性研究报告_V1.0": ("feasibility", "可行性研究报告"),
    "招标文件_V1.0": ("tender", "招标文件"),
    "合同签约准备版_V1.0": ("contract", "合同"),
}


def _find_soffice() -> str | None:
    configured = os.environ.get("LIBREOFFICE_BIN")
    candidates = [
        configured,
        shutil.which("soffice"),
        shutil.which("libreoffice"),
        "C:/Program Files/LibreOffice/program/soffice.exe",
    ]
    return next((str(value) for value in candidates if value and Path(value).is_file()), None)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _assert(condition: bool, message: str, checks: list[dict[str, str]]) -> None:
    checks.append({"status": "passed" if condition else "failed", "message": message})
    if not condition:
        raise AssertionError(message)


def verify() -> None:
    expected_fields = json.loads((CASE_DIR / "expected_fields.json").read_text(encoding="utf-8"))
    expected_sections = json.loads((CASE_DIR / "expected_sections.json").read_text(encoding="utf-8"))
    manifest = json.loads((ARTIFACT_DIR / "artifact-manifest.json").read_text(encoding="utf-8"))
    validation = json.loads((ARTIFACT_DIR / "validation-report.json").read_text(encoding="utf-8"))
    checks: list[dict[str, str]] = []
    artifact_results: dict[str, Any] = {}
    soffice = _find_soffice()
    _assert(soffice is not None, "LibreOffice 可执行文件可用", checks)

    for stem, (stage, document_title) in STAGES.items():
        docx_path = ARTIFACT_DIR / f"{stem}.docx"
        pdf_path = ARTIFACT_DIR / f"{stem}.pdf"
        _assert(docx_path.is_file() and docx_path.stat().st_size > 0, f"{docx_path.name} 存在且非空", checks)
        _assert(pdf_path.is_file() and pdf_path.stat().st_size > 0, f"{pdf_path.name} 存在且非空", checks)
        with zipfile.ZipFile(docx_path) as archive:
            _assert(archive.testzip() is None, f"{docx_path.name} OOXML ZIP 完整", checks)
            _assert("word/document.xml" in archive.namelist(), f"{docx_path.name} 包含 document.xml", checks)
            document_xml = archive.read("word/document.xml").decode("utf-8")
            footer_xml = "\n".join(
                archive.read(name).decode("utf-8")
                for name in archive.namelist()
                if name.startswith("word/footer") and name.endswith(".xml")
            )
        word = WordDocument(docx_path)
        word_text = "\n".join(p.text for p in word.paragraphs)
        for table in word.tables:
            word_text += "\n" + "\n".join(cell.text for row in table.rows for cell in row.cells)
        _assert(document_title in word_text, f"{docx_path.name} 含文档标题", checks)
        for section_title in expected_sections[stage]:
            _assert(section_title in word_text, f"{docx_path.name} 含章节 {section_title}", checks)
        _assert(not PLACEHOLDER.search(word_text), f"{docx_path.name} 未替换变量为 0", checks)
        for word_section in word.sections:
            _assert(abs(word_section.page_width.cm - 21.0) < 0.05, f"{docx_path.name} 页面宽度为 A4", checks)
            _assert(abs(word_section.page_height.cm - 29.7) < 0.05, f"{docx_path.name} 页面高度为 A4", checks)
        reader = PdfReader(str(pdf_path))
        _assert(len(reader.pages) > 0, f"{pdf_path.name} 页数大于 0", checks)
        pdf_text = "\n".join(page.extract_text() or "" for page in reader.pages)
        _assert(document_title in pdf_text, f"{pdf_path.name} 可提取关键标题", checks)
        _assert(not PLACEHOLDER.search(pdf_text), f"{pdf_path.name} 未替换变量为 0", checks)
        if stage == "tender":
            _assert(len(word.tables) >= 20, "招标 DOCX 包含完整前附表、评审表和可填写投标格式表格", checks)
            _assert('TOC \\o "1-3"' in document_xml, "招标 DOCX 使用可更新目录域", checks)
            _assert("PAGEREF toc_" in document_xml, "招标 DOCX 目录页码使用可更新引用域", checks)
            _assert(
                document_xml.count("w:tblHeader") >= len(word.tables),
                "招标 DOCX 表格均设置跨页重复表头",
                checks,
            )
            _assert(" PAGE " in footer_xml, "招标 DOCX 页脚包含页码域", checks)
            _assert("目录将在" not in word_text + pdf_text, "招标 DOCX/PDF 目录不含更新占位提示", checks)
            _assert("D9EAF7" not in document_xml, "招标 DOCX 不使用平台蓝色表头装饰", checks)
            _assert(not re.search(r"。\s*。|；\s*。", word_text), "招标 DOCX 正文无重复句末标点", checks)
            _assert(not re.search(r"。\s*。|；\s*。", pdf_text), "招标 PDF 正文无重复句末标点", checks)
            for name in TENDER_TABLES_AND_FORMS:
                _assert(name in word_text, f"招标 DOCX 包含 {name}", checks)
                _assert(name in pdf_text, f"招标 PDF 包含 {name}", checks)
            for marker in FINAL_PROHIBITED:
                _assert(marker not in word_text, f"招标 DOCX 不含内部标记 {marker}", checks)
                _assert(marker not in pdf_text, f"招标 PDF 不含内部标记 {marker}", checks)
            for internal_appendix in ("字段来源清单", "一致性校验报告", "风险清单"):
                _assert(internal_appendix not in word_text, f"招标正文未附带 {internal_appendix}", checks)
                _assert(internal_appendix not in pdf_text, f"招标 PDF 未附带 {internal_appendix}", checks)
            for section_title in expected_sections[stage]:
                _assert(
                    pdf_text.count(section_title) >= 2, f"招标 PDF 目录和正文均显示 {section_title}", checks
                )
            _assert(len(reader.pages) >= 12, "招标 PDF 具有正式分页正文", checks)
            for page_number, page in enumerate(reader.pages, 1):
                page_text = page.extract_text() or ""
                body_text = page_text.replace(document_title, "")
                body_text = re.sub(rf"第\s*{page_number}\s*页", "", body_text)
                body_text = re.sub(r"\s+", "", body_text)
                _assert(len(body_text) >= 10, f"招标 PDF 第 {page_number} 页不是空白页", checks)
        for page in reader.pages:
            width = float(page.mediabox.width)
            height = float(page.mediabox.height)
            _assert(
                abs(width - 595.28) < 2 and abs(height - 841.89) < 2, f"{pdf_path.name} 页面为 A4", checks
            )
        _assert(
            _sha256(docx_path) == manifest["artifacts"][docx_path.name]["sha256"],
            f"{docx_path.name} SHA-256 匹配",
            checks,
        )
        _assert(
            _sha256(pdf_path) == manifest["artifacts"][pdf_path.name]["sha256"],
            f"{pdf_path.name} SHA-256 匹配",
            checks,
        )
        _assert(
            validation["validations"][stage]["status"] == "passed",
            f"{document_title} 文档校验通过",
            checks,
        )
        artifact_results[docx_path.name] = {"sha256": _sha256(docx_path), "paragraphs": len(word.paragraphs)}
        artifact_results[pdf_path.name] = {"sha256": _sha256(pdf_path), "pages": len(reader.pages)}

    with tempfile.TemporaryDirectory(prefix="docchain-lo-") as temp_dir:
        for stem in STAGES:
            docx_path = ARTIFACT_DIR / f"{stem}.docx"
            result = subprocess.run(  # noqa: S603 - executable is resolved from trusted local paths.
                [str(soffice), "--headless", "--convert-to", "pdf", "--outdir", temp_dir, str(docx_path)],
                check=False,
                capture_output=True,
                text=True,
                timeout=120,
            )
            converted = Path(temp_dir) / f"{docx_path.stem}.pdf"
            _assert(
                result.returncode == 0 and converted.is_file(),
                f"{docx_path.name} 可由 LibreOffice 打开并转换",
                checks,
            )

    pre_draft_path = ARTIFACT_DIR / "合同预草案_V0.1.docx"
    _assert(pre_draft_path.is_file() and pre_draft_path.stat().st_size > 0, "合同预草案存在且非空", checks)
    with zipfile.ZipFile(pre_draft_path) as archive:
        _assert(archive.testzip() is None, "合同预草案 OOXML ZIP 完整", checks)
    _assert(
        _sha256(pre_draft_path) == manifest["artifacts"][pre_draft_path.name]["sha256"],
        "合同预草案 SHA-256 匹配",
        checks,
    )
    artifact_results[pre_draft_path.name] = {"sha256": _sha256(pre_draft_path)}

    xlsx_path = ARTIFACT_DIR / "field-traceability.xlsx"
    _assert(
        xlsx_path.is_file() and xlsx_path.stat().st_size > 0,
        "field-traceability.xlsx 存在且非空",
        checks,
    )
    workbook = load_workbook(xlsx_path, data_only=False)
    _assert(workbook.sheetnames == ["字段来源追溯", "合同付款校验"], "追溯工作簿工作表完整", checks)
    trace = workbook["字段来源追溯"]
    rows = list(trace.iter_rows(min_row=2, values_only=True))
    indexed = {(row[0], row[1]): row[3] for row in rows}
    for stage, fields in expected_fields.items():
        for key, value in fields.items():
            actual = indexed[(stage, key)]
            if isinstance(value, (dict, list)):
                actual = json.loads(actual)
            _assert(actual == value, f"追溯字段 {stage}.{key} 与 expected_fields 一致", checks)
    payment_sheet = workbook["合同付款校验"]
    ratio_total = sum(float(payment_sheet.cell(row=row, column=2).value) for row in range(2, 5))
    amount_total = sum(float(payment_sheet.cell(row=row, column=3).value) for row in range(2, 5))
    _assert(ratio_total == 100, "合同付款比例合计 100%", checks)
    _assert(
        amount_total == expected_fields["contract"]["final_contract_amount"],
        "合同付款金额合计等于最终合同金额",
        checks,
    )
    _assert(
        expected_fields["feasibility"]["total_investment"] != expected_fields["tender"]["procurement_budget"],
        "可研总投资未映射为招标预算",
        checks,
    )
    _assert(
        expected_fields["tender"]["maximum_price"] != expected_fields["contract"]["final_contract_amount"],
        "最高限价未映射为最终合同金额",
        checks,
    )
    _assert(expected_fields["requirement"]["project_period"] != 12, "项目总周期未映射为合同履行期限", checks)
    _assert(
        _sha256(xlsx_path) == manifest["artifacts"][xlsx_path.name]["sha256"],
        "追溯 XLSX SHA-256 匹配",
        checks,
    )
    artifact_results[xlsx_path.name] = {"sha256": _sha256(xlsx_path), "sheets": workbook.sheetnames}

    report = {
        "case": "demo_001",
        "status": "passed",
        "profile": "demo_enterprise_report_cn",
        "strict_customer_template_compliance": False,
        "checks": checks,
        "artifacts": artifact_results,
        "external_boundary": (
            "Demo 通用模板已验证；客户正式模板和授权字体尚未提供，不能声明客户模板严格合规。"
        ),
    }
    (ARTIFACT_DIR / "format-compliance-report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8"
    )
    print(f"Golden Case verified: {len(checks)} checks passed")


if __name__ == "__main__":
    verify()
