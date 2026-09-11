from __future__ import annotations

import io
import zipfile
from datetime import UTC, datetime

from docx import Document as WordDocument
from sqlalchemy import func, select

from backend.app.config import get_settings
from backend.app.db import SessionLocal
from backend.app.field_catalog import BASE_FIELD_DEFINITIONS, default_rules_for_field, merge_missing_aliases
from backend.app.models import (
    DocumentFormatProfile,
    FieldDefinition,
    Organization,
    Permission,
    ProcurementRuleSet,
    ProjectCodeRule,
    ProjectType,
    Role,
    RolePermission,
    Template,
    TemplateSection,
    TemplateVariable,
    TemplateVersion,
    User,
    UserRole,
)
from backend.app.security import hash_password, verify_password
from backend.app.services.feasibility_rules import BUILTIN_FEASIBILITY_RULES, FEASIBILITY_RULES_VERSION
from backend.app.services.providers import SECTION_PLANS
from backend.app.services.section_tree import section_nodes_from_pairs
from backend.app.services.storage import get_storage, sha256_bytes
from backend.app.services.templates import (
    DEMO_TEMPLATE_CONTENT_TYPE,
    build_demo_template,
    build_template_document,
)
from backend.app.template_catalog import (
    BUILTIN_TEMPLATE_CATALOG,
    PLATFORM_REFERENCE_TEMPLATE,
    TEMPLATE_SOURCE_LABELS,
    section_plan_items_for,
)

PERMISSIONS = {
    "system.admin": "系统管理",
    "project.create": "创建项目",
    "project.read": "查看项目",
    "project.write": "编辑项目",
    "template.manage": "管理模板",
    "template.publish": "发布模板",
    "document.edit": "编辑文档",
    "document.finalize": "定稿文档",
    "export.create": "导出文档",
    "audit.read": "查看审计日志",
    "user.manage": "管理用户",
}
ROLE_PERMISSIONS = {
    "system_admin": set(PERMISSIONS),
    "template_admin": {"project.read", "template.manage", "template.publish", "export.create"},
    "project_editor": {
        "project.create",
        "project.read",
        "project.write",
        "document.edit",
        "export.create",
    },
    "reviewer": {"project.read", "document.edit", "document.finalize", "export.create"},
    "viewer": {"project.read"},
}
ROLE_NAMES = {
    "system_admin": "系统管理员",
    "template_admin": "模板管理员",
    "project_editor": "项目编制人员",
    "reviewer": "审核人员",
    "viewer": "只读人员",
}
DEMO_TEMPLATES = {
    "requirement": "政府投资项目建议书通用参考模板",
    "feasibility": "一般投资项目可行性研究报告参考模板",
    "tender": "通用采购招标文件参考模板",
    "contract": "通用采购合同参考模板",
}
FIELD_DEFINITIONS = BASE_FIELD_DEFINITIONS
DEFAULT_PROJECT_TYPES = (
    ("government_investment", "政府投资项目", 10),
    ("enterprise_investment", "企业投资项目", 20),
)
DEFAULT_PROJECT_CODE_RULE = {
    "pattern": "{project_type}_{date}_{seq}",
    "date_format": "YYYYMMDD",
    "seq_width": 4,
    "reset_scope": "type_day",
}


def _is_legacy_demo_source(storage_key: str | None) -> bool:
    if not storage_key:
        return False
    storage = get_storage()
    if not storage.exists(storage_key):
        return False
    try:
        document = WordDocument(io.BytesIO(storage.get(storage_key)))
        text = [paragraph.text for paragraph in document.paragraphs]
        for section in document.sections:
            text.extend(paragraph.text for paragraph in section.header.paragraphs)
            text.extend(paragraph.text for paragraph in section.footer.paragraphs)
    except (OSError, ValueError, zipfile.BadZipFile):
        return False
    return any("DEMO 通用模板" in paragraph for paragraph in text)


def seed() -> None:
    settings = get_settings()
    with SessionLocal() as db:
        organization = db.scalar(
            select(Organization).where(Organization.name == settings.demo_organization_name)
        )
        if organization is None:
            organization = Organization(name=settings.demo_organization_name)
            db.add(organization)
            db.flush()
        if not organization.brand_name:
            organization.brand_name = "智能招标管理"
        if not organization.brand_subtitle:
            organization.brand_subtitle = "受控生成与定稿平台"
        if not organization.brand_mark:
            organization.brand_mark = "智"

        permission_rows: dict[str, Permission] = {}
        for code, name in PERMISSIONS.items():
            permission = db.scalar(select(Permission).where(Permission.code == code))
            if permission is None:
                permission = Permission(code=code, name=name)
                db.add(permission)
                db.flush()
            permission_rows[code] = permission

        roles: dict[str, Role] = {}
        for key, name in ROLE_NAMES.items():
            role = db.scalar(select(Role).where(Role.organization_id == organization.id, Role.key == key))
            if role is None:
                role = Role(
                    organization_id=organization.id,
                    key=key,
                    name=name,
                    is_system=True,
                )
                db.add(role)
                db.flush()
            roles[key] = role
            for code in ROLE_PERMISSIONS[key]:
                exists = db.scalar(
                    select(RolePermission).where(
                        RolePermission.role_id == role.id,
                        RolePermission.permission_id == permission_rows[code].id,
                    )
                )
                if exists is None:
                    db.add(
                        RolePermission(
                            organization_id=organization.id,
                            role_id=role.id,
                            permission_id=permission_rows[code].id,
                        )
                    )

        admin_account = settings.demo_admin_account.strip().lower()
        admin = db.scalar(select(User).where(func.lower(User.email) == admin_account))
        if admin is None and admin_account != "admin@example.com":
            admin = db.scalar(select(User).where(func.lower(User.email) == "admin@example.com"))
        if admin is None:
            admin = User(
                organization_id=organization.id,
                email=admin_account,
                display_name="Demo 管理员",
                password_hash=hash_password(settings.demo_admin_password, minimum_length=8),
            )
            db.add(admin)
            db.flush()
        else:
            credentials_changed = False
            if admin.email != admin_account:
                admin.email = admin_account
                credentials_changed = True
            if not verify_password(settings.demo_admin_password, admin.password_hash):
                admin.password_hash = hash_password(settings.demo_admin_password, minimum_length=8)
                credentials_changed = True
            if credentials_changed:
                admin.session_version += 1
        if (
            db.scalar(
                select(UserRole).where(
                    UserRole.user_id == admin.id, UserRole.role_id == roles["system_admin"].id
                )
            )
            is None
        ):
            db.add(
                UserRole(
                    organization_id=organization.id,
                    user_id=admin.id,
                    role_id=roles["system_admin"].id,
                )
            )

        for code, name, sort_order in DEFAULT_PROJECT_TYPES:
            project_type = db.scalar(
                select(ProjectType).where(
                    ProjectType.organization_id == organization.id,
                    ProjectType.code == code,
                )
            )
            if project_type is None:
                db.add(
                    ProjectType(
                        organization_id=organization.id,
                        code=code,
                        name=name,
                        sort_order=sort_order,
                        is_active=True,
                        is_system=True,
                        created_by=admin.id,
                        updated_by=admin.id,
                    )
                )
            else:
                if not project_type.name:
                    project_type.name = name
                if project_type.sort_order is None:
                    project_type.sort_order = sort_order
                project_type.is_system = True

        code_rule = db.scalar(
            select(ProjectCodeRule).where(ProjectCodeRule.organization_id == organization.id)
        )
        if code_rule is None:
            db.add(
                ProjectCodeRule(
                    organization_id=organization.id,
                    pattern=DEFAULT_PROJECT_CODE_RULE["pattern"],
                    date_format=DEFAULT_PROJECT_CODE_RULE["date_format"],
                    seq_width=DEFAULT_PROJECT_CODE_RULE["seq_width"],
                    reset_scope=DEFAULT_PROJECT_CODE_RULE["reset_scope"],
                    created_by=admin.id,
                    updated_by=admin.id,
                )
            )

        for spec in BUILTIN_TEMPLATE_CATALOG:
            template = db.scalar(
                select(Template).where(
                    Template.organization_id == organization.id,
                    Template.name == spec.name,
                )
            )
            if template is None and spec.legacy_name:
                template = db.scalar(
                    select(Template).where(
                        Template.organization_id == organization.id,
                        Template.name == spec.legacy_name,
                    )
                )
            if template is None:
                template = Template(
                    organization_id=organization.id,
                    name=spec.name,
                    stage=spec.stage,
                    source_kind=spec.source_kind,
                    scope="organization",
                    status="published",
                    current_version=1,
                )
                db.add(template)
                db.flush()
            template.name = spec.name
            template.stage = spec.stage
            template.source_kind = spec.source_kind
            template.issuing_authority = spec.issuing_authority
            template.document_number = spec.document_number
            template.publish_year = spec.publish_year
            template.source_url = spec.source_url
            template.applicability = spec.applicability
            template.procurement_type = spec.procurement_type
            template.contract_type = spec.contract_type
            template.is_builtin = True
            template.generation_enabled = spec.generation_enabled
            template.status = "published"
            target_version = template.current_version
            current_template_version = db.scalar(
                select(TemplateVersion).where(
                    TemplateVersion.template_id == template.id,
                    TemplateVersion.version == target_version,
                )
            )
            if (
                spec.source_kind == PLATFORM_REFERENCE_TEMPLATE
                and target_version == 1
                and current_template_version is not None
                and _is_legacy_demo_source(current_template_version.storage_key)
            ):
                current_template_version.status = "superseded"
                target_version = 2
            template.current_version = target_version
            template_version = db.scalar(
                select(TemplateVersion).where(
                    TemplateVersion.template_id == template.id,
                    TemplateVersion.version == target_version,
                )
            )
            if template_version is None:
                template_version = TemplateVersion(
                    organization_id=organization.id,
                    template_id=template.id,
                    version=target_version,
                    status="published",
                    published_at=datetime.now(UTC),
                    format_profile={
                        "standard": "demo_enterprise_report_cn",
                        "page_size": "A4",
                        "cjk_font": "Noto Sans CJK SC or customer-authorized font",
                        "strict_compliance": False,
                    },
                )
                db.add(template_version)
                db.flush()
            template_version.status = "published"
            if template_version.published_at is None:
                template_version.published_at = datetime.now(UTC)
            if spec.generation_enabled:
                if spec.source_kind == PLATFORM_REFERENCE_TEMPLATE:
                    template_content = build_demo_template(spec.stage, spec.name)
                else:
                    template_content = build_template_document(
                        spec.stage,
                        spec.name,
                        source_label=TEMPLATE_SOURCE_LABELS[spec.source_kind],
                        disclaimer=spec.applicability,
                    )
                template_key = (
                    f"{organization.id}/templates/{template.id}/"
                    f"v{target_version}/template.docx"
                )
                storage = get_storage()
                if not storage.exists(template_key):
                    storage.put(template_key, template_content, DEMO_TEMPLATE_CONTENT_TYPE)
                stored_template_content = storage.get(template_key)
                template_version.storage_key = template_key
                template_version.sha256 = sha256_bytes(stored_template_content)
            section_items = section_plan_items_for(
                spec, section_nodes_from_pairs(SECTION_PLANS[spec.stage])
            )
            key_to_section: dict[str, TemplateSection] = {}
            for item in section_items:
                section = db.scalar(
                    select(TemplateSection).where(
                        TemplateSection.template_version_id == template_version.id,
                        TemplateSection.key == item.key,
                    )
                )
                parent = key_to_section.get(item.parent_key) if item.parent_key else None
                if section is None:
                    section = TemplateSection(
                        organization_id=organization.id,
                        template_version_id=template_version.id,
                        sequence=item.sequence,
                        key=item.key,
                        title=item.title,
                        parent_id=parent.id if parent else None,
                        level=item.level,
                        section_type="editable",
                        required=True,
                        content=None,
                    )
                    db.add(section)
                    db.flush()
                else:
                    section.sequence = item.sequence
                    section.title = item.title
                    section.parent_id = parent.id if parent else None
                    section.level = item.level
                key_to_section[item.key] = section
            from backend.app.tender_field_profiles import profile_fields, resolve_profile_name

            allowed_variable_keys: set[str] | None = None
            if spec.stage == "tender":
                profile_name = resolve_profile_name(
                    procurement_category=spec.procurement_type,
                    group_name=spec.name,
                )
                allowed_variable_keys = {item.field_key for item in profile_fields(profile_name)}
            for key, _label, _data_type, _unit, _criticality, required in FIELD_DEFINITIONS[
                spec.stage
            ]:
                # Avoid registering the entire tender catalog on every template; that made
                # applicable-field resolution treat every file as needing all ~37 fields.
                if allowed_variable_keys is not None and key not in allowed_variable_keys:
                    continue
                variable = db.scalar(
                    select(TemplateVariable).where(
                        TemplateVariable.template_version_id == template_version.id,
                        TemplateVariable.variable_key == key,
                    )
                )
                if variable is None:
                    db.add(
                        TemplateVariable(
                            organization_id=organization.id,
                            template_version_id=template_version.id,
                            variable_key=key,
                            field_key=key,
                            required=required,
                        )
                    )

        profile = db.scalar(
            select(DocumentFormatProfile).where(
                DocumentFormatProfile.organization_id == organization.id,
                DocumentFormatProfile.key == "demo_enterprise_report_cn",
                DocumentFormatProfile.version == 1,
            )
        )
        if profile is None:
            db.add(
                DocumentFormatProfile(
                    organization_id=organization.id,
                    key="demo_enterprise_report_cn",
                    name="Demo 中文企业报告格式",
                    version=1,
                    status="published",
                    rules={
                        "page_size": "A4",
                        "margins_cm": {"top": 2.6, "bottom": 2.4, "left": 2.8, "right": 2.6},
                        "body_font_role": "宋体兼容",
                        "heading_font_role": "黑体兼容",
                    },
                    required_fonts=["Noto Sans CJK SC or customer-authorized font"],
                    strict_compliance=False,
                )
            )

        for stage, definitions in FIELD_DEFINITIONS.items():
            for key, label, data_type, unit, criticality, required in definitions:
                exists = db.scalar(
                    select(FieldDefinition).where(
                        FieldDefinition.organization_id == organization.id,
                        FieldDefinition.stage == stage,
                        FieldDefinition.field_key == key,
                    )
                )
                if exists is None:
                    db.add(
                        FieldDefinition(
                            organization_id=organization.id,
                            stage=stage,
                            field_key=key,
                            field_label=label,
                            data_type=data_type,
                            unit=unit,
                            criticality=criticality,
                            required=required,
                            is_base=True,
                            is_active=True,
                            rules=default_rules_for_field(key, label),
                        )
                    )
                else:
                    exists.is_base = True
                    exists.rules = merge_missing_aliases(exists.rules, key, exists.field_label or label)

        builtin_key = "feasibility-procurement-core"
        existing_rule = db.scalar(
            select(ProcurementRuleSet).where(
                ProcurementRuleSet.organization_id == organization.id,
                ProcurementRuleSet.key == builtin_key,
                ProcurementRuleSet.status == "published",
            )
        )
        admin_user = db.scalar(select(User).where(User.organization_id == organization.id).limit(1))
        actor_id = admin_user.id if admin_user else None
        if existing_rule is None:
            db.add(
                ProcurementRuleSet(
                    organization_id=organization.id,
                    key=builtin_key,
                    name="可研采购划分核心规则",
                    version=1,
                    status="published",
                    applicable_subject=None,
                    region=None,
                    funding_nature=None,
                    source_name="平台内置 feasibility-procurement-v1",
                    source_url=None,
                    effective_date=None,
                    expiry_date=None,
                    rules_json=BUILTIN_FEASIBILITY_RULES,
                    created_by=actor_id,
                    updated_by=actor_id,
                )
            )
        elif (existing_rule.rules_json or {}).get("version") != FEASIBILITY_RULES_VERSION:
            existing_rule.rules_json = BUILTIN_FEASIBILITY_RULES
            existing_rule.revision += 1
            existing_rule.updated_by = actor_id

        db.commit()
        print(f"Seed complete: {admin_account} and builtin template catalog are ready")


if __name__ == "__main__":
    seed()
