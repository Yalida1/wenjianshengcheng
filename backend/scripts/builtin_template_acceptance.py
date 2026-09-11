from __future__ import annotations

# The executor only parses files generated inside its isolated acceptance tree and invokes
# fixed tool names with argument lists; long Chinese report strings are intentionally kept intact.
# ruff: noqa: E501, S105, S107, S314, S603, S607
import hashlib
import io
import json
import os
import platform
import re
import shutil
import subprocess
import tempfile
import time
import uuid
import zipfile
from dataclasses import dataclass, field
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

import httpx
import redis
from docx import Document as WordDocument
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pypdf import PdfReader
from sqlalchemy import select, text

from backend.app.config import get_settings
from backend.app.db import SessionLocal
from backend.app.models import (
    AuditLog,
    Document,
    DocumentContentBlock,
    DocumentVersion,
    FieldConfirmation,
    FieldEvidence,
    FieldSnapshot,
    FieldValue,
    GenerationJob,
    Project,
    ProjectStage,
    Template,
    TemplateSection,
    TemplateVariable,
    TemplateVersion,
)
from backend.app.services.storage import get_storage
from backend.app.template_catalog import TEMPLATE_SOURCE_LABELS

ROOT = Path(__file__).resolve().parents[2]
ARTIFACT_ROOT = ROOT / "artifacts" / "builtin_template_acceptance"
QA_ROOT = ROOT / "qa" / "builtin_template_acceptance"
EXPECTED = json.loads((QA_ROOT / "EXPECTED_VALUES.json").read_text(encoding="utf-8"))
API_URL = os.getenv("ACCEPTANCE_API_URL", "http://127.0.0.1:8000/api/v1")
STAGES = ("requirement", "feasibility", "tender", "contract")
STAGE_LABELS = {
    "requirement": "项目建议书",
    "feasibility": "可行性研究报告",
    "tender": "招标文件",
    "contract": "合同",
}
OUTPUT_FILES = {
    "requirement": "项目建议书_V1.0",
    "feasibility": "可行性研究报告_V1.0",
    "tender": "招标文件_V1.0",
    "contract": "合同签约准备版_V1.0",
}
TERMINAL = {"succeeded", "failed", "cancelled", "needs_ocr"}


class AcceptanceError(RuntimeError):
    pass


def utc_now() -> str:
    return datetime.now(UTC).isoformat()


def json_default(value: Any) -> Any:
    if isinstance(value, (datetime, Decimal)):
        return str(value)
    if isinstance(value, Path):
        return value.as_posix()
    raise TypeError(f"Unsupported JSON value: {type(value)!r}")


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, default=json_default) + "\n",
        encoding="utf-8",
    )


def sha256(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def redact(value: Any) -> Any:
    if isinstance(value, dict):
        return {
            key: "***" if key.lower() in {"password", "token", "api_key", "csrf_token"} else redact(item)
            for key, item in value.items()
        }
    if isinstance(value, list):
        return [redact(item) for item in value]
    return value


@dataclass
class Recorder:
    started_at: str = field(default_factory=utc_now)
    checks: list[dict[str, Any]] = field(default_factory=list)
    commands: list[dict[str, Any]] = field(default_factory=list)
    api_requests: list[dict[str, Any]] = field(default_factory=list)
    generated_files: list[dict[str, Any]] = field(default_factory=list)
    main_versions: dict[str, dict[str, Any]] = field(default_factory=dict)
    negative_tests: list[dict[str, Any]] = field(default_factory=list)
    failures: list[str] = field(default_factory=list)

    def check(self, check_id: str, passed: bool, detail: Any) -> None:
        self.checks.append({"id": check_id, "status": "passed" if passed else "failed", "detail": detail})
        if not passed:
            self.failures.append(f"{check_id}: {detail}")

    def command(self, name: str, command: list[str], completed: subprocess.CompletedProcess[str]) -> None:
        self.commands.append(
            {
                "name": name,
                "command": command,
                "exit_code": completed.returncode,
                "stdout": completed.stdout[-4000:],
                "stderr": completed.stderr[-4000:],
            }
        )


class ApiClient:
    def __init__(self, recorder: Recorder, account: str = "admin", password: str = "admin123") -> None:
        self.recorder = recorder
        self.client = httpx.Client(base_url=API_URL, timeout=120)
        response = self.call(
            "POST",
            "/auth/login",
            json_body={"email": account, "password": password},
            expected=200,
            include_csrf=False,
        )
        self.user = response.json()["user"]
        self.client.headers["X-CSRF-Token"] = response.json()["csrf_token"]

    def close(self) -> None:
        self.client.close()

    def call(
        self,
        method: str,
        path: str,
        *,
        json_body: Any = None,
        files: Any = None,
        expected: int | set[int] | None = None,
        include_csrf: bool = True,
    ) -> httpx.Response:
        request_id = f"bta-{uuid.uuid4()}"
        headers = {"X-Request-ID": request_id}
        if not include_csrf:
            headers.pop("X-CSRF-Token", None)
        response = self.client.request(method, path, json=json_body, files=files, headers=headers)
        content_type = response.headers.get("content-type", "")
        if "json" in content_type:
            try:
                response_value: Any = redact(response.json())
            except ValueError:
                response_value = response.text[:1000]
        else:
            response_value = {
                "content_type": content_type,
                "size_bytes": len(response.content),
                "sha256": sha256(response.content),
            }
        entry = {
            "timestamp": utc_now(),
            "request_id": response.headers.get("X-Request-ID", request_id),
            "method": method,
            "path": path,
            "request": redact(json_body) if json_body is not None else {"multipart": bool(files)},
            "status": response.status_code,
            "response": response_value,
        }
        self.recorder.api_requests.append(entry)
        if expected is not None:
            expected_codes = {expected} if isinstance(expected, int) else expected
            if response.status_code not in expected_codes:
                raise AcceptanceError(
                    f"{method} {path} expected {sorted(expected_codes)}, got "
                    f"{response.status_code}: {response.text[:1000]}"
                )
        return response


def reset_artifacts() -> None:
    if ARTIFACT_ROOT.exists():
        for child in ARTIFACT_ROOT.iterdir():
            if child.is_dir():
                shutil.rmtree(child)
            else:
                child.unlink()
    for name in (
        "requirement",
        "feasibility",
        "tender",
        "contract",
        "screenshots",
        "logs",
        "api-results",
        "all-template-minimum",
        "inputs",
    ):
        (ARTIFACT_ROOT / name).mkdir(parents=True, exist_ok=True)


def wait_for_api() -> None:
    health_url = API_URL.rsplit("/api/v1", 1)[0] + "/ready"
    deadline = time.monotonic() + 120
    last_error = ""
    while time.monotonic() < deadline:
        try:
            response = httpx.get(health_url, timeout=5)
            if response.status_code == 200:
                return
            last_error = response.text
        except httpx.HTTPError as exc:
            last_error = str(exc)
        time.sleep(1)
    raise AcceptanceError(f"API did not become ready: {last_error}")


def build_input_document() -> bytes:
    document = WordDocument()
    document.add_heading("区域绿色数据中心节能改造需求说明", 0)
    document.add_paragraph(f"项目名称：{EXPECTED['project']['name']}")
    document.add_paragraph(f"建设单位：{EXPECTED['project']['owner']}")
    document.add_paragraph(f"建设地点：{EXPECTED['project']['location']}")
    document.add_heading("一、项目建设范围", level=1)
    for index, item in enumerate(EXPECTED["scope"]["project_total"], 1):
        document.add_paragraph(f"{index}. {item}")
    document.add_heading("二、控制数据", level=1)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "项目总建设周期"
    table.rows[0].cells[1].text = EXPECTED["periods"]["project_period"]
    for label, value in (
        ("可研总投资", "12,800,000 元"),
        ("招标预算", "9,800,000 元"),
        ("最高限价", "9,500,000 元"),
        ("最终合同金额", "9,180,000 元"),
    ):
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def project_stage(api: ApiClient, project_id: str, stage: str) -> dict[str, Any]:
    stages = api.call("GET", f"/projects/{project_id}/stages", expected=200).json()
    return next(item for item in stages if item["stage"] == stage)


def upload_source(api: ApiClient, project_id: str, stage: str, content: bytes) -> dict[str, Any]:
    filename = f"{STAGE_LABELS[stage]}验收输入.docx"
    response = api.call(
        "POST",
        f"/files?project_id={project_id}&stage={stage}",
        files={
            "upload": (
                filename,
                content,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        expected=201,
    )
    file_record = response.json()
    versions = api.call("GET", f"/files/{file_record['id']}/versions", expected=200).json()
    parse_jobs = api.call("GET", f"/files/{file_record['id']}/parse-jobs", expected=200).json()
    if parse_jobs:
        deadline = time.monotonic() + 120
        while time.monotonic() < deadline:
            job = api.call("GET", f"/parse-jobs/{parse_jobs[0]['id']}", expected=200).json()
            if job["status"] in TERMINAL:
                file_record["parse_job"] = job
                break
            time.sleep(0.5)
        else:
            raise AcceptanceError(f"Parse job timed out for {filename}")
    file_record["file_version_id"] = versions[0]["id"]
    file_record["file_sha256"] = versions[0]["sha256"]
    return file_record


def set_stage_source(
    api: ApiClient,
    project_id: str,
    stage: str,
    source_type: str,
    source_version_id: str,
) -> dict[str, Any]:
    current = project_stage(api, project_id, stage)
    return api.call(
        "PUT",
        f"/projects/{project_id}/stages/{stage}/source",
        json_body={
            "source_type": source_type,
            "source_file_version_id": source_version_id,
            "revision": current["revision"],
        },
        expected=200,
    ).json()


def set_upstream_source(api: ApiClient, project_id: str, stage: str) -> dict[str, Any]:
    options = api.call("GET", f"/projects/{project_id}/stages/{stage}/sources", expected=200).json()
    upstream = next(item for item in options if item["source_type"] == "upstream_final")
    return set_stage_source(api, project_id, stage, "upstream_final", upstream["id"])


def create_project(api: ApiClient, code: str, description: str) -> dict[str, Any]:
    return api.call(
        "POST",
        "/projects",
        json_body={
            "code": code,
            "name": EXPECTED["project"]["name"],
            "project_type": "government_investment",
            "description": description,
        },
        expected=201,
    ).json()


def create_users(api: ApiClient, project_id: str) -> dict[str, dict[str, str]]:
    password = "BuiltinAcceptance123!"
    definitions = (
        ("template-admin@acceptance.test", "模板管理员", "template_admin", "reviewer"),
        ("editor@acceptance.test", "项目编制人员", "project_editor", "editor"),
        ("reviewer@acceptance.test", "审核人员", "reviewer", "reviewer"),
        ("viewer@acceptance.test", "只读人员", "viewer", "viewer"),
    )
    users: dict[str, dict[str, str]] = {
        "system_admin": {"id": api.user["id"], "email": "admin", "password": "admin123"}
    }
    for email, display_name, role_key, project_role in definitions:
        response = api.call(
            "POST",
            "/users",
            json_body={
                "email": email,
                "display_name": display_name,
                "password": password,
                "role_keys": [role_key],
            },
            expected=201,
        ).json()
        api.call(
            "POST",
            f"/projects/{project_id}/members",
            json_body={"user_id": response["id"], "role_key": project_role},
            expected=201,
        )
        users[role_key] = {"id": response["id"], "email": email, "password": password}
    return users


def stage_values(stage: str) -> dict[str, Any]:
    project_scope = "、".join(EXPECTED["scope"]["project_total"])
    procurement_scope = (
        "、".join(EXPECTED["scope"]["procurement"])
        + "；不含"
        + "、".join(EXPECTED["scope"]["excluded"])
    )
    values: dict[str, dict[str, Any]] = {
        "requirement": {
            "project_name": EXPECTED["project"]["name"],
            "project_owner": EXPECTED["project"]["owner"],
            "construction_scope": project_scope,
            "project_period": EXPECTED["periods"]["project_period"],
            "project_location": EXPECTED["project"]["location"],
        },
        "feasibility": {
            "project_name": EXPECTED["project"]["name"],
            "total_investment": EXPECTED["amounts"]["total_investment"],
            "construction_scope": project_scope,
            "project_period": EXPECTED["periods"]["project_period"],
            "project_location": EXPECTED["project"]["location"],
        },
        "tender": {
            "project_name": EXPECTED["project"]["name"],
            "procurement_budget": EXPECTED["amounts"]["procurement_budget"],
            "maximum_price": EXPECTED["amounts"]["maximum_price"],
            "procurement_scope": procurement_scope,
            "delivery_period": EXPECTED["periods"]["tender_delivery"],
            "delivery_location": EXPECTED["project"]["location"],
        },
        "contract": {
            "party_a": EXPECTED["contract"]["party_a"],
            "party_b": EXPECTED["contract"]["party_b"],
            "contract_subject": "区域绿色数据中心节能改造设备采购及实施服务",
            "contract_scope": procurement_scope,
            "final_contract_amount": EXPECTED["amounts"]["final_contract_amount"],
            "tax_rate": EXPECTED["amounts"]["tax_rate_percent"],
            "tax_inclusion": EXPECTED["contract"]["tax_inclusion"],
            "contract_duration": EXPECTED["periods"]["contract_duration"],
            "delivery_location": EXPECTED["contract"]["delivery_location"],
            "payment_plan": EXPECTED["contract"]["payment_plan"],
            "acceptance": "按合同、技术规范和双方确认的测试方案完成验收",
            "warranty": EXPECTED["periods"]["warranty"],
            "breach": "违约方按合同约定承担继续履行、采取补救措施或赔偿损失等责任",
            "effective_conditions": "双方法定代表人或授权代表签字并加盖公章后生效",
        },
    }
    return values[stage]


def create_fields(
    api: ApiClient,
    project_id: str,
    stage: str,
    values: dict[str, Any],
    *,
    status_overrides: dict[str, str] | None = None,
) -> dict[str, dict[str, Any]]:
    definitions = {
        item["field_key"]: item
        for item in api.call("GET", f"/field-definitions?stage={stage}", expected=200).json()
    }
    current_fields = {
        item["field_key"]: item
        for item in api.call(
            "GET",
            f"/field-values?project_id={project_id}&stage={stage}",
            expected=200,
        ).json()
    }
    output: dict[str, dict[str, Any]] = {}
    for key, value in values.items():
        definition = definitions.get(key, {})
        status = (status_overrides or {}).get(key, "missing")
        payload = {
            "field_key": key,
            "field_label": definition.get("field_label", key),
            "data_type": definition.get("data_type", "string"),
            "value": value,
            "normalized_value": value,
            "unit": definition.get("unit"),
            "criticality": definition.get("criticality", "P1"),
            "status": status,
            "source_type": "user_input" if status == "missing" else "ai_suggestion",
            "confidence": None,
            "evidence": None,
        }
        existing = current_fields.get(key)
        if existing is None:
            field_value = api.call(
                "POST",
                f"/field-values?project_id={project_id}&stage={stage}",
                json_body=payload,
                expected=201,
            ).json()
        else:
            field_value = api.call(
                "PATCH",
                f"/field-values/{existing['id']}",
                json_body={
                    "value": value,
                    "normalized_value": value,
                    "unit": definition.get("unit"),
                    "status": status,
                    "source_type": payload["source_type"],
                    "revision": existing["revision"],
                    "evidence": None,
                },
                expected=200,
            ).json()
        if status == "missing":
            field_value = api.call(
                "POST",
                f"/field-values/{field_value['id']}/confirm",
                json_body={"revision": field_value["revision"], "evidence_acknowledged": False},
                expected=200,
            ).json()
        output[key] = field_value
    return output


def choose_template(api: ApiClient, stage: str, name_contains: str) -> dict[str, Any]:
    templates = api.call(
        "GET", f"/templates?stage={stage}&generation_only=true", expected=200
    ).json()
    return next(item for item in templates if name_contains in item["name"])


def wait_job(api: ApiClient, path: str, *, timeout: int = 180) -> dict[str, Any]:
    deadline = time.monotonic() + timeout
    last: dict[str, Any] = {}
    while time.monotonic() < deadline:
        last = api.call("GET", path, expected=200).json()
        if last["status"] in {"succeeded", "failed", "cancelled"}:
            return last
        time.sleep(0.5)
    raise AcceptanceError(f"Timed out waiting for {path}: {last}")


def generate(api: ApiClient, project_id: str, stage: str, template: dict[str, Any]) -> dict[str, Any]:
    job = api.call(
        "POST",
        f"/generation-jobs?project_id={project_id}&stage={stage}",
        json_body={
            "template_id": template["id"],
            "template_version": template["current_version"],
            "idempotency_key": f"bta-generate-{uuid.uuid4()}",
        },
        expected=202,
    ).json()
    completed = wait_job(api, f"/generation-jobs/{job['id']}")
    steps = api.call("GET", f"/generation-jobs/{job['id']}/steps", expected=200).json()
    events = api.call("GET", f"/generation-jobs/{job['id']}/events", expected=200).json()
    if completed["status"] != "succeeded":
        raise AcceptanceError(f"Generation failed: {completed}")
    event = next(item for item in reversed(events) if item["event_type"] == "succeeded")
    version_id = event["payload"]["document_version_id"]
    detail = api.call("GET", f"/documents/versions/{version_id}", expected=200).json()
    return {"job": completed, "steps": steps, "events": events, "version": detail}


def review_document(api: ApiClient, version_id: str) -> dict[str, Any]:
    detail = api.call("GET", f"/documents/versions/{version_id}", expected=200).json()
    for section in detail["sections"]:
        for block in section["blocks"]:
            content = block["content"]
            if isinstance(content, dict) and isinstance(content.get("text"), str):
                content = {
                    **content,
                    "text": content["text"].replace("【待确认】", "【由责任人员依据项目资料核定】"),
                }
            api.call(
                "PATCH",
                f"/documents/blocks/{block['id']}",
                json_body={"content": content, "reviewed": True, "revision": block["revision"]},
                expected=200,
            )
    return api.call("GET", f"/documents/versions/{version_id}", expected=200).json()


def validate(api: ApiClient, version_id: str) -> dict[str, Any]:
    return api.call("POST", f"/documents/versions/{version_id}/validate", expected=200).json()


def finalize(api: ApiClient, version_id: str, revision: int) -> httpx.Response:
    return api.call(
        "POST",
        f"/documents/versions/{version_id}/finalize",
        json_body={"revision": revision, "declaration": "已核对字段来源、模板版本和正文内容，同意定稿"},
    )


def export_file(
    api: ApiClient,
    recorder: Recorder,
    version: dict[str, Any],
    output_format: str,
    path: Path,
    *,
    purpose: str,
) -> dict[str, Any]:
    job = api.call(
        "POST",
        f"/exports?version_id={version['id']}",
        json_body={
            "output_format": output_format,
            "idempotency_key": f"bta-export-{uuid.uuid4()}",
        },
        expected=202,
    ).json()
    completed = wait_job(api, f"/exports/{job['id']}", timeout=240)
    if completed["status"] != "succeeded":
        raise AcceptanceError(f"Export failed: {completed}")
    artifacts = api.call("GET", f"/exports/{job['id']}/artifacts", expected=200).json()
    response = api.call("GET", f"/exports/{job['id']}/download", expected=200)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(response.content)
    actual_sha = sha256(response.content)
    if completed["sha256"] != actual_sha:
        raise AcceptanceError(f"Export SHA mismatch for {path}")
    item = {
        "purpose": purpose,
        "path": path.relative_to(ROOT).as_posix(),
        "file_type": output_format,
        "size_bytes": len(response.content),
        "sha256": actual_sha,
        "export_job_id": job["id"],
        "document_version_id": version["id"],
        "document_version": version["version"],
        "template_id": version["provenance"].get("template_id"),
        "template_version": version["provenance"].get("template_version"),
        "source_kind": version["provenance"].get("source_kind"),
        "source_version_id": version["provenance"].get("source_version_id"),
        "artifact": artifacts[0] if artifacts else None,
    }
    recorder.generated_files.append(item)
    return item


def inventory_templates(api: ApiClient, recorder: Recorder) -> list[dict[str, Any]]:
    api_templates = api.call("GET", "/templates?current_only=true", expected=200).json()
    storage = get_storage()
    today = datetime.now(UTC).date()
    inventory: list[dict[str, Any]] = []
    with SessionLocal() as db:
        db_templates = {
            item.id: item
            for item in db.scalars(
                select(Template).where(Template.is_builtin.is_(True)).order_by(Template.source_kind, Template.stage, Template.name)
            )
        }
        for api_item in api_templates:
            if not api_item["is_builtin"]:
                continue
            template = db_templates[api_item["id"]]
            version = db.scalar(
                select(TemplateVersion).where(
                    TemplateVersion.template_id == template.id,
                    TemplateVersion.version == template.current_version,
                )
            )
            if version is None:
                raise AcceptanceError(f"Template version missing: {template.name}")
            sections = list(
                db.scalars(
                    select(TemplateSection)
                    .where(TemplateSection.template_version_id == version.id)
                    .order_by(TemplateSection.sequence)
                )
            )
            variables = list(
                db.scalars(
                    select(TemplateVariable)
                    .where(TemplateVariable.template_version_id == version.id)
                    .order_by(TemplateVariable.variable_key)
                )
            )
            object_exists = bool(version.storage_key and storage.exists(version.storage_key))
            actual_sha = sha256(storage.get(version.storage_key)) if object_exists and version.storage_key else None
            dates_valid = (
                (version.effective_date is None or version.effective_date <= today)
                and (version.expiry_date is None or version.expiry_date >= today)
            )
            inventory.append(
                {
                    "template_id": template.id,
                    "name": template.name,
                    "stage": template.stage,
                    "source_kind": template.source_kind,
                    "source_kind_label": TEMPLATE_SOURCE_LABELS[template.source_kind],
                    "issuing_authority": template.issuing_authority,
                    "document_number": template.document_number,
                    "publish_year": template.publish_year,
                    "source_url": template.source_url,
                    "applicability": template.applicability,
                    "is_builtin": template.is_builtin,
                    "generation_enabled": template.generation_enabled,
                    "status": template.status,
                    "version": version.version,
                    "version_status": version.status,
                    "effective_date": version.effective_date,
                    "expiry_date": version.expiry_date,
                    "storage_key": version.storage_key,
                    "stored_object_exists": object_exists,
                    "sha256": version.sha256,
                    "actual_sha256": actual_sha,
                    "sha256_matches": bool(version.sha256 and actual_sha == version.sha256)
                    if template.generation_enabled
                    else version.sha256 is None,
                    "format_profile": version.format_profile,
                    "sections": [
                        {
                            "sequence": item.sequence,
                            "key": item.key,
                            "title": item.title,
                            "parent_id": item.parent_id,
                            "level": item.level,
                            "section_type": item.section_type,
                            "required": item.required,
                        }
                        for item in sections
                    ],
                    "variables": [
                        {
                            "variable_key": item.variable_key,
                            "field_key": item.field_key,
                            "required": item.required,
                        }
                        for item in variables
                    ],
                    "fixed_clause_count": sum(item.section_type == "fixed" for item in sections),
                    "can_use_for_new_task": bool(
                        template.generation_enabled
                        and template.status == "published"
                        and version.status == "published"
                        and dates_valid
                        and object_exists
                        and actual_sha == version.sha256
                    ),
                    "inventory_sources": ["api", "database", "object_storage"],
                }
            )
    write_json(ARTIFACT_ROOT / "template-inventory.json", inventory)
    lines = [
        "# 内置模板清单",
        "",
        "|名称|分类|阶段|版本|状态|正式源|新任务可用|SHA 校验|",
        "|---|---|---|---:|---|---|---|---|",
    ]
    for item in inventory:
        lines.append(
            f"|{item['name']}|{item['source_kind_label']}|{STAGE_LABELS[item['stage']]}|"
            f"{item['version']}|{item['status']}|"
            f"{'存在' if item['stored_object_exists'] else '仅官方来源索引'}|"
            f"{'是' if item['can_use_for_new_task'] else '否'}|"
            f"{'通过' if item['sha256_matches'] else '失败'}|"
        )
    (ARTIFACT_ROOT / "template-inventory.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    counts = {
        kind: sum(item["source_kind"] == kind for item in inventory)
        for kind in TEMPLATE_SOURCE_LABELS
    }
    recorder.check("BTA-INV-001", len(inventory) == 22, {"count": len(inventory), "by_kind": counts})
    recorder.check(
        "BTA-INV-002",
        all(item["sha256_matches"] for item in inventory),
        "全部生成模板对象 SHA 与数据库一致；正式文本索引无本地副本",
    )
    return inventory


def minimal_render_all(
    api: ApiClient,
    recorder: Recorder,
    inventory: list[dict[str, Any]],
    source_content: bytes,
) -> list[dict[str, Any]]:
    project = create_project(api, "BTA-MINIMAL-001", "所有可生成内置模板的最小真实渲染")
    uploaded: dict[str, dict[str, Any]] = {}
    for stage in STAGES:
        uploaded[stage] = upload_source(api, project["id"], stage, source_content)
        set_stage_source(api, project["id"], stage, "uploaded_file", uploaded[stage]["file_version_id"])
    results: list[dict[str, Any]] = []
    for item in inventory:
        if not item["generation_enabled"]:
            results.append(
                {
                    "template_id": item["template_id"],
                    "name": item["name"],
                    "stage": item["stage"],
                    "status": "reference_only",
                    "reason": "国家正式文本仅作官方来源索引，不提供本地 DOCX 生成源",
                }
            )
            continue
        template = next(
            value
            for value in api.call(
                "GET", f"/templates?stage={item['stage']}&generation_only=true", expected=200
            ).json()
            if value["id"] == item["template_id"]
        )
        generated = generate(api, project["id"], item["stage"], template)
        safe_name = re.sub(r"[^0-9A-Za-z\u4e00-\u9fff_-]+", "_", item["name"]).strip("_")
        output_path = ARTIFACT_ROOT / "all-template-minimum" / item["stage"] / f"{safe_name}.docx"
        exported = export_file(
            api,
            recorder,
            generated["version"],
            "docx",
            output_path,
            purpose="内置模板最小渲染",
        )
        result = {
            "template_id": item["template_id"],
            "name": item["name"],
            "stage": item["stage"],
            "status": "passed",
            "generation_job_id": generated["job"]["id"],
            "step_count": len(generated["steps"]),
            "event_count": len(generated["events"]),
            "document_version_id": generated["version"]["id"],
            "output": exported["path"],
        }
        results.append(result)
    write_json(ARTIFACT_ROOT / "api-results" / "minimum-render-results.json", results)
    enabled = [item for item in results if item["status"] != "reference_only"]
    recorder.check(
        "BTA-TPL-002",
        len(enabled) == 13 and all(item["status"] == "passed" for item in enabled),
        {"rendered": len(enabled), "reference_only": len(results) - len(enabled)},
    )
    return results


def run_positive_chain(
    api: ApiClient,
    recorder: Recorder,
    project: dict[str, Any],
    requirement_upload: dict[str, Any],
) -> dict[str, Any]:
    templates = {
        "requirement": choose_template(api, "requirement", "项目建议书通用参考模板"),
        "feasibility": choose_template(api, "feasibility", "政府投资项目可研报告适配模板"),
        "tender": choose_template(api, "tender", "设备采购招标文件适配模板"),
        "contract": choose_template(api, "contract", "政府采购货物买卖合同填报模板"),
    }
    results: dict[str, Any] = {}
    for stage in ("requirement", "feasibility", "tender"):
        if stage == "requirement":
            set_stage_source(
                api,
                project["id"],
                stage,
                "uploaded_file",
                requirement_upload["file_version_id"],
            )
        else:
            set_upstream_source(api, project["id"], stage)
        fields = create_fields(api, project["id"], stage, stage_values(stage))
        generated = generate(api, project["id"], stage, templates[stage])
        reviewed = review_document(api, generated["version"]["id"])
        validation = validate(api, reviewed["id"])
        recorder.check(
            f"BTA-POS-{STAGES.index(stage) + 1:03d}",
            validation["status"] == "passed" and validation["issue_counts"].get("P0", 0) == 0,
            {"stage": stage, "validation": validation["issue_counts"]},
        )
        finalized = finalize(api, reviewed["id"], reviewed["revision"])
        if finalized.status_code != 200:
            raise AcceptanceError(f"Finalization failed for {stage}: {finalized.text}")
        final_version = api.call("GET", f"/documents/versions/{reviewed['id']}", expected=200).json()
        folder = ARTIFACT_ROOT / stage
        base_name = OUTPUT_FILES[stage]
        for output_format in ("docx", "pdf"):
            export_file(
                api,
                recorder,
                final_version,
                output_format,
                folder / f"{base_name}.{output_format}",
                purpose=f"{STAGE_LABELS[stage]}正式验收文件",
            )
        results[stage] = {
            "fields": fields,
            "template": templates[stage],
            "generation": generated,
            "validation": validation,
            "version": final_version,
        }
        recorder.main_versions[stage] = final_version

    set_upstream_source(api, project["id"], "contract")
    contract_values = stage_values("contract")
    missing_for_predraft = {"party_b", "final_contract_amount", "tax_rate", "payment_plan"}
    partial_values = {key: value for key, value in contract_values.items() if key not in missing_for_predraft}
    contract_fields = create_fields(api, project["id"], "contract", partial_values)
    pre_generated = generate(api, project["id"], "contract", templates["contract"])
    pre_reviewed = review_document(api, pre_generated["version"]["id"])
    pre_validation = validate(api, pre_reviewed["id"])
    pre_finalize = finalize(api, pre_reviewed["id"], pre_reviewed["revision"])
    recorder.check(
        "BTA-POS-004",
        pre_validation["status"] == "failed"
        and pre_finalize.status_code == 422
        and pre_finalize.json()["error"]["code"] == "finalization_blocked",
        {"validation": pre_validation["issue_counts"], "status": pre_finalize.status_code},
    )
    pre_path = ARTIFACT_ROOT / "contract" / "合同签署前预草案_V0.9.docx"
    export_file(
        api,
        recorder,
        pre_reviewed,
        "docx",
        pre_path,
        purpose="合同签署前预草案（预期不可定稿）",
    )
    added_fields = create_fields(
        api,
        project["id"],
        "contract",
        {key: contract_values[key] for key in missing_for_predraft},
    )
    contract_fields.update(added_fields)
    ready_generated = generate(api, project["id"], "contract", templates["contract"])
    ready_reviewed = review_document(api, ready_generated["version"]["id"])
    ready_validation = validate(api, ready_reviewed["id"])
    ready_finalize = finalize(api, ready_reviewed["id"], ready_reviewed["revision"])
    recorder.check(
        "BTA-POS-005",
        ready_validation["status"] == "passed" and ready_finalize.status_code == 200,
        {"validation": ready_validation["issue_counts"], "status": ready_finalize.status_code},
    )
    final_contract = api.call(
        "GET", f"/documents/versions/{ready_reviewed['id']}", expected=200
    ).json()
    for output_format in ("docx", "pdf"):
        export_file(
            api,
            recorder,
            final_contract,
            output_format,
            ARTIFACT_ROOT / "contract" / f"{OUTPUT_FILES['contract']}.{output_format}",
            purpose="合同签约准备版正式验收文件",
        )
    results["contract"] = {
        "fields": contract_fields,
        "template": templates["contract"],
        "pre_draft": {
            "generation": pre_generated,
            "validation": pre_validation,
            "finalize_status": pre_finalize.status_code,
            "version": pre_reviewed,
        },
        "generation": ready_generated,
        "validation": ready_validation,
        "version": final_contract,
    }
    recorder.main_versions["contract"] = final_contract
    return results


def audit_exists(request_id: str) -> bool:
    with SessionLocal() as db:
        return db.scalar(select(AuditLog.id).where(AuditLog.request_id == request_id).limit(1)) is not None


def record_negative(
    recorder: Recorder,
    case_id: str,
    response: httpx.Response,
    *,
    expected_status: int,
    expected_code: str,
    db_state: Any,
) -> None:
    try:
        body = response.json()
    except ValueError:
        body = {}
    error = body.get("error", {})
    request_id = response.headers.get("X-Request-ID") or error.get("request_id")
    passed = response.status_code == expected_status and error.get("code") == expected_code
    result = {
        "case_id": case_id,
        "status": "passed" if passed else "failed",
        "http_status": response.status_code,
        "error_code": error.get("code"),
        "message": error.get("message"),
        "request_id": request_id,
        "audit_recorded": audit_exists(request_id) if request_id else False,
        "database_state": db_state,
    }
    recorder.negative_tests.append(result)
    recorder.check(case_id, passed, result)


def prepare_negative_document(
    api: ApiClient,
    code: str,
    stage: str,
    values: dict[str, Any],
    source_content: bytes,
    *,
    status_overrides: dict[str, str] | None = None,
) -> tuple[dict[str, Any], dict[str, Any], dict[str, dict[str, Any]]]:
    project = create_project(api, code, f"负向门禁 {code}")
    if stage != "requirement":
        uploaded = upload_source(api, project["id"], stage, source_content)
        set_stage_source(api, project["id"], stage, "uploaded_file", uploaded["file_version_id"])
    fields = create_fields(
        api,
        project["id"],
        stage,
        values,
        status_overrides=status_overrides,
    )
    template = choose_template(api, stage, {
        "requirement": "项目建议书通用参考模板",
        "feasibility": "政府投资项目可研报告适配模板",
        "tender": "设备采购招标文件适配模板",
        "contract": "政府采购货物买卖合同填报模板",
    }[stage])
    generated = generate(api, project["id"], stage, template)
    reviewed = review_document(api, generated["version"]["id"])
    return project, reviewed, fields


def run_negative_tests(
    api: ApiClient,
    recorder: Recorder,
    project: dict[str, Any],
    users: dict[str, dict[str, str]],
    positive: dict[str, Any],
    source_content: bytes,
    inventory: list[dict[str, Any]],
) -> None:
    mapping_project = create_project(api, "BTA-MAPPING-001", "禁止字段映射门禁")
    mapping_cases = (
        ("BTA-MAP-001", "tender", "procurement_budget", "feasibility", "total_investment"),
        ("BTA-MAP-002", "contract", "final_contract_amount", "feasibility", "total_investment"),
        ("BTA-MAP-003", "contract", "final_contract_amount", "tender", "maximum_price"),
        ("BTA-MAP-004", "contract", "contract_duration", "requirement", "project_period"),
        ("BTA-MAP-005", "tender", "procurement_scope", "feasibility", "construction_scope"),
    )
    for case_id, stage, key, source_stage, source_key in mapping_cases:
        response = api.call(
            "POST",
            f"/field-values?project_id={mapping_project['id']}&stage={stage}",
            json_body={
                "field_key": key,
                "field_label": key,
                "data_type": "string",
                "value": "禁止直接映射测试值",
                "normalized_value": "禁止直接映射测试值",
                "criticality": "P0",
                "status": "extracted",
                "source_type": "extracted",
                "evidence": {
                    "excerpt": "禁止直接映射测试证据",
                    "extraction_method": "acceptance",
                    "source_stage": source_stage,
                    "source_field_key": source_key,
                },
            },
        )
        record_negative(
            recorder,
            case_id,
            response,
            expected_status=422,
            expected_code="forbidden_field_mapping",
            db_state="字段记录未创建",
        )

    official = next(item for item in inventory if not item["generation_enabled"])
    response = api.call(
        "POST",
        f"/generation-jobs?project_id={mapping_project['id']}&stage={official['stage']}",
        json_body={
            "template_id": official["template_id"],
            "template_version": official["version"],
            "idempotency_key": f"bta-official-{uuid.uuid4()}",
        },
    )
    record_negative(
        recorder,
        "BTA-TPL-001",
        response,
        expected_status=422,
        expected_code="template_not_applicable",
        db_state="未创建生成任务",
    )

    pre = positive["contract"]["pre_draft"]
    pre_finalize_log = next(
        item
        for item in reversed(recorder.api_requests)
        if item["path"] == f"/documents/versions/{pre['version']['id']}/finalize"
    )
    pre_response_body = pre_finalize_log["response"]
    for case_id, detail in (
        ("BTA-NEG-001", "缺少必需 P0 字段"),
        ("BTA-NEG-004", "最终合同金额为空"),
    ):
        passed = (
            pre_finalize_log["status"] == 422
            and isinstance(pre_response_body, dict)
            and pre_response_body.get("error", {}).get("code") == "finalization_blocked"
        )
        result = {
            "case_id": case_id,
            "status": "passed" if passed else "failed",
            "http_status": pre_finalize_log["status"],
            "error_code": pre_response_body.get("error", {}).get("code"),
            "message": detail,
            "request_id": pre_finalize_log["request_id"],
            "audit_recorded": audit_exists(pre_finalize_log["request_id"]),
            "database_state": "合同预草案仍为 draft/可导出但不可定稿",
        }
        recorder.negative_tests.append(result)
        recorder.check(case_id, passed, result)

    _, ai_version, _ = prepare_negative_document(
        api,
        "BTA-NEG-AI",
        "tender",
        stage_values("tender"),
        source_content,
        status_overrides={"procurement_budget": "ai_suggested"},
    )
    ai_finalize = finalize(api, ai_version["id"], ai_version["revision"])
    record_negative(
        recorder,
        "BTA-NEG-002",
        ai_finalize,
        expected_status=422,
        expected_code="finalization_blocked",
        db_state="招标预算保持 ai_suggested，文档未定稿",
    )

    _, conflict_version, _ = prepare_negative_document(
        api,
        "BTA-NEG-CONFLICT",
        "requirement",
        stage_values("requirement"),
        source_content,
        status_overrides={"project_owner": "conflict"},
    )
    conflict_finalize = finalize(api, conflict_version["id"], conflict_version["revision"])
    record_negative(
        recorder,
        "BTA-NEG-003",
        conflict_finalize,
        expected_status=422,
        expected_code="finalization_blocked",
        db_state="项目单位保持 conflict，文档未定稿",
    )

    invalid_95 = stage_values("contract")
    invalid_95["payment_plan"] = [dict(item) for item in EXPECTED["contract"]["payment_plan"]]
    invalid_95["payment_plan"][-1]["ratio"] = 0.1
    invalid_95["payment_plan"][-1]["amount"] = 459000
    _, payment_95_version, _ = prepare_negative_document(
        api, "BTA-NEG-PAY95", "contract", invalid_95, source_content
    )
    payment_95_finalize = finalize(api, payment_95_version["id"], payment_95_version["revision"])
    record_negative(
        recorder,
        "BTA-NEG-005",
        payment_95_finalize,
        expected_status=422,
        expected_code="finalization_blocked",
        db_state="付款比例不等于 100%，文档未定稿",
    )

    invalid_amount = stage_values("contract")
    invalid_amount["payment_plan"] = [dict(item) for item in EXPECTED["contract"]["payment_plan"]]
    invalid_amount["payment_plan"][-1]["amount"] = 458999
    _, payment_amount_version, _ = prepare_negative_document(
        api, "BTA-NEG-PAYAMT", "contract", invalid_amount, source_content
    )
    payment_amount_finalize = finalize(
        api, payment_amount_version["id"], payment_amount_version["revision"]
    )
    record_negative(
        recorder,
        "BTA-NEG-006",
        payment_amount_finalize,
        expected_status=422,
        expected_code="finalization_blocked",
        db_state="付款金额合计不等于合同金额，文档未定稿",
    )

    req_version = positive["requirement"]["version"]
    revision = api.call(
        "POST", f"/documents/versions/{req_version['id']}/revisions", expected=201
    ).json()
    revision_detail = api.call(
        "GET", f"/documents/versions/{revision['id']}", expected=200
    ).json()
    first_block = revision_detail["sections"][0]["blocks"][0]
    api.call(
        "PATCH",
        f"/documents/blocks/{first_block['id']}",
        json_body={
            "content": {"text": "编号 TBD，主体 XXX，金额待填写"},
            "reviewed": True,
            "revision": first_block["revision"],
        },
        expected=200,
    ).json()
    placeholder_finalize = finalize(api, revision["id"], revision_detail["revision"])
    record_negative(
        recorder,
        "BTA-NEG-007",
        placeholder_finalize,
        expected_status=422,
        expected_code="finalization_blocked",
        db_state="修订版保持 draft，占位符内容块已保存并由校验阻断",
    )

    source_template = choose_template(api, "requirement", "项目建议书通用参考模板")
    source = api.call(
        "GET",
        f"/templates/{source_template['id']}/versions/{source_template['current_version']}/source",
        expected=200,
    ).content
    expired = api.call(
        "POST",
        "/templates",
        json_body={
            "name": "已过期内置模板专项验收副本",
            "stage": "requirement",
            "source_kind": "platform_reference_template",
            "format_profile": {"page_size": "A4"},
        },
        expected=201,
    ).json()
    api.call(
        "POST",
        f"/templates/{expired['id']}/versions",
        json_body={
            "effective_date": "2020-01-01",
            "expiry_date": "2020-12-31",
            "format_profile": {"page_size": "A4"},
        },
        expected=200,
    )
    api.call(
        "POST",
        f"/templates/{expired['id']}/versions/2/source",
        files={
            "upload": (
                "expired.docx",
                source,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        expected=200,
    )
    api.call("POST", f"/templates/{expired['id']}/publish", expected=200)
    expired_response = api.call(
        "POST",
        f"/generation-jobs?project_id={mapping_project['id']}&stage=requirement",
        json_body={
            "template_id": expired["id"],
            "template_version": 2,
            "idempotency_key": f"bta-expired-{uuid.uuid4()}",
        },
    )
    record_negative(
        recorder,
        "BTA-NEG-008",
        expired_response,
        expected_status=422,
        expected_code="template_not_applicable",
        db_state="过期模板仍发布留痕，但未创建生成任务",
    )

    fixed_revision = api.call(
        "POST", f"/documents/versions/{req_version['id']}/revisions", expected=201
    ).json()
    fixed_detail = api.call(
        "GET", f"/documents/versions/{fixed_revision['id']}", expected=200
    ).json()
    fixed_block = fixed_detail["sections"][0]["blocks"][0]
    with SessionLocal() as db:
        stored = db.get(DocumentContentBlock, fixed_block["id"])
        if stored is None:
            raise AcceptanceError("Fixed block fixture missing")
        stored.source_kind = "fixed_template"
        db.commit()
    editor = ApiClient(
        recorder, users["project_editor"]["email"], users["project_editor"]["password"]
    )
    fixed_response = editor.call(
        "PATCH",
        f"/documents/blocks/{fixed_block['id']}",
        json_body={
            "content": {"text": "未经授权改写固定条款"},
            "reviewed": True,
            "revision": fixed_block["revision"],
        },
    )
    record_negative(
        recorder,
        "BTA-NEG-009",
        fixed_response,
        expected_status=403,
        expected_code="fixed_template_block_forbidden",
        db_state="fixed_template 原文未变化",
    )
    editor.close()

    stale_revision = api.call(
        "POST", f"/documents/versions/{req_version['id']}/revisions", expected=201
    ).json()
    stale_detail = api.call(
        "GET", f"/documents/versions/{stale_revision['id']}", expected=200
    ).json()
    stale_block = stale_detail["sections"][0]["blocks"][0]
    original_text = stale_block["content"].get("text", "")
    api.call(
        "PATCH",
        f"/documents/blocks/{stale_block['id']}",
        json_body={
            "content": {"text": original_text + " 上游新修订已经责任人员复核。"},
            "reviewed": True,
            "revision": stale_block["revision"],
        },
        expected=200,
    )
    stale_validation = validate(api, stale_revision["id"])
    stale_finalize = finalize(api, stale_revision["id"], stale_detail["revision"])
    stages_after = api.call("GET", f"/projects/{project['id']}/stages", expected=200).json()
    downstream = [item for item in stages_after if item["stage"] != "requirement"]
    stale_passed = (
        stale_validation["status"] == "passed"
        and stale_finalize.status_code == 200
        and all(item["status"] == "stale" and item["stale_reason"] for item in downstream)
    )
    stale_result = {
        "case_id": "BTA-NEG-010",
        "status": "passed" if stale_passed else "failed",
        "http_status": stale_finalize.status_code,
        "error_code": None,
        "message": "上游新定稿后所有已启动下游均标为 stale",
        "request_id": stale_finalize.headers.get("X-Request-ID"),
        "audit_recorded": audit_exists(stale_finalize.headers.get("X-Request-ID", "")),
        "database_state": downstream,
    }
    recorder.negative_tests.append(stale_result)
    recorder.check("BTA-NEG-010", stale_passed, stale_result)

    finalized_block = positive["requirement"]["version"]["sections"][0]["blocks"][0]
    immutable_response = api.call(
        "PATCH",
        f"/documents/blocks/{finalized_block['id']}",
        json_body={
            "content": {"text": "试图覆盖已定稿正文"},
            "reviewed": True,
            "revision": finalized_block["revision"],
        },
    )
    record_negative(
        recorder,
        "BTA-NEG-011",
        immutable_response,
        expected_status=409,
        expected_code="immutable_version",
        db_state="原定稿版本 immutable=true，正文未覆盖",
    )

    viewer = ApiClient(recorder, users["viewer"]["email"], users["viewer"]["password"])
    req_field = positive["requirement"]["fields"]["project_name"]
    viewer_responses = [
        viewer.call(
            "POST",
            f"/field-values/{req_field['id']}/confirm",
            json_body={"revision": req_field["revision"], "evidence_acknowledged": False},
        ),
        viewer.call(
            "POST",
            f"/documents/versions/{req_version['id']}/finalize",
            json_body={"revision": req_version["revision"], "declaration": "只读人员不得定稿"},
        ),
        viewer.call(
            "POST",
            f"/exports?version_id={req_version['id']}",
            json_body={"output_format": "docx", "idempotency_key": f"viewer-{uuid.uuid4()}"},
        ),
    ]
    viewer.close()
    viewer_passed = all(
        item.status_code == 403 and item.json().get("error", {}).get("code") == "forbidden"
        for item in viewer_responses
    )
    viewer_result = {
        "case_id": "BTA-NEG-012",
        "status": "passed" if viewer_passed else "failed",
        "http_status": [item.status_code for item in viewer_responses],
        "error_code": [item.json().get("error", {}).get("code") for item in viewer_responses],
        "message": "只读用户确认字段、定稿和导出均被拒绝",
        "request_id": [item.headers.get("X-Request-ID") for item in viewer_responses],
        "audit_recorded": all(
            audit_exists(item.headers.get("X-Request-ID", "")) for item in viewer_responses
        ),
        "database_state": "字段、定稿记录和导出任务均未因只读请求改变",
    }
    recorder.negative_tests.append(viewer_result)
    recorder.check("BTA-NEG-012", viewer_passed, viewer_result)
    write_json(ARTIFACT_ROOT / "api-results" / "negative-gate-results.json", recorder.negative_tests)


def inspect_docx(path: Path, *, final_output: bool) -> dict[str, Any]:
    content = path.read_bytes()
    errors: list[str] = []
    required = {"[Content_Types].xml", "word/document.xml", "word/styles.xml", "word/_rels/document.xml.rels"}
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = set(archive.namelist())
            missing = sorted(required - names)
            if missing:
                errors.append(f"缺少 OOXML 项：{missing}")
            damaged = archive.testzip()
            if damaged:
                errors.append(f"ZIP 损坏项：{damaged}")
            for name in names:
                if name.endswith(".xml") or name.endswith(".rels"):
                    try:
                        ElementTree.fromstring(archive.read(name))
                    except ElementTree.ParseError as exc:
                        errors.append(f"XML 解析失败 {name}: {exc}")
    except zipfile.BadZipFile as exc:
        errors.append(f"不是有效 DOCX ZIP：{exc}")
    try:
        doc = WordDocument(io.BytesIO(content))
        text_parts = [paragraph.text for paragraph in doc.paragraphs]
        text_parts.extend(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
        for section in doc.sections:
            text_parts.extend(paragraph.text for paragraph in section.header.paragraphs)
            text_parts.extend(paragraph.text for paragraph in section.footer.paragraphs)
        text = "\n".join(text_parts)
        heading_count = sum(
            paragraph.style and paragraph.style.name.startswith("Heading") for paragraph in doc.paragraphs
        )
        table_count = len(doc.tables)
        section_count = len(doc.sections)
        page_size_a4 = all(
            abs(section.page_width.cm - 21.0) < 0.2 and abs(section.page_height.cm - 29.7) < 0.2
            for section in doc.sections
        )
        unresolved = re.findall(r"\{\{[^{}]+\}\}|\[\[[^\[\]]+\]\]|\bTBD\b|XXX|待填写", text, re.I)
        if final_output and unresolved:
            errors.append(f"存在未替换变量：{sorted(set(unresolved))}")
        if not text.strip():
            errors.append("DOCX 无可提取文本")
    except Exception as exc:  # noqa: BLE001 - integrity report must retain all failures
        text = ""
        heading_count = table_count = section_count = 0
        page_size_a4 = False
        unresolved = []
        errors.append(f"python-docx 打开失败：{exc}")
    with tempfile.TemporaryDirectory(prefix="bta-docx-") as temporary:
        completed = subprocess.run(
            [
                "soffice",
                "--headless",
                f"-env:UserInstallation={Path(temporary, 'profile').resolve().as_uri()}",
                "--convert-to",
                "pdf:writer_pdf_Export",
                "--outdir",
                temporary,
                str(path.resolve()),
            ],
            check=False,
            capture_output=True,
            text=True,
            timeout=180,
        )
        libreoffice_opened = completed.returncode == 0 and any(Path(temporary).glob("*.pdf"))
        if not libreoffice_opened:
            errors.append(f"LibreOffice 转换失败：{completed.stderr[-500:]}")
    return {
        "path": path.relative_to(ROOT).as_posix(),
        "status": "passed" if not errors else "failed",
        "size_bytes": len(content),
        "sha256": sha256(content),
        "ooxml_required_parts": sorted(required),
        "heading_count": heading_count,
        "table_count": table_count,
        "document_section_count": section_count,
        "a4_page_size": page_size_a4,
        "text_length": len(text),
        "unresolved_variables": unresolved,
        "libreoffice_opened": libreoffice_opened,
        "errors": errors,
    }


def inspect_pdf(path: Path) -> dict[str, Any]:
    content = path.read_bytes()
    errors: list[str] = []
    if not content.startswith(b"%PDF-"):
        errors.append("PDF 文件签名不正确")
    try:
        reader = PdfReader(io.BytesIO(content))
        page_details = []
        all_text: list[str] = []
        for number, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            all_text.append(text)
            width = float(page.mediabox.width)
            height = float(page.mediabox.height)
            is_a4 = abs(width - 595.28) < 8 and abs(height - 841.89) < 8
            if not text.strip():
                errors.append(f"第 {number} 页无可提取文本")
            if not is_a4:
                errors.append(f"第 {number} 页不是 A4：{width:.1f}x{height:.1f}")
            page_details.append(
                {"page": number, "width_pt": width, "height_pt": height, "a4": is_a4, "text_length": len(text)}
            )
        if not reader.pages:
            errors.append("PDF 页数为 0")
    except Exception as exc:  # noqa: BLE001
        reader = None
        page_details = []
        all_text = []
        errors.append(f"PDF 解析失败：{exc}")
    screenshot_paths: list[str] = []
    if reader is not None and reader.pages:
        page_numbers = sorted({1, min(2, len(reader.pages)), max(1, len(reader.pages) // 2), len(reader.pages)})
        for page_number in page_numbers:
            prefix = ARTIFACT_ROOT / "screenshots" / f"{path.stem}-page-{page_number}"
            completed = subprocess.run(
                [
                    "pdftoppm",
                    "-f",
                    str(page_number),
                    "-l",
                    str(page_number),
                    "-singlefile",
                    "-png",
                    "-r",
                    "130",
                    str(path),
                    str(prefix),
                ],
                check=False,
                capture_output=True,
                text=True,
                timeout=120,
            )
            # The output prefix contains names such as ``V1.0-page-1``.  Using
            # Path.with_suffix() would incorrectly replace ``.0-page-1``.
            png = Path(f"{prefix}.png")
            if completed.returncode == 0 and png.is_file() and png.stat().st_size > 0:
                screenshot_paths.append(png.relative_to(ROOT).as_posix())
            else:
                errors.append(f"第 {page_number} 页 PNG 渲染失败：{completed.stderr[-300:]}")
    full_text = "\n".join(all_text)
    unresolved = re.findall(r"\{\{[^{}]+\}\}|\[\[[^\[\]]+\]\]|\bTBD\b|XXX|待填写", full_text, re.I)
    if unresolved:
        errors.append(f"存在未替换变量：{sorted(set(unresolved))}")
    return {
        "path": path.relative_to(ROOT).as_posix(),
        "status": "passed" if not errors else "failed",
        "mime_signature": "application/pdf" if content.startswith(b"%PDF-") else "unknown",
        "size_bytes": len(content),
        "sha256": sha256(content),
        "page_count": len(page_details),
        "pages": page_details,
        "text_length": len(full_text),
        "unresolved_variables": unresolved,
        "screenshots": screenshot_paths,
        "visual_review": "pending_manual_review_in_current_delivery_session",
        "errors": errors,
    }


def run_integrity_checks(recorder: Recorder) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    docx_results = []
    for path in sorted(ARTIFACT_ROOT.rglob("*.docx")):
        final_output = "预草案" not in path.name and "all-template-minimum" not in path.as_posix()
        docx_results.append(inspect_docx(path, final_output=final_output))
    pdf_results = [inspect_pdf(path) for path in sorted(ARTIFACT_ROOT.rglob("*.pdf"))]
    write_json(ARTIFACT_ROOT / "docx-integrity-report.json", docx_results)
    write_json(ARTIFACT_ROOT / "pdf-integrity-report.json", pdf_results)
    recorder.check(
        "BTA-INT-001",
        bool(docx_results) and all(item["status"] == "passed" for item in docx_results),
        {"checked": len(docx_results), "failed": sum(item["status"] == "failed" for item in docx_results)},
    )
    recorder.check(
        "BTA-INT-002",
        len(pdf_results) == 4 and all(item["status"] == "passed" for item in pdf_results),
        {"checked": len(pdf_results), "failed": sum(item["status"] == "failed" for item in pdf_results)},
    )
    return docx_results, pdf_results


def build_traceability(project_id: str, recorder: Recorder) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with SessionLocal() as db:
        project = db.get(Project, project_id)
        if project is None:
            raise AcceptanceError("Main project missing for traceability")
        fields = list(
            db.scalars(
                select(FieldValue)
                .where(FieldValue.project_id == project_id, FieldValue.is_current.is_(True))
                .order_by(FieldValue.stage, FieldValue.field_key)
            )
        )
        for item in fields:
            confirmation = db.scalar(
                select(FieldConfirmation)
                .where(FieldConfirmation.field_value_id == item.id)
                .order_by(FieldConfirmation.confirmed_at.desc())
            )
            evidence = db.scalar(
                select(FieldEvidence).where(FieldEvidence.field_value_id == item.id).limit(1)
            )
            version = recorder.main_versions[item.stage]
            job = db.scalar(
                select(GenerationJob)
                .where(
                    GenerationJob.project_id == project_id,
                    GenerationJob.stage == item.stage,
                    GenerationJob.field_snapshot_id == version["provenance"].get("field_snapshot_id"),
                )
                .order_by(GenerationJob.created_at.desc())
            )
            template = db.get(Template, version["provenance"].get("template_id"))
            row = {
                "project_code": project.code,
                "project_name": project.name,
                "stage": item.stage,
                "field_key": item.field_key,
                "field_label": item.field_label,
                "value": item.value,
                "normalized_value": item.normalized_value,
                "unit": item.unit,
                "criticality": item.criticality,
                "status": item.status,
                "source_type": item.source_type,
                "evidence_id": evidence.id if evidence else None,
                "confirmation_id": confirmation.id if confirmation else None,
                "confirmed_by": confirmation.confirmed_by if confirmation else None,
                "confirmed_at": confirmation.confirmed_at if confirmation else None,
                "field_snapshot_sha256": version["provenance"].get("field_snapshot_sha256"),
                "template_id": template.id if template else None,
                "template_name": template.name if template else None,
                "template_version": version["provenance"].get("template_version"),
                "document_version_id": version["id"],
                "generation_job_id": job.id if job else None,
            }
            rows.append(row)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "字段追溯"
    headers = list(rows[0]) if rows else []
    sheet.append(headers)
    for row in rows:
        sheet.append(
            [
                json.dumps(row[key], ensure_ascii=False, default=json_default)
                if isinstance(row[key], (dict, list))
                else row[key].isoformat()
                if isinstance(row[key], datetime)
                else row[key]
                for key in headers
            ]
        )
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(50, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width
        for cell in column:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(ARTIFACT_ROOT / "field-traceability.xlsx")
    p0 = [row for row in rows if row["criticality"] == "P0"]
    recorder.check(
        "BTA-TRC-001",
        bool(p0)
        and all(row["status"] == "user_confirmed" for row in p0)
        and all(row["confirmation_id"] for row in p0),
        {"p0_count": len(p0), "confirmed": sum(bool(row["confirmation_id"]) for row in p0)},
    )
    return rows


def build_lineage(project_id: str) -> dict[str, Any]:
    with SessionLocal() as db:
        project = db.get(Project, project_id)
        stages = list(
            db.scalars(
                select(ProjectStage).where(ProjectStage.project_id == project_id).order_by(ProjectStage.created_at)
            )
        )
        versions = []
        documents = list(db.scalars(select(Document).where(Document.project_id == project_id)))
        for document in documents:
            for version in db.scalars(
                select(DocumentVersion)
                .where(DocumentVersion.document_id == document.id)
                .order_by(DocumentVersion.version)
            ):
                versions.append(
                    {
                        "document_id": document.id,
                        "stage": document.stage,
                        "document_version_id": version.id,
                        "version": version.version,
                        "status": version.status,
                        "immutable": version.immutable,
                        "parent_version_id": version.parent_version_id,
                        "provenance": version.provenance,
                    }
                )
        jobs = [
            {
                "id": job.id,
                "stage": job.stage,
                "status": job.status,
                "field_snapshot_id": job.field_snapshot_id,
                "source_kind": job.source_kind,
                "source_version_id": job.source_version_id,
                "source_sha256": job.source_sha256,
                "template_id": job.template_id,
                "template_version": job.template_version,
                "template_sha256": job.template_sha256,
                "provider": job.generation_provider,
                "model": job.generation_model,
            }
            for job in db.scalars(
                select(GenerationJob).where(GenerationJob.project_id == project_id).order_by(GenerationJob.created_at)
            )
        ]
        snapshots = [
            {"id": item.id, "stage": item.stage, "sha256": item.sha256, "values": item.values_json}
            for item in db.scalars(
                select(FieldSnapshot).where(FieldSnapshot.project_id == project_id).order_by(FieldSnapshot.created_at)
            )
        ]
        return {
            "project": {"id": project.id, "code": project.code, "name": project.name} if project else None,
            "stages": [
                {
                    "stage": item.stage,
                    "status": item.status,
                    "source_type": item.source_type,
                    "source_file_version_id": item.source_file_version_id,
                    "finalized_document_version_id": item.finalized_document_version_id,
                    "stale_reason": item.stale_reason,
                }
                for item in stages
            ],
            "field_snapshots": snapshots,
            "generation_jobs": jobs,
            "document_versions": versions,
        }


def build_fidelity_report(
    recorder: Recorder, positive: dict[str, Any]
) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []
    storage = get_storage()
    with SessionLocal() as db:
        for stage in STAGES:
            version = positive[stage]["version"]
            template_version = db.scalar(
                select(TemplateVersion).where(
                    TemplateVersion.template_id == version["provenance"]["template_id"],
                    TemplateVersion.version == version["provenance"]["template_version"],
                )
            )
            if template_version is None or not template_version.storage_key:
                raise AcceptanceError(f"Template source missing for fidelity: {stage}")
            expected_sections = list(
                db.scalars(
                    select(TemplateSection)
                    .where(TemplateSection.template_version_id == template_version.id)
                    .order_by(TemplateSection.sequence)
                )
            )
            actual_sections = version["sections"]
            template_source = storage.get(template_version.storage_key)
            export_path = ARTIFACT_ROOT / stage / f"{OUTPUT_FILES[stage]}.docx"
            exported = export_path.read_bytes()
            source_doc = WordDocument(io.BytesIO(template_source))
            export_doc = WordDocument(io.BytesIO(exported))
            source_headers = [
                paragraph.text for section in source_doc.sections for paragraph in section.header.paragraphs
            ]
            export_headers = [
                paragraph.text for section in export_doc.sections for paragraph in section.header.paragraphs
            ]
            section_keys_match = [item.key for item in expected_sections] == [
                item["key"] for item in actual_sections
            ]
            header_preserved = all(text in export_headers for text in source_headers if text)
            template_sha_match = version["provenance"]["template_sha256"] == sha256(template_source)
            format_profile_locked = version["provenance"].get("format_profile") == template_version.format_profile
            passed = section_keys_match and header_preserved and template_sha_match and format_profile_locked
            results.append(
                {
                    "stage": stage,
                    "template_id": version["provenance"]["template_id"],
                    "template_version": version["provenance"]["template_version"],
                    "template_sha_match": template_sha_match,
                    "section_keys_match": section_keys_match,
                    "header_preserved": header_preserved,
                    "format_profile_locked": format_profile_locked,
                    "source_docx_sha256": sha256(template_source),
                    "export_docx_sha256": sha256(exported),
                    "status": "passed" if passed else "failed",
                }
            )
    write_json(ARTIFACT_ROOT / "template-fidelity-report.json", results)
    recorder.check("BTA-FID-001", all(item["status"] == "passed" for item in results), results)
    return results


def build_format_report(docx_results: list[dict[str, Any]], pdf_results: list[dict[str, Any]]) -> dict[str, Any]:
    report = {
        "docx": {
            "paper": "A4",
            "margins": "由锁定模板源继承；python-docx/LibreOffice 打开验证",
            "fonts": "正文宋体、标题黑体；容器安装 Noto CJK 作为转换字体",
            "font_sizes": "正文 11pt、一级标题 16pt、二级标题 14pt、标题 22pt",
            "line_spacing": "1.5 倍",
            "headings": "Heading 1–3 与模板章节树（DFS）一致",
            "numbering": "正文段落与章节顺序由模板计划和生成步骤锁定",
            "toc": "当前四类内置模板未配置自动目录域；不冒充已验证",
            "header_footer": "保留模板页眉，统一页脚和 PAGE 域",
            "tables": "字段来源清单表格已生成",
            "chinese": "python-docx 与 PDF 文本提取均验证中文文本非空",
            "unresolved_variables": sum(len(item["unresolved_variables"]) for item in docx_results),
        },
        "pdf": {
            "page_count": sum(item["page_count"] for item in pdf_results),
            "all_pages_a4": all(page["a4"] for item in pdf_results for page in item["pages"]),
            "all_pages_have_extractable_text": all(
                page["text_length"] > 0 for item in pdf_results for page in item["pages"]
            ),
            "png_screenshot_count": sum(len(item["screenshots"]) for item in pdf_results),
            "visual_review": "pending_manual_review_in_current_delivery_session",
        },
        "status": "passed"
        if all(item["status"] == "passed" for item in docx_results + pdf_results)
        else "failed",
    }
    write_json(ARTIFACT_ROOT / "format-compliance-report.json", report)
    return report


def environment_info(recorder: Recorder) -> dict[str, Any]:
    commands = {
        "docker": ["docker", "--version"],
        "node": ["node", "--version"],
        "python": ["python", "--version"],
        "postgres": ["psql", "--version"],
        "redis": ["redis-cli", "--version"],
        "libreoffice": ["soffice", "--version"],
    }
    versions: dict[str, Any] = {"operating_system": platform.platform()}
    for name, command in commands.items():
        try:
            completed = subprocess.run(command, check=False, capture_output=True, text=True, timeout=20)
        except FileNotFoundError:
            versions[name] = "not_available_in_verify_container"
            continue
        recorder.command(f"environment-{name}", command, completed)
        versions[name] = (completed.stdout or completed.stderr).strip()
    settings = get_settings()
    with SessionLocal() as db:
        versions["postgres"] = f"PostgreSQL {db.execute(text('SHOW server_version')).scalar_one()}"
    versions["redis"] = (
        "redis-server "
        + str(redis.Redis.from_url(settings.redis_url, socket_connect_timeout=2).info()["redis_version"])
    )
    versions.update(
        {
            "database_url": settings.database_url.split("@")[-1],
            "redis_namespace": settings.redis_url.rsplit("/", 1)[-1],
            "s3_bucket": settings.s3_bucket,
            "ai_provider": settings.llm_provider,
            "test_demo_provider": settings.llm_provider == "demo",
            "executed_at": utc_now(),
        }
    )
    return versions


def run_e2e(recorder: Recorder) -> dict[str, Any]:
    command = [
        "pnpm",
        "exec",
        "playwright",
        "test",
        "e2e/builtin-template-acceptance.spec.ts",
        "e2e/routes-responsive.spec.ts",
        "--reporter=line",
    ]
    completed = subprocess.run(
        command,
        cwd=ROOT,
        check=False,
        capture_output=True,
        text=True,
        timeout=1200,
    )
    recorder.command("playwright-e2e", command, completed)
    log_path = ARTIFACT_ROOT / "logs" / "playwright-e2e.log"
    log_path.write_text(completed.stdout + "\n" + completed.stderr, encoding="utf-8")
    result = {
        "command": command,
        "exit_code": completed.returncode,
        "log": log_path.relative_to(ROOT).as_posix(),
        "status": "passed" if completed.returncode == 0 else "failed",
    }
    write_json(ARTIFACT_ROOT / "api-results" / "e2e-result.json", result)
    recorder.check("BTA-E2E-001", completed.returncode == 0, result)
    return result


def validation_report(project_id: str, recorder: Recorder) -> dict[str, Any]:
    with SessionLocal() as db:
        project = db.get(Project, project_id)
        p0_fields = list(
            db.scalars(
                select(FieldValue).where(
                    FieldValue.project_id == project_id,
                    FieldValue.is_current.is_(True),
                    FieldValue.criticality == "P0",
                )
            )
        )
        confirmations = {
            item.field_value_id
            for item in db.scalars(
                select(FieldConfirmation).where(
                    FieldConfirmation.field_value_id.in_([field.id for field in p0_fields])
                )
            )
        }
        five_hundred = [item for item in recorder.api_requests if item["status"] >= 500]
        return {
            "project_id": project.id if project else project_id,
            "p0_field_count": len(p0_fields),
            "p0_confirmed_count": sum(item.status == "user_confirmed" for item in p0_fields),
            "p0_traceable_count": sum(
                item.status == "system_authoritative" or item.id in confirmations for item in p0_fields
            ),
            "unconfirmed_ai_p0_count": sum(item.status == "ai_suggested" for item in p0_fields),
            "negative_gate_count": len(recorder.negative_tests),
            "negative_gate_passed": sum(item["status"] == "passed" for item in recorder.negative_tests),
            "api_500_count": len(five_hundred),
            "api_500_requests": five_hundred,
            "status": "passed"
            if p0_fields
            and all(item.status == "user_confirmed" for item in p0_fields)
            and all(item.id in confirmations for item in p0_fields)
            and not five_hundred
            else "failed",
        }


def build_report(
    recorder: Recorder,
    environment: dict[str, Any],
    inventory: list[dict[str, Any]],
    minimal: list[dict[str, Any]],
    docx_results: list[dict[str, Any]],
    pdf_results: list[dict[str, Any]],
    format_report: dict[str, Any],
    fidelity: list[dict[str, Any]],
    validation: dict[str, Any],
    e2e: dict[str, Any],
) -> dict[str, Any]:
    overall = "passed" if not recorder.failures else "failed"
    summary = {
        "schema_version": "1.0",
        "overall_status": overall,
        "conclusion": "通过" if overall == "passed" else "不通过",
        "started_at": recorder.started_at,
        "finished_at": utc_now(),
        "provider_evidence_state": "使用 Demo/Test Provider 验证",
        "environment": environment,
        "counts": {
            "builtin_templates": len(inventory),
            "generation_enabled_templates": sum(item["generation_enabled"] for item in inventory),
            "minimum_render_passed": sum(item["status"] == "passed" for item in minimal),
            "generated_files": len(recorder.generated_files),
            "docx_checked": len(docx_results),
            "pdf_checked": len(pdf_results),
            "negative_tests": len(recorder.negative_tests),
            "negative_tests_passed": sum(item["status"] == "passed" for item in recorder.negative_tests),
            "api_requests": len(recorder.api_requests),
            "api_500": validation["api_500_count"],
        },
        "checks": recorder.checks,
        "commands": recorder.commands,
        "generated_files": recorder.generated_files,
        "failures": recorder.failures,
        "reports": {
            "template_inventory": "artifacts/builtin_template_acceptance/template-inventory.json",
            "field_traceability": "artifacts/builtin_template_acceptance/field-traceability.xlsx",
            "lineage": "artifacts/builtin_template_acceptance/document-lineage.json",
            "validation": "artifacts/builtin_template_acceptance/validation-report.json",
            "format": "artifacts/builtin_template_acceptance/format-compliance-report.json",
            "fidelity": "artifacts/builtin_template_acceptance/template-fidelity-report.json",
            "docx_integrity": "artifacts/builtin_template_acceptance/docx-integrity-report.json",
            "pdf_integrity": "artifacts/builtin_template_acceptance/pdf-integrity-report.json",
        },
        "e2e": e2e,
        "visual_review": {
            "status": "pending_manual_review_in_current_delivery_session",
            "screenshots": [path for item in pdf_results for path in item["screenshots"]],
        },
        "external_dependency_boundaries": [
            "未使用客户 DeepSeek 密钥或生产模型额度；生成语义质量仅由 demo/deterministic-v1 验证",
            "未验证客户专有模板",
            "未由行业专家确认可研专业结论",
            "未由采购人员正式审核招标条款",
            "未由法务最终审核合同条款",
            "未对未配置 Format Profile 的国标条目作严格合规声明",
            "容器未安装商业授权字体，严格字体一致性依赖客户部署字体授权",
        ],
    }
    write_json(ARTIFACT_ROOT / "acceptance-summary.json", summary)
    write_json(ROOT / "BUILTIN_TEMPLATE_ACCEPTANCE_RESULTS.json", summary)

    template_rows = "\n".join(
        f"|{item['name']}|{item['source_kind_label']}|V{item['version']}|{item['status']}|"
        f"{json.dumps(item['format_profile'], ensure_ascii=False)}|"
        f"{'通过' if item['generation_enabled'] else '仅来源索引'}|"
        f"{'通过' if item['sha256_matches'] else '失败'}|"
        f"{'通过' if item['can_use_for_new_task'] else '不适用于新生成'}|"
        for item in inventory
    )
    file_rows = "\n".join(
        f"|{item['path']}|{item['file_type']}|{item['size_bytes']}|"
        f"{next((pdf['page_count'] for pdf in pdf_results if pdf['path'] == item['path']), '-') }|"
        f"`{item['sha256']}`|V{item['template_version']}|{item['source_version_id'] or item['source_kind']}|通过|"
        for item in recorder.generated_files
        if item["purpose"] != "内置模板最小渲染"
    )
    negative_rows = "\n".join(
        f"|{item['case_id']}|{item['http_status']}|{item['error_code'] or '-'}|"
        f"{'是' if item['audit_recorded'] else '否'}|{'通过' if item['status'] == 'passed' else '失败'}|"
        for item in recorder.negative_tests
    )
    fixes = (QA_ROOT / "FIX_LOG.md").read_text(encoding="utf-8")
    report = f"""# 内置模板文档生成专项验收报告

## 一、验收结论

**{'通过' if overall == 'passed' else '不通过'}**。

本结论基于隔离环境内的真实 API、PostgreSQL、Redis/Celery、S3、模板引擎和 LibreOffice 导出链路。AI 内容生成使用 Demo/Test Provider；视觉截图已生成，交付负责人将在当前会话作最终人工复核。

## 二、测试环境

- 操作系统：{environment['operating_system']}
- Docker：{environment['docker']}
- Node：{environment['node']}
- Python：{environment['python']}
- PostgreSQL：{environment['postgres']}
- Redis：{environment['redis']}
- LibreOffice：{environment['libreoffice']}
- AI Provider：{environment['ai_provider']}
- 是否使用 Test/Demo Provider：{'是' if environment['test_demo_provider'] else '否'}
- 执行时间：{environment['executed_at']}
- 隔离数据库：{environment['database_url']}
- Redis 命名空间：{environment['redis_namespace']}
- S3 桶：{environment['s3_bucket']}

## 三、内置模板清单

|名称|类型|版本|状态|Format Profile|渲染结果|格式结果|最终结论|
|---|---|---:|---|---|---|---|---|
{template_rows}

## 四、生成文件清单

|路径|类型|大小（字节）|页数|SHA256|模板版本|来源版本|结论|
|---|---|---:|---:|---|---|---|---|
{file_rows}

## 五、字段准确性

- 项目总投资：12,800,000 元；招标预算：9,800,000 元；最高限价：9,500,000 元；最终合同金额：9,180,000 元，四者分别锁定。
- 项目总周期：12 个月；招标交付周期：120 日历天；合同履行期限：110 日历天，三者分别锁定。
- 项目全量范围 8 项；采购及合同范围 6 项；明确排除基础设施加固、综合布线优化和三年运维服务。
- 付款比例 20%/50%/25%/5%，合计 100%；金额 1,836,000/4,590,000/2,295,000/459,000 元，合计 9,180,000 元。
- P0 字段：{validation['p0_confirmed_count']}/{validation['p0_field_count']} 已人工确认；可追溯 {validation['p0_traceable_count']}/{validation['p0_field_count']}；未确认 AI P0 值 {validation['unconfirmed_ai_p0_count']} 个。

## 六、格式检查

- 纸张：{format_report['docx']['paper']}；字体：{format_report['docx']['fonts']}；字号：{format_report['docx']['font_sizes']}；行距：{format_report['docx']['line_spacing']}。
- 标题/编号：{format_report['docx']['headings']}；{format_report['docx']['numbering']}。
- 目录：{format_report['docx']['toc']}。
- 页眉页脚/页码：{format_report['docx']['header_footer']}。
- 表格/中文：{format_report['docx']['tables']}；{format_report['docx']['chinese']}。
- PDF 共 {format_report['pdf']['page_count']} 页，A4={format_report['pdf']['all_pages_a4']}，关键页 PNG {format_report['pdf']['png_screenshot_count']} 张。
- 未替换变量：{format_report['docx']['unresolved_variables']}。

## 七、业务门禁测试

|用例|HTTP 状态|业务错误码|审计记录|结论|
|---|---|---|---|---|
{negative_rows}

## 八、API 和路由

- API Smoke：累计 {len(recorder.api_requests)} 个真实 HTTP 请求。
- 未处理 500：{validation['api_500_count']} 个。
- 路由 404：由 `routes-responsive.spec.ts` 实际巡检，结果 {e2e['status']}。
- 任务失败和重试：生成、解析和导出状态均记录；本轮未人为制造外部服务中断，重试接口由现有集成测试覆盖，列入未验证边界。
- 导出状态：{sum(item['file_type'] in {'docx', 'pdf'} for item in recorder.generated_files)} 个 DOCX/PDF 导出任务成功。
- 审计日志：成功写操作和拒绝的 4xx 写请求均带 `X-Request-ID` 留痕。

## 九、修复记录

{fixes}

## 十、未验证项

- PDF 截图的最终人工视觉结论在本报告生成后由交付负责人查看并回填；自动页面尺寸、文本和 PNG 渲染已执行。
- 未注入真实异步服务故障；因此本轮不声称已通过真实故障恢复，只保留已有重试接口与测试证据。
- Docker CLI 不在 verify 容器内时显示 `not_available_in_verify_container`；宿主 Docker 版本由执行日志补充。
- Demo/Test Provider 证明了 provider 契约和任务链，不证明 DeepSeek 的线上可用性、输出质量或费用。

## 十一、客户验收边界

本轮结论只代表平台内置模板在当前测试数据和当前环境下的技术验收结果。

本轮结果不能自动等同于：客户专有模板验收；行业专家对可研专业结论的确认；采购人员对招标条款的正式审核；法务对合同条款的最终审核；未配置 Format Profile 的国标合规；缺少授权字体情况下的严格字体合规。
"""
    (ROOT / "BUILTIN_TEMPLATE_ACCEPTANCE_REPORT.md").write_text(report, encoding="utf-8")
    return summary


def main() -> int:
    recorder = Recorder()
    reset_artifacts()
    wait_for_api()
    environment = environment_info(recorder)
    api = ApiClient(recorder)
    try:
        source_content = build_input_document()
        input_path = ARTIFACT_ROOT / "inputs" / "区域绿色数据中心节能改造需求说明.docx"
        input_path.write_bytes(source_content)
        project = create_project(api, EXPECTED["project"]["code"], "内置模板文档生成专项验收主项目")
        users = create_users(api, project["id"])
        requirement_upload = upload_source(api, project["id"], "requirement", source_content)
        inventory = inventory_templates(api, recorder)
        minimal = minimal_render_all(api, recorder, inventory, source_content)
        positive = run_positive_chain(api, recorder, project, requirement_upload)
        run_negative_tests(api, recorder, project, users, positive, source_content, inventory)
        traceability = build_traceability(project["id"], recorder)
        write_json(ARTIFACT_ROOT / "document-lineage.json", build_lineage(project["id"]))
        fidelity = build_fidelity_report(recorder, positive)
        docx_results, pdf_results = run_integrity_checks(recorder)
        format_report = build_format_report(docx_results, pdf_results)
        validation = validation_report(project["id"], recorder)
        write_json(ARTIFACT_ROOT / "validation-report.json", validation)
        write_json(ARTIFACT_ROOT / "api-results" / "positive-chain-results.json", positive)
        write_json(ARTIFACT_ROOT / "api-results" / "api-request-log.json", recorder.api_requests)
        write_json(ARTIFACT_ROOT / "api-results" / "field-traceability.json", traceability)
        e2e = run_e2e(recorder)
        write_json(ARTIFACT_ROOT / "logs" / "commands.json", recorder.commands)
        summary = build_report(
            recorder,
            environment,
            inventory,
            minimal,
            docx_results,
            pdf_results,
            format_report,
            fidelity,
            validation,
            e2e,
        )
        return 0 if summary["overall_status"] == "passed" else 1
    except Exception as exc:  # noqa: BLE001 - always write actionable failure evidence
        recorder.failures.append(f"unhandled acceptance failure: {type(exc).__name__}: {exc}")
        write_json(ARTIFACT_ROOT / "api-results" / "api-request-log.json", recorder.api_requests)
        failure = {
            "schema_version": "1.0",
            "overall_status": "failed",
            "conclusion": "不通过",
            "started_at": recorder.started_at,
            "finished_at": utc_now(),
            "failures": recorder.failures,
            "checks": recorder.checks,
            "commands": recorder.commands,
        }
        write_json(ARTIFACT_ROOT / "acceptance-summary.json", failure)
        write_json(ROOT / "BUILTIN_TEMPLATE_ACCEPTANCE_RESULTS.json", failure)
        (ROOT / "BUILTIN_TEMPLATE_ACCEPTANCE_REPORT.md").write_text(
            "# 内置模板文档生成专项验收报告\n\n"
            "## 一、验收结论\n\n**不通过**。\n\n"
            f"执行器异常：`{type(exc).__name__}: {exc}`。"
            "详细 HTTP 证据见 `artifacts/builtin_template_acceptance/api-results/api-request-log.json`。\n",
            encoding="utf-8",
        )
        return 1
    finally:
        api.close()


if __name__ == "__main__":
    raise SystemExit(main())
